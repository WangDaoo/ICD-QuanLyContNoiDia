from __future__ import annotations

from pathlib import Path
from docx import Document


ROOT = Path(r"D:\Project\Đồ Án 4 +Mobile")
SOURCE = ROOT / "12523040 - Le Tien Linh - Bao cao Mobile ICD.docx"
OUTPUT = ROOT / "10123264 - Dao Minh Quang - Bao cao Mobile ICD.docx"

REPLACEMENTS = (
    ("ĐỒ ÁN 4", "PHÁT TRIỂN ỨNG DỤNG MOBILE ĐA NỀN TẢNG"),
    ("Đồ án 4", "Phát triển ứng dụng mobile đa nền tảng"),
    ("LÊ TIẾN LINH", "ĐÀO MINH QUANG"),
    ("Lê Tiến Linh", "Đào Minh Quang"),
    ("12523040", "10123264"),
    ("TS. HOÀNG QUỐC VIỆT", "VŨ XUÂN THẮNG"),
    ("TS. Hoàng Quốc Việt", "Vũ Xuân Thắng"),
    (
        "Xây dựng ứng dụng di động hỗ trợ vận hành ICD",
        "Xây dựng hệ thống quản lý kho bãi container nội địa (ICD)",
    ),
    (
        "ỨNG DỤNG DI ĐỘNG HỖ TRỢ VẬN HÀNH ICD",
        "HỆ THỐNG QUẢN LÝ KHO BÃI CONTAINER NỘI ĐỊA (ICD)",
    ),
)


def edit_paragraph(paragraph):
    for run in paragraph.runs:
        text = run.text
        for old, new in REPLACEMENTS:
            text = text.replace(old, new)
        if text != run.text:
            run.text = text


def edit_container(container):
    for paragraph in container.paragraphs:
        edit_paragraph(paragraph)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                edit_container(cell)


doc = Document(SOURCE)
edit_container(doc)
for section in doc.sections:
    edit_container(section.header)
    edit_container(section.footer)

# Keep the original academic cover composition, with the requested title on two lines.
assert doc.paragraphs[9].text == "XÂY DỰNG ỨNG DỤNG DI ĐỘNG"
assert doc.paragraphs[10].text == "HỖ TRỢ VẬN HÀNH ICD"
doc.paragraphs[9].runs[0].text = "XÂY DỰNG HỆ THỐNG QUẢN LÝ KHO BÃI"
doc.paragraphs[10].runs[0].text = "CONTAINER NỘI ĐỊA (ICD)"

# The old class code belongs to the previous student's cover and was not supplied
# for the new author, so leave that line empty instead of attributing it to them.
assert doc.paragraphs[20].text == "MÃ LỚP: 12523W.4"
for run in doc.paragraphs[20].runs:
    run.text = ""

# A short running header avoids two wrapped lines across the 83-page report.
header = doc.sections[1].header.paragraphs[0]
header.runs[0].text = "PHÁT TRIỂN ỨNG DỤNG MOBILE ĐA NỀN TẢNG"

doc.core_properties.title = "Xây dựng hệ thống quản lý kho bãi container nội địa (ICD)"
doc.core_properties.author = "Đào Minh Quang"
doc.core_properties.subject = "Phát triển ứng dụng mobile đa nền tảng - Mobile ICD"

doc.save(OUTPUT)
print(OUTPUT)
