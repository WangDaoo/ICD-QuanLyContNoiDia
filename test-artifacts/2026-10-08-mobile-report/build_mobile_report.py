from __future__ import annotations

import hashlib
import json
import math
import re
import textwrap
from collections import defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r'D:\Project\Đồ Án 4 +Mobile\icd-management')
WORK = ROOT / 'test-artifacts/2026-10-08-mobile-report'
TEMPLATE = WORK / 'DoAn3-mobile-reference.docx'
OUTPUT = ROOT.parent / '12523040 - Le Tien Linh - Bao cao Mobile ICD.docx'
INV = json.loads((ROOT / 'docs/system/system-inventory.json').read_text(encoding='utf-8'))
DIAGRAM = WORK / 'diagrams'
DIAGRAM.mkdir(exist_ok=True)
SHA = 'FA00C74F131462DDC97D675DE8087F82F84EF4F54ABB2EBEDF11B4164FB148EA'
assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest().upper() == SHA

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(80, 80, 80)
LIGHT = 'EAF1F7'
PALE = 'F7FAFC'
FONT = r'C:\Windows\Fonts\arial.ttf'
FONT_BOLD = r'C:\Windows\Fonts\arialbd.ttf'


def font(size=30, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def place_wrapped(draw, rect, label, size=28, bold=False, fill='#102130'):
    x0,y0,x1,y1 = rect
    f = font(size,bold)
    lines=[]
    for raw in label.split('\n'):
        tokens=raw.split()
        line=''
        for token in tokens:
            candidate=(line+' '+token).strip()
            if draw.textlength(candidate,font=f)>x1-x0-24 and line:
                lines.append(line); line=token
            else: line=candidate
        lines.append(line)
    h=len(lines)*(size+7)
    y=(y0+y1-h)/2
    for line in lines:
        w=draw.textlength(line,font=f)
        draw.text(((x0+x1-w)/2,y),line,font=f,fill=fill)
        y += size+7


def line_arrow(draw, start, end, color='#344b62', width=4):
    draw.line([start,end],fill=color,width=width)
    angle=math.atan2(end[1]-start[1],end[0]-start[0])
    back=16
    left=(end[0]-back*math.cos(angle-0.55),end[1]-back*math.sin(angle-0.55))
    right=(end[0]-back*math.cos(angle+0.55),end[1]-back*math.sin(angle+0.55))
    draw.polygon([end,left,right],fill=color)


def graph(name, nodes, edges, height=650, title=''):
    image=Image.new('RGB',(1600,height),'white')
    d=ImageDraw.Draw(image)
    if title: d.text((35,24),title,font=font(30,True),fill='#1d3850')
    coords={n[0]:(n[2],n[3],n[4],n[5]) for n in nodes}
    for a,b,label in edges:
        aa=coords[a]; bb=coords[b]
        ax=(aa[0]+aa[2])/2; ay=(aa[1]+aa[3])/2
        bx=(bb[0]+bb[2])/2; by=(bb[1]+bb[3])/2
        if abs(ax-bx)>=abs(ay-by):
            p=(aa[2],ay) if bx>ax else (aa[0],ay)
            q=(bb[0],by) if bx>ax else (bb[2],by)
        else:
            p=(ax,aa[3]) if by>ay else (ax,aa[1])
            q=(bx,bb[1]) if by>ay else (bx,bb[3])
        line_arrow(d,p,q)
        if label:
            mx=(p[0]+q[0])/2; my=(p[1]+q[1])/2
            tw=d.textlength(label,font=font(21))
            d.rectangle((mx-tw/2-5,my-30,mx+tw/2+5,my),fill='white')
            d.text((mx-tw/2,my-29),label,font=font(21),fill='#30465d')
    for ident,label,x0,y0,x1,y1 in nodes:
        d.rounded_rectangle((x0,y0,x1,y1),radius=18,outline='#315778',width=4,fill='#eaf3f9')
        place_wrapped(d,(x0,y0,x1,y1),label,size=27,bold=True)
    p=DIAGRAM/f'{name}.png'; image.save(p); return p


def sequence(name, lanes, steps):
    W=1600; H=max(560,145+len(steps)*88)
    image=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(image)
    xs=[110+i*(1380/(len(lanes)-1)) for i in range(len(lanes))]
    for i,(lane,x) in enumerate(zip(lanes,xs)):
        rect=(int(x-140),30,int(x+140),95)
        d.rounded_rectangle(rect,radius=15,outline='#315778',fill='#eaf3f9',width=3)
        place_wrapped(d,rect,lane,size=25,bold=True)
        d.line((x,95,x,H-32),fill='#adbcca',width=3)
    for k,(a,b,label) in enumerate(steps):
        y=155+k*88
        start=(xs[a]+(15 if b>a else -15),y+18); end=(xs[b]-(15 if b>a else -15),y+18)
        line_arrow(d,start,end,width=3)
        f=font(23); maxw=abs(xs[b]-xs[a])-25
        text=label
        while d.textlength(text,font=f)>maxw and len(text)>10: text=text[:-2]
        if text!=label:text=text.rstrip()+'…'
        tx=(xs[a]+xs[b]-d.textlength(text,font=f))/2
        d.text((tx,y-16),text,font=f,fill='#22384d')
    p=DIAGRAM/f'{name}.png';image.save(p);return p


def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader');el.set(qn('w:val'),'true');trPr.append(el)


def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)


def cell_border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); b=tcPr.first_child_found_in('w:tcBorders')
    if b is None: b=OxmlElement('w:tcBorders');tcPr.append(b)
    for edge in ('top','left','bottom','right'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');b.append(el)


def set_cell_margin(cell, v=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar');tcPr.append(mar)
    for edge in ('top','start','bottom','end'):
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');mar.append(el)


def no_split(row):
    trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trPr.append(el)


def field(p,instruction):
    r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    r=p.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' '+instruction+' ';r._r.append(i)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    p.add_run('Cập nhật trường trong Word')
    r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)


doc=Document(TEMPLATE)
body=doc.element.body
cover=[p._p for p in doc.paragraphs[:29]]
for child in list(body):
    if child not in cover and child.tag!=qn('w:sectPr'): body.remove(child)

def replace_cover(i,text):
    p=doc.paragraphs[i]
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:r.text=''
    else:p.add_run(text)

replace_cover(7,'ĐỒ ÁN 4')
replace_cover(9,'XÂY DỰNG ỨNG DỤNG DI ĐỘNG')
replace_cover(10,'HỖ TRỢ VẬN HÀNH ICD')
replace_cover(15,'HƯỚNG ĐỀ TÀI: ỨNG DỤNG DI ĐỘNG')
replace_cover(28,'HƯNG YÊN – 2026')
# The reference cover has six empty spacer paragraphs before the location.
# The new heading renders slightly taller in Word; remove two spacers so its
# bottom border and location stay on the same A4 page.
for ix in (26,27):
    body.remove(cover[ix])

# The original heading styles contain stale automatic numbering. Reuse their
# typography while removing those numbering definitions for this report.
for style_name in ('Heading 1','Heading 2','Heading 3'):
    st=doc.styles[style_name]
    st.font.name='Times New Roman';st.font.color.rgb=BLACK
    st.font.size=Pt(14 if style_name=='Heading 1' else 13)
    st.font.bold=True
    ppr=st.element.pPr
    if ppr is not None:
        for num in ppr.xpath('./w:numPr'):ppr.remove(num)
    st.paragraph_format.space_before=Pt(18 if style_name=='Heading 1' else 10)
    st.paragraph_format.space_after=Pt(8 if style_name=='Heading 1' else 5)
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 2'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Heading 3'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Normal'].font.name='Times New Roman'
doc.styles['Normal'].font.size=Pt(13)
doc.styles['Normal'].font.color.rgb=BLACK
doc.styles['Normal'].paragraph_format.line_spacing=1.5
doc.styles['Normal'].paragraph_format.space_after=Pt(3)
for toc_name in ('TOC 1','TOC 2','TOC 3'):
    if toc_name in doc.styles:
        toc_style=doc.styles[toc_name]
        toc_style.font.name='Times New Roman';toc_style.font.color.rgb=BLACK
        toc_style.font.size=Pt(11)
        toc_style.paragraph_format.space_before=Pt(0)
        toc_style.paragraph_format.space_after=Pt(0)
        toc_style.paragraph_format.line_spacing=1.0
for name in ('HinhVe','Bang'):
    doc.styles[name].font.name='Times New Roman';doc.styles[name].font.size=Pt(11)
    doc.styles[name].font.color.rgb=BLACK
    doc.styles[name].paragraph_format.keep_with_next=False

sec=doc.sections[1]
sec.different_first_page_header_footer=False
for p in sec.header.paragraphs:p.clear()
hp=sec.header.paragraphs[0]; hp.text='ĐỒ ÁN 4  ỨNG DỤNG DI ĐỘNG HỖ TRỢ VẬN HÀNH ICD'
hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in hp.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.font.color.rgb=BLACK
for p in sec.footer.paragraphs:p.clear()
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('Trang ');field(fp,'PAGE')
for r in fp.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.font.color.rgb=BLACK


def para(text='',style='Normal',italic=False,bold=False,center=False,keep=False):
    p=doc.add_paragraph(style=style)
    if text:
        r=p.add_run(text);r.italic=italic;r.bold=bold
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=keep
    if style=='Normal':p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def heading(number,text,level=1):
    if level==1:
        p=para(f'CHƯƠNG {number}  {text.upper()}','Heading 1')
        p.paragraph_format.page_break_before=True
    else:p=para(f'{number} {text}','Heading '+str(level))
    return p


def cap(kind,number,text):
    s='HinhVe' if kind=='Hình' else 'Bang'
    p=para(f'{kind} {number}. {text}',s,center=(kind=='Hình'))
    p.paragraph_format.keep_with_next=True
    return p


def table(number,title,headers,rows,widths=None,font_size=10.5):
    caption=cap('Bảng',number,title)
    if number in ('3.5','C.3'):
        caption.paragraph_format.page_break_before=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=widths or [15.4/len(headers)]*len(headers)
    for i,(h,w) in enumerate(zip(headers,widths)):
        c=t.rows[0].cells[i];c.width=Cm(w);c.text=h;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(c,LIGHT);cell_border(c);set_cell_margin(c,110)
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
            for r in p.runs:r.bold=True;r.font.size=Pt(font_size);r.font.name='Times New Roman'
    set_repeat_header(t.rows[0]);no_split(t.rows[0])
    for idx,values in enumerate(rows):
        cells=t.add_row().cells
        for i,(value,w) in enumerate(zip(values,widths)):
            c=cells[i];c.width=Cm(w);c.text=str(value);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_border(c);set_cell_margin(c,105)
            if idx%2:set_cell_shading(c,PALE)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(font_size)
        no_split(t.rows[-1])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


def figure(number,title,path,note='',width=14.2):
    p=doc.add_paragraph(style='Normal');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.keep_with_next=True
    with Image.open(path) as im:
        actual_width=min(width,15.0*im.width/im.height)
    r=p.add_run();r.add_picture(str(path),width=Cm(actual_width))
    cap('Hình',number,title)
    if note:para(note,italic=True)


def bullet(text):
    p=para('•  '+text)
    p.paragraph_format.left_indent=Cm(.65)
    p.paragraph_format.first_line_indent=Cm(-.4)


def page_title(text):
    p=para(text,'Heading 1');p.paragraph_format.page_break_before=True
    return p


def note(label,text):
    p=doc.add_paragraph(style='Normal')
    r=p.add_run(label+' ');r.bold=True
    p.add_run(text)
    return p


def add_source(lines):
    para('Nguồn đối chiếu: '+lines,italic=True)


FIG={}
FIG['architecture']=graph('architecture',[
    ('device','Ứng dụng Mobile\nExpo / React Native',70,200,340,360),
    ('api','Backend ICD\nNestJS / JWT / RBAC',470,200,790,360),
    ('service','Service nghiệp vụ\nCổng / Bãi / Phiếu',930,200,1220,360),
    ('db','Prisma / MySQL\ntrạng thái bền vững',1320,200,1560,360),
    ('web','Web điều hành\nquản trị / billing',470,440,790,580),
],[('device','api','HTTPS / JSON'),('api','service','validation'),('service','db','transaction'),('web','api','cùng API')],height=660)
FIG['navigation']=graph('navigation',[
    ('root','Khôi phục phiên / Đăng nhập',50,75,420,185),
    ('tabs','Tabs theo quyền / vai trò',570,75,940,185),
    ('gate','Cổng\nnhập và xuất',60,355,320,490),
    ('yard','Bãi và tác nghiệp',345,355,600,490),
    ('survey','Giám định',625,355,880,490),
    ('lookup','Tra cứu',905,355,1160,490),
    ('queue','Việc ca / Monitor',1185,355,1540,490),
    ('more','Thông báo / Tài khoản',1095,75,1530,185),
],[('root','tabs','JWT + quyền'),('tabs','gate',''),('tabs','yard',''),('tabs','survey',''),('tabs','lookup',''),('tabs','queue',''),('tabs','more','')],height=580)
FIG['usecases']=graph('usecases',[
    ('gate','Nhân viên Cổng',40,115,300,225),
    ('yard','Nhân viên Bãi',40,325,300,435),
    ('operator','Điều độ / Quản lý',40,520,300,630),
    ('in','Tiếp nhận container\nGate-in',500,95,840,205),
    ('out','Quét phiếu và\nxác nhận Gate-out',1020,95,1480,205),
    ('assign','Kiểm tra và\nxếp vị trí',500,320,840,430),
    ('operation','Di chuyển / Giám định\nBooking',1020,320,1480,430),
    ('read','Tra cứu / Việc ca\nThông báo',640,560,1180,670),
],[('gate','in',''),('gate','out',''),('yard','assign',''),('yard','operation',''),('operator','read',''),('in','read','dữ liệu'),('assign','read','cập nhật')],height=760)
FIG['container_state']=graph('container_state',[
    ('pending','PENDING',40,230,260,350),
    ('auth','AUTHORIZED',300,230,530,350),
    ('yard','IN_YARD',575,230,780,350),
    ('pass','GATE_PASS_ISSUED',830,230,1205,350),
    ('exited','EXITED',1285,230,1530,350),
    ('cancel','CANCELLED',300,465,560,565),
],[('pending','auth','lệnh hợp lệ'),('auth','yard','Gate-in'),('yard','pass','đủ readiness'),('pass','exited','Gate-out'),('pending','cancel','hủy'),('auth','cancel','hủy')],height=620)
FIG['slot']=graph('slot_assignment',[
    ('select','Chọn container\nIN_YARD',45,80,335,195),
    ('read','Đọc slot / Legend\ntrạng thái hiện tại',445,80,760,195),
    ('check','POST yard/check\nslot đã chọn',875,80,1170,195),
    ('blocked','Từ chối và hiển thị\nblocker / warning',1240,70,1555,185),
    ('confirm','Xác nhận xếp\nMANUAL',875,360,1170,475),
    ('persist','POST yard/assign\nbackend commit',445,360,760,475),
    ('refresh','Đọc vị trí mới\nhiển thị kết quả',45,360,335,475),
],[('select','read',''),('read','check',''),('check','blocked','không đủ'),('check','confirm','đủ'),('confirm','persist',''),('persist','refresh','')],height=560)
FIG['operations']=graph('operations',[
    ('pending','PENDING\nđã yêu cầu',80,220,350,345),
    ('progress','IN_PROGRESS\nđã bắt đầu',510,220,850,345),
    ('done','COMPLETED\nđã hoàn tất',1115,95,1530,215),
    ('cancel','CANCELLED\ncó lý do',1115,420,1530,545),
],[('pending','progress','start'),('progress','done','complete'),('pending','cancel','cancel'),('progress','cancel','cancel theo rule')],height=640)
FIG['erd']=graph('erd_mobile',[
    ('user','User / Role / Permission',30,70,430,185),
    ('visit','ContainerVisit\ntrục nghiệp vụ',605,70,1030,185),
    ('truck','TruckVisit / Reception',1170,70,1570,185),
    ('slot','YardSlot / LocationLog',30,325,430,445),
    ('ops','Movement / Inspection / Booking',605,325,1030,445),
    ('exit','Hold / GatePass / Handover',1170,325,1570,445),
    ('notif','Notification / AuditLog',600,580,1040,685),
],[('user','visit','actor'),('visit','truck','vào cổng'),('visit','slot','vị trí'),('visit','ops','tác nghiệp'),('visit','exit','điều kiện'),('visit','notif','sự kiện')],height=750)
FIG['offline']=graph('offline_read_write',[
    ('open','Mở màn / phục hồi',45,170,305,290),
    ('conn','Kiểm trạng thái API',410,170,720,290),
    ('cache','Đọc cache theo\nuser + API + ICD',850,65,1210,185),
    ('label','Dữ liệu stale có\nmốc cập nhật',1320,65,1560,185),
    ('fresh','Tải dữ liệu / quyền\nmới',850,350,1210,470),
    ('write','Cho phép write sau\nrevalidate',1320,350,1560,470),
],[('open','conn',''),('conn','cache','mất kết nối'),('cache','label',''),('conn','fresh','online'),('fresh','write','')],height=570)
FIG['deployment']=graph('deployment',[
    ('phone','Android Expo Go\napp RN',35,160,360,285),
    ('web','Expo Web preview\nchỉ smoke trình duyệt',35,370,360,500),
    ('api','API NestJS\nJWT / modules',600,250,980,385),
    ('mysql','Prisma / MySQL\ntest riêng',1190,250,1550,385),
],[('phone','api','LAN / API URL'),('web','api','localhost API'),('api','mysql','DB URL')],height=600)
FIG['seq_login']=sequence('sequence_login',['Nhân viên','Ứng dụng','Auth API','MySQL'],[
    (0,1,'Nhập email và mật khẩu'),(1,2,'POST /auth/login'),(2,3,'Đọc user / role / permission'),
    (3,2,'Kết quả xác thực'),(2,1,'JWT và user trong data'),(1,0,'Mở tab được phép')])
FIG['seq_gatein']=sequence('sequence_gatein',['Nhân viên Cổng','Ứng dụng','Gate API','MySQL'],[
    (0,1,'Quét hoặc nhập số container'),(1,2,'GET container / gate-in-context'),
    (2,1,'Visit và chuyến hợp lệ'),(0,1,'Xác nhận seal, cân, tình trạng'),
    (1,2,'POST /containers/:visitId/gate-in'),(2,3,'Ghi reception + event'),
    (3,2,'Commit'),(2,1,'Kết quả từ server'),(1,0,'Biên nhận Gate-in')])
FIG['seq_gateout']=sequence('sequence_gateout',['Nhân viên Cổng','Ứng dụng','Gate API','MySQL'],[
    (0,1,'Quét QR token'),(1,2,'POST /gate-pass/scan'),(2,3,'Kiểm pass và readiness'),
    (3,2,'Blockers / trạng thái'),(2,1,'Thông tin để kiểm tra'),
    (0,1,'Xác nhận khi xe rời cổng'),(1,2,'POST /gate-out'),
    (2,3,'Recheck + commit EXITED / USED'),(3,2,'Commit'),(2,1,'Biên nhận'),(1,0,'Hiển thị kết quả')])
FIG['seq_notification']=sequence('sequence_notification',['Sự kiện nghiệp vụ','Backend','Notification','Ứng dụng'],[
    (0,1,'Gate-in / HOLD / Gate-out commit'),(1,2,'Tạo bản ghi người nhận'),
    (3,2,'GET /notifications/history'),(2,3,'Danh sách và unread count'),
    (3,2,'PATCH /:id/read hoặc read-all'),(2,3,'Đếm và trạng thái mới')])


def frontmatter():
    p=page_title('NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN')
    para('Họ và tên sinh viên: Lê Tiến Linh. Mã sinh viên: 12523040. Đề tài: Xây dựng ứng dụng di động hỗ trợ vận hành ICD.')
    for label in ['Nội dung và phạm vi', 'Kết quả thực hiện', 'Hình thức trình bày', 'Đánh giá chung']:
        para(label+':',bold=True)
        para('..........................................................................................................................')
        para('..........................................................................................................................')
    para('Hưng Yên, ngày ..... tháng ..... năm 2026',center=True)
    para('GIẢNG VIÊN HƯỚNG DẪN',bold=True,center=True)
    para('(Ký và ghi rõ họ tên)',italic=True,center=True)

    page_title('LỜI CAM ĐOAN')
    para('Em là Lê Tiến Linh, mã sinh viên 12523040. Báo cáo trình bày cấu trúc, chức năng và các kết quả kiểm thử của ứng dụng Mobile trong dự án ICD. Những nội dung kế thừa từ tài liệu dự án, mã nguồn và các báo cáo audit được ghi trong phần tài liệu tham khảo. Ảnh chụp, số liệu kiểm thử và trạng thái chức năng được dẫn về bằng chứng tương ứng; nội dung chưa kiểm trên thiết bị được nêu rõ thay vì coi là kết quả đã đạt.')
    para('Em chịu trách nhiệm kiểm tra lại các thông tin trong báo cáo, bổ sung kết quả nghiệm thu trên Android và iOS khi có bằng chứng thực tế, và cập nhật các hình đã thay đổi trước khi nộp bản cuối.')
    para('Hưng Yên, ngày ..... tháng ..... năm 2026',center=True)
    para('SINH VIÊN',center=True,bold=True)
    para('Lê Tiến Linh',center=True)

    page_title('LỜI CẢM ƠN')
    para('Em cảm ơn Khoa Công nghệ thông tin, Trường Đại học Sư phạm Kỹ thuật Hưng Yên và giảng viên hướng dẫn TS. Hoàng Quốc Việt. Phần hướng dẫn, góp ý và điều kiện học tập của thầy cô giúp em có cơ sở để triển khai đề tài, kiểm tra các luồng vận hành và viết lại tài liệu theo hệ thống hiện có.')
    para('Trong quá trình thực hiện, em đã phải đối chiếu nhiều lớp của cùng một nghiệp vụ: màn hình Mobile, API, trạng thái cơ sở dữ liệu và kết quả sau thao tác. Các ca chưa đủ bằng chứng sẽ tiếp tục được kiểm thử trên thiết bị trước khi đánh giá hoàn thiện. Em mong nhận được góp ý để chỉnh sửa báo cáo và sản phẩm.')

    page_title('TÓM TẮT BÁO CÁO')
    para('Đề tài xây dựng ứng dụng di động phục vụ nhân viên hiện trường của hệ thống quản lý cảng cạn ICD. Ứng dụng React Native/Expo truy cập backend NestJS bằng JWT, hiển thị chức năng theo permission và thao tác trên hồ sơ container đang có trong MySQL. Báo cáo đi từ bài toán và phạm vi đến mô hình dữ liệu, ca sử dụng, các màn hình Cổng/Bãi/Giám định/Tra cứu/Việc ca, rồi kiểm thử theo ảnh và dữ liệu audit.')
    para('Bản README cập nhật ngày 04/10/2026 ghi 17 file màn hình Mobile, 53 model Prisma và 205 endpoint cho toàn hệ thống. Các số này là inventory mã nguồn, không phải số chức năng đã nghiệm thu. Audit ngày 02/10 và smoke Expo Web ngày 04/10 cho thấy nhiều tác vụ hiện trường đã làm việc với API và hồ sơ QA. Camera native, TalkBack, cỡ chữ hệ thống và iOS vẫn cần kiểm chứng riêng; báo cáo không dùng ảnh Expo Web để kết luận những hạng mục đó.')
    para('Từ khóa: ICD; container; Gate-in; Gate-out; xếp bãi; React Native; Expo; REST API; kiểm thử nghiệp vụ.')

    page_title('MỤC LỤC')
    p=doc.add_paragraph();field(p,'TOC \\o "1-3" \\h \\z')

    page_title('DANH MỤC CÁC THUẬT NGỮ')
    table('0.1','Thuật ngữ dùng trong báo cáo',['Thuật ngữ','Giải thích ngắn'],[
        ('ICD','Inland Container Depot: cảng cạn, điểm xử lý container nội địa.'),
        ('Gate-in / Gate-out','Tiếp nhận container vào / xác nhận container rời cổng.'),
        ('Truck Visit','Lượt xe ra/vào gắn với một hoặc nhiều Container Visit.'),
        ('Container Visit','Lượt lưu container tại ICD; trục liên kết các trạng thái tác nghiệp.'),
        ('Yard Slot','Ô/vị trí bãi xác định bằng Block, Bay, Row, Tier và mã vị trí.'),
        ('Inspection / Hold','Giám định / trạng thái ngăn một số thao tác khi chưa giải quyết.'),
        ('Gate Pass / Readiness','Phiếu ra cổng / phép kiểm điều kiện ra cổng do backend quyết định.'),
        ('JWT / RBAC','Token xác thực / phân quyền theo role và permission.'),
        ('API / DTO','Giao diện dịch vụ / kiểu dữ liệu nhận hoặc trả của endpoint.'),
        ('Expo Web / Native','Ứng dụng RN chạy trong trình duyệt / ứng dụng chạy trên Android hoặc iOS.'),
    ],[4.1,11.3],10.5)

    page_title('DANH MỤC CÁC HÌNH VẼ')
    field(doc.add_paragraph(),'TOC \\h \\z \\t "HinhVe,1"')
    page_title('DANH MỤC CÁC BẢNG')
    field(doc.add_paragraph(),'TOC \\h \\z \\t "Bang,1"')


frontmatter()


heading('1','TỔNG QUAN VỀ ĐỀ TÀI')
heading('1.1','Bối cảnh vận hành cảng cạn',2)
para('Một container đi qua ICD kéo theo nhiều bản ghi: hồ sơ nhập, lệnh, chuyến xe, biên nhận cổng, vị trí bãi, tác nghiệp nội bộ, kiểm tra điều kiện ra cổng và phiếu ra cổng. Nhân viên tại cổng và trong bãi cần biết ngay container nào đang chờ xử lý, điều kiện nào đang chặn và hành động nào đã được hệ thống ghi nhận. Khi thông tin chỉ nằm trên màn hình điều hành bàn làm việc, quá trình đối chiếu ở hiện trường dễ bị gián đoạn.')
para('Ứng dụng Mobile được xây dựng cho phần công việc ở hiện trường. Nhân viên quét mã hoặc nhập số container, kiểm tra bối cảnh trước khi gửi lệnh, rồi đọc kết quả từ backend. Quyết định về trạng thái, quyền thao tác và điều kiện nghiệp vụ thuộc các service của ICD. Mobile không tự biến một ảnh QR hay thông báo giao diện thành xác nhận giao dịch.')
figure('1.1','Vị trí của ứng dụng Mobile trong hệ thống ICD',FIG['architecture'])
add_source('[1] README mục 1–5; [2] đặc tả Mobile; [3] mã nguồn apps/mobile và apps/api.')

heading('1.2','Mục tiêu và kết quả cần có',2)
para('Đề tài tập trung vào một ứng dụng di động có thể theo dõi việc ca và xử lý các tác vụ chính của nhân viên cổng, bãi và giám định. Mỗi thao tác phải thể hiện rõ hồ sơ đang xử lý, điều kiện được phép, trạng thái đang gửi, kết quả đã xác nhận và cách quay lại khi lỗi. Phân quyền phải dựa trên permission mà backend trả cho phiên hiện tại.')
table('1.1','Mục tiêu gắn với bằng chứng đánh giá',['Mục tiêu','Bằng chứng cần có','Giới hạn đánh giá'],[
    ('Cổng','Quét/nhập container và Gate Pass, Gate-in/Gate-out, biên nhận và trạng thái persisted.','Không xác nhận bằng ảnh camera Web.'),
    ('Bãi','Vị trí và Legend khớp dữ liệu, check slot, assign MANUAL, tác nghiệp movement/booking.','Thuật toán tối ưu vị trí chưa nghiệm thu.'),
    ('Giám định','Loại kiểm tra, ghi chú, kết quả PASS/FAIL/HOLD, lịch sử và ảnh hưởng readiness.','Ảnh đính kèm cần contract media riêng.'),
    ('Tra cứu và việc ca','Danh sách phân trang, chi tiết, nhiệm vụ đúng ID và quyền.','Empty khác lỗi tải hoặc forbidden.'),
    ('Chất lượng','Typecheck/test, API và dữ liệu bền vững, ảnh UI và ma trận quyền.','Native Android/iOS có cổng nghiệm thu riêng.'),
],[3.2,7.5,4.7])

heading('1.3','Đối tượng và phạm vi',2)
para('Đối tượng trực tiếp gồm GATE_STAFF, YARD_STAFF, OPERATOR và MANAGER; ADMIN được dùng cho quản trị và kiểm thử theo permission. Tài khoản AGENT/CONSIGNEE chỉ có thể dùng phần đã được gắn phạm vi dữ liệu hợp lệ. README mô tả customer scope chưa cấu hình cho loại tài khoản external-only, nên báo cáo không trình bày các màn đó như một portal chủ hàng đã hoàn tất.')
para('Phạm vi chính là apps/mobile trong monorepo icd-management. Web và API xuất hiện khi chúng cung cấp dữ liệu hoặc xác định quy tắc nghiệp vụ của Mobile. Những màn quản trị user/role, tạo Manifest, thanh toán và cấu hình bãi nằm ở Web. Các giao dịch Partner qua API key không được thực hiện trong app của nhân viên ICD.')
table('1.2','Ranh giới hệ thống và trạng thái nguồn',['Thành phần','Vai trò với Mobile','Nguồn'],[
    ('Mobile React Native','Màn và thao tác hiện trường; token/theme/cache cục bộ.','apps/mobile, README mục 5'),
    ('NestJS API','Xác thực, permission, validation và transaction.','apps/api, README mục 7'),
    ('Prisma/MySQL','Lưu trạng thái visit, vị trí, tác nghiệp, hold và pass.','schema.prisma, README mục 9'),
    ('Web React','Điều hành và những nghiệp vụ ngoài phạm vi mobile.','apps/web, README mục 4'),
    ('Partner API','Bàn giao server-to-server bằng key và scope.','Partner API Spec v1.7'),
],[3.4,8.0,4.0])

heading('1.4','Phương pháp thực hiện và nguồn bằng chứng',2)
para('Báo cáo đọc README và đặc tả v1.7 để xác định chuỗi nghiệp vụ. Từng luồng Mobile được đối chiếu với navigation, màn TSX, API caller, endpoint backend và model Prisma. Khi nhận xét kết quả kiểm thử, báo cáo dùng ảnh trong audit cùng các file verify sau thao tác. Một ảnh hiển thị form chỉ chứng minh form đã mở; trạng thái giao dịch sau commit cần thêm phản hồi API hoặc dữ liệu đọc lại.')
para('Inventory hệ thống chụp ngày 04/10/2026 ghi 17 màn Mobile, 22 nhóm API có endpoint, 205 endpoint toàn backend, 53 model, 41 enum, 7 role và 56 permission. Các số này có thể đổi sau khi mã nguồn đổi. Chương 4 ghi phạm vi test theo thời điểm đã có bằng chứng thay vì ngoại suy tất cả 205 route đã chạy.')

heading('2','CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ')
heading('2.1','Kiến trúc ứng dụng di động kết nối API',2)
para('Ứng dụng dùng React Native 0.86.3 với Expo 57 và TypeScript. React Navigation 7 quản lý stack/tab; StyleSheet và theme nội bộ định nghĩa màu, kiểu chữ, spacing. Expo Camera xử lý quét mã trên native, DateTimePicker dùng cho thời gian booking, SecureStore lưu thông tin phiên phù hợp, AsyncStorage giữ theme và cache đọc. API client dùng chung gửi request tới NestJS 12.0.3; backend dùng Prisma 7.10.0 với MySQL 8/InnoDB.')
table('2.1','Stack triển khai đọc từ mã nguồn',['Lớp','Công nghệ','Vai trò'],[
    ('Mobile UI','React Native 0.86.3, Expo 57, TypeScript','Màn và điều hướng hiện trường'),
    ('Điều hướng','React Navigation 7','Tab theo quyền, stack Cổng/Bãi, route ID'),
    ('Thiết bị','Expo Camera, Image Picker, DateTimePicker','Quét, chọn ảnh và giờ booking; tùy nền tảng'),
    ('Lưu cục bộ','SecureStore, AsyncStorage','Phiên, theme và cache đọc'),
    ('API','NestJS 12, JWT, DTO, PermissionsGuard','Xác thực và quyết định tác nghiệp'),
    ('Dữ liệu','Prisma 7.10, MySQL 8','Quan hệ và transaction'),
],[2.8,6.5,6.1])
note('Lưu ý tài liệu.', 'Đặc tả Mobile v1.7 có nhắc TanStack Query như phương án cache. Package Mobile đang triển khai không khai báo thư viện này; phần kiến trúc của báo cáo theo mã nguồn hiện tại.')

heading('2.2','Dữ liệu, trạng thái và quyền thao tác',2)
para('Client đọc JSON qua API. Payload nội bộ dùng camelCase; enum backend giữ nguyên khi gửi request. Thời điểm lưu ở UTC và hiển thị theo giờ Việt Nam UTC+7. Các kết quả dạng danh sách có data/meta hoặc envelope đặc thù của một số route bãi, nên caller phải đọc đúng response của endpoint. Dữ liệu tiền ở Prisma Decimal có thể về dưới dạng chuỗi; Mobile không tự tính số tiền phải thu từ bản tóm tắt readiness.')
para('Quyền của user là tập hợp permission active từ các role. Tab dùng điều kiện cho phép hiển thị; trước khi thực hiện lệnh, backend tiếp tục kiểm JWT, permission và state nghiệp vụ. Tên role không cấp quyền ngầm. Ví dụ GATE_STAFF có thể quét phiếu và xác nhận Gate-out nhưng không cấp phiếu; YARD_STAFF được xếp bãi/giám định nhưng không thu tiền.')
figure('2.1','Điều hướng Mobile theo phiên và nhóm quyền',FIG['navigation'])

heading('2.3','Thiết kế cho môi trường kết nối yếu',2)
para('Cache chỉ phục vụ đọc: việc ca giữ tối đa 5 phút, container đã xem 10 phút và profile/permission 30 phút theo README. Khóa cache gắn với user, địa chỉ API và ICD để không hiển thị nhầm dữ liệu sau đổi tài khoản hoặc môi trường. Khi offline, app ghi thời điểm làm mới và chặn write quan trọng. Sau reconnect, ứng dụng đọc lại dữ liệu và quyền trước khi thao tác; giao dịch chưa rõ kết quả không được replay tự động.')
figure('2.2','Luồng đọc cache và mở khóa tác nghiệp sau khi kết nối lại',FIG['offline'])

heading('2.4','Các tiêu chí sử dụng và kiểm thử trên thiết bị',2)
para('Một thao tác ở cổng cần vùng chạm đủ lớn, nhãn rõ, có fallback khi camera bị từ chối và hiển thị trạng thái đang gửi. Với bãi, nhân viên phải phân biệt ô trống, ô đã chiếm và ô bảo trì theo Legend, kể cả khi màu màn hình thay đổi. Lỗi nhập seal, ngày booking hoặc lý do hủy nên xuất hiện gần trường tương ứng. Các hành vi này được kiểm bằng thao tác thực tế, kích thước vùng chạm, thứ tự focus và dữ liệu backend.')
table('2.2','Ma trận chất lượng cần nghiệm thu',['Nhóm','Điều cần quan sát','Bằng chứng phù hợp'],[
    ('Khả năng tiếp cận','Touch target Android ≥48dp, label/state, TalkBack, cỡ chữ hệ thống.','Native hierarchy, ghi nhận lời đọc, ảnh font scale'),
    ('Phản hồi','Loading, pending, success, lỗi có cách xử lý, double-submit.','Video/ảnh và log request'),
    ('Điều hướng','Tab theo quyền, Back đóng dialog, deep link đến ID thật.','Trace màn và route state'),
    ('Mạng','Offline chặn write, cache có nhãn, reconnect revalidate.','Điều kiện mạng và request trace'),
    ('Dữ liệu','Vị trí, trạng thái, hold và pass khớp API/DB.','Response và snapshot kiểm tra'),
],[3.0,7.3,5.1])


heading('3','PHÂN TÍCH VÀ THIẾT KẾ ỨNG DỤNG MOBILE ICD')
heading('3.1','Bài toán và chuỗi nghiệp vụ',2)
para('Container bắt đầu từ hồ sơ do bộ phận điều hành lập. Sau lệnh và chuyến xe, nhân viên cổng tiếp nhận theo visit. Bãi chọn vị trí hợp lệ và ghi lịch sử di chuyển; việc kiểm tra, rút hàng hoặc đóng hàng có thể làm phát sinh blocker. Khi hồ sơ và thanh toán đủ điều kiện, backend cấp Gate Pass trên Web. Nhân viên cổng quét token từ phiếu, xem trạng thái hiện tại và xác nhận xe đã rời ICD. Mobile hỗ trợ những điểm tác nghiệp có mặt tại hiện trường, trong khi dữ liệu giữa các bước vẫn gắn cùng Container Visit.')
figure('3.1','Nhóm tác nhân và ca sử dụng Mobile theo nghiệp vụ',FIG['usecases'])
table('3.1','Tác nhân và giới hạn quyền trong app',['Tác nhân','Nhiệm vụ chính','Điều kiện'],[
    ('GATE_STAFF','Gate-in, scan Gate Pass, Gate-out, tra cứu và việc ca.','gate_in.create / gate_pass.use'),
    ('YARD_STAFF','Xếp slot, chuyển vị trí, giám định, booking, tra cứu.','yard.read / update / move / inspect / booking'),
    ('OPERATOR','Theo dõi việc ca, tra cứu và điều phối theo permission.','Tab hiển thị Việc ca/Tra cứu/Monitor'),
    ('MANAGER','Giám sát, tác nghiệp bãi được cấp, kiểm tra dữ liệu.','Quyền từ backend, không dựa tên role'),
    ('ADMIN','Kiểm thử/điều hành các luồng theo permission active.','Bị ràng buộc state và validation như user khác'),
],[3.1,6.8,5.5])

heading('3.2','Yêu cầu chức năng',2)
functional=[
    ('F01','Đăng nhập, phục hồi phiên và đăng xuất','Phiên chỉ mở tab có quyền; lỗi mạng cho phép thử lại.'),
    ('F02','Tiếp nhận xe và container vào cổng','Tìm chuyến/visit, arrive, nhập seal/cân/tình trạng rồi xác nhận Gate-in.'),
    ('F03','Quét và xác nhận ra cổng','Đọc QR token, recheck pass/readiness, xác nhận Gate-out.'),
    ('F04','Xem bãi và gán vị trí bằng tay','Slot và Legend đồng nhất; check trước assign; hiển thị blocker.'),
    ('F05','Yêu cầu và thực hiện tác nghiệp bãi','Movement, inspection, booking có list/detail/start/complete/cancel.'),
    ('F06','Giám định hiện trường','Loại kiểm tra, notes, severity khi phù hợp và kết quả PASS/FAIL/HOLD.'),
    ('F07','Tra cứu container','Tìm theo số; đọc visit, vị trí, events, hold, readiness/pass/handover theo quyền.'),
    ('F08','Việc ca và điều phối','Lọc nhiệm vụ, số quá hạn, mở đúng entity và quyền.'),
    ('F09','Thông báo','Lịch sử, unread count, đọc một/đọc tất cả, mở đúng record.'),
    ('F10','Tài khoản và kết nối','Xem vai trò/quyền, sáng/tối, logout, nhãn cache và offline.'),
]
table('3.2','Danh mục yêu cầu chức năng',['Mã','Yêu cầu','Điều kiện nghiệm thu'],functional,[1.3,5.3,8.8],10)
heading('3.2.1','Quy tắc nghiệp vụ xuyên suốt',3)
para('Mỗi request có session hợp lệ và ICD đúng phạm vi. Backend từ chối chuyển trạng thái không được phép, ví dụ Gate-out với phiếu hết hạn hoặc container còn Hold. Một lệnh xếp bãi không được ghi chỉ vì UI vừa hiển thị ô AVAILABLE: backend cần kiểm lại ô trong transaction. Điều kiện có thể đổi giữa thời điểm đọc và thời điểm xác nhận, nhất là với phiếu ra cổng và các thao tác cùng visit.')
figure('3.2','Các trạng thái chính của Container Visit',FIG['container_state'])
note('Quy ước sơ đồ.', 'Đường trạng thái ở Hình 3.2 là mô hình tóm tắt. Các điều kiện cụ thể nằm trong service. Xếp bãi và giám định có thể cập nhật vị trí hoặc blocker khi visit vẫn ở IN_YARD. Chương này không tự thêm trạng thái mới vào enum backend.')

heading('3.3','Đặc tả ca sử dụng',2)
para('Mỗi ca sử dụng dưới đây gắn với màn hình hiện có. Nhánh lỗi mô tả kết quả cần hiển thị khi backend hoặc thiết bị không đáp ứng; đây là đặc tả để nghiệm thu, không mặc nhiên coi mọi nhánh đã được test native.')

USECASES=[
('UC01','Đăng nhập và khôi phục phiên','Nhân viên ICD','LoginScreen / AuthProvider','Tài khoản active, API có thể truy cập.','Nhập email/mật khẩu; POST /auth/login; lưu phiên; GET /auth/me khi mở lại; tính tab theo permission.','Sai thông tin báo lỗi; lỗi mạng giữ khả năng thử lại; 401 khi restore xóa phiên cũ.','Mở màn được phép với đúng user/ICD; logout xóa dữ liệu phiên.','/auth/login; /auth/me; /auth/refresh; /auth/logout'),
('UC02','Đánh dấu xe đến','GATE_STAFF có truck_visit.arrive','GateInFormScreen / TruckArrival','Truck Visit SCHEDULED gắn container hợp lệ.','Tìm chuyến/biển số; đối chiếu visit; xác nhận xe có mặt; POST /gate/truck-visits/:id/arrive.','Chuyến sai/hết hiệu lực hoặc thiếu quyền: giữ bối cảnh và nêu lý do.','Truck Visit sang ARRIVED theo phản hồi backend.','GET /gate/truck-visits; POST /gate/truck-visits/:id/arrive'),
('UC03','Tiếp nhận container vào cổng','GATE_STAFF có gate_in.create','GateInScanScreen → GateInFormScreen → GateInSuccessScreen','Visit có lệnh/chuyến phù hợp, chưa có reception.','Quét/nhập container; đọc gate-in-context; nhập seal thực tế, gross weight và condition; xem lại; POST gate-in.','Container sai, seal thiếu hoặc visit đã xử lý: hiển thị lỗi gần trường/luồng; timeout giữ dữ liệu để đối chiếu trước khi gửi lại.','Tạo reception/event, visit IN_YARD; kết quả xuất hiện sau server xác nhận.','GET /containers/:visitId/gate-in-context; POST /containers/:visitId/gate-in'),
('UC04','Quét phiếu ra cổng','GATE_STAFF có gate_pass.use','GatePassScanScreen','Token phiếu hợp lệ và còn thời hạn.','Mở camera hoặc nhập token; POST /gate-pass/scan; đọc container, người nhận, hạn và blockers.','Thiếu camera cho phép fallback; phiếu USED/EXPIRED/CANCELLED bị từ chối.','Chỉ mở màn xác nhận khi scan trả visit/pass phù hợp.','POST /gate-pass/scan'),
('UC05','Xác nhận Gate-out','GATE_STAFF có gate_pass.use','GateOutConfirmScreen','Xe thực tế tại cổng; phiếu chưa dùng, readiness đủ tại thời điểm xác nhận.','Xem lại thông tin; backend recheck token và blocker; POST /gate-out; đọc biên nhận.','409/timeout không tự replay; phiếu bị thu hồi/hết hạn hiển thị nguyên nhân.','Visit EXITED, pass USED, vị trí bãi giải phóng theo transaction backend.','POST /gate-pass/scan; POST /gate-out'),
('UC06','Kiểm tra và xếp vị trí bãi','YARD_STAFF có yard.update','YardHomeScreen → YardAssignmentScreen','Container IN_YARD chưa có vị trí hoặc cần điều chỉnh hợp lệ.','Đọc danh sách slot/Legend; chọn ô; POST check; xem blockers/warnings; xác nhận assign MANUAL; tải lại vị trí.','Ô đã chiếm/bảo trì, block không hoạt động hoặc container không phù hợp: ngăn assign và giữ lựa chọn.','Vị trí canonical và location log được backend ghi; UI hiển thị slot mới.','GET /yard/slots; POST /containers/:visitId/yard/check; POST /containers/:visitId/yard/assign'),
('UC07','Yêu cầu di chuyển nội bộ','YARD_STAFF có yard.move','YardOperationsScreen / YardOperationDetailScreen','Visit đang có vị trí và đích được phép.','Chọn visit và slot đích; POST movement; mở detail của operation mới.','Đích chiếm/khóa hoặc nguồn sai: trình bày lý do, không giả đã chuyển.','Operation PENDING gắn visit và toSlotId.','POST /containers/:visitId/yard/movements'),
('UC08','Bắt đầu và hoàn tất di chuyển','YARD_STAFF có yard.move','YardOperationDetailScreen','Movement PENDING; container và slot liên kết còn đúng.','POST start; thao tác ngoài bãi; POST complete; tải lại detail/location.','Request đến muộn hoặc status đổi: thông báo xung đột và refetch.','Movement COMPLETED và location log có vị trí mới.','POST /yard/movements/:id/start; POST /yard/movements/:id/complete'),
('UC09','Yêu cầu giám định','YARD_STAFF có yard.inspect','SurveyHomeScreen','Visit thuộc ICD và có quyền giám định.','Chọn loại kiểm tra; nhập ghi chú thực tế; POST inspections; mở lịch sử.','Trường bắt buộc thiếu thì báo tại chỗ. Với kiểm tra seal, không yêu cầu mức hư hỏng.','Inspection PENDING được tạo với loại và notes đúng.','POST /containers/:visitId/inspections'),
('UC10','Kết thúc giám định','YARD_STAFF có yard.inspect','SurveyHomeScreen / YardOperationDetailScreen','Inspection PENDING hoặc IN_PROGRESS.','Start nếu cần; ghi kết quả PASS/FAIL/HOLD và notes; POST complete; tải lại readiness.','FAIL/HOLD thiếu lý do: từ chối; request đã hoàn thành: hiển thị trạng thái hiện tại.','Inspection COMPLETED; HOLD có thể chặn Gate Pass.','POST /inspections/:id/start; POST /inspections/:id/complete'),
('UC11','Đặt lịch tác nghiệp bãi','YARD_STAFF có yard.booking','YardOperationsScreen / BookingDateField','Visit còn trong bãi; booking type được hỗ trợ.','Chọn STRIPPING/STUFFING/INSPECTION; chọn ngày giờ địa phương; đổi sang ISO; POST booking.','Ngày sai/đã qua hoặc thiếu notes theo DTO: nêu lỗi và giữ draft.','Booking PENDING có scheduledAt đúng instant.','POST /containers/:visitId/yard-bookings'),
('UC12','Hoàn tất hoặc hủy booking','YARD_STAFF có yard.booking','YardOperationDetailScreen','Booking PENDING/IN_PROGRESS.','Start; nhập kết quả thực tế; complete hoặc hủy cùng lý do; tải lại detail.','Số kiện/cân âm, thiếu lý do hủy hoặc status đã đổi: không báo thành công.','Booking COMPLETED/CANCELLED được lưu cùng kết quả/lý do.','POST /yard/bookings/:id/start|complete|cancel'),
('UC13','Tra cứu container','Người có container.read','ContainerSearchScreen / ContainerDetailScreen','Đã xác thực, có quyền đọc visit.','Tìm số container; phân trang; mở detail; đọc event, slot, holds, readiness, pass và handover theo quyền.','Không có kết quả khác lỗi tải; 403 không được hiển thị dữ liệu cache không phù hợp.','Thông tin màn khớp API từng resource.','GET /containers; GET /containers/:visitId và các route detail'),
('UC14','Theo dõi việc ca','Nhân viên có quyền task tương ứng','WorkQueueScreen','Backend trả queue cho phạm vi user.','Đọc tổng/quá hạn; lọc loại và mức ưu tiên; mở task theo visit/operation ID và permission.','Task đổi trạng thái/mất quyền: màn hướng quay lại queue; không mở nhầm record khác.','Destination mở đúng entity; nhiệm vụ không được cập nhật bằng suy đoán UI.','GET /containers/work-queue'),
('UC15','Đọc và đánh dấu thông báo','Nhân viên có phiên hợp lệ','NotificationsScreen','Có notification thuộc user và ICD.','Đọc history/unread count; mở bản ghi; mark-one hoặc mark-all; tải lại.','Request lỗi giữ count cũ đã xác nhận và nút retry; list empty sau read thành công.','Thông báo được đọc theo response backend; không tạo count âm.','GET /notifications/history; PATCH /:id/read; POST /read-all'),
('UC16','Xem tài khoản và thay theme','Nhân viên có phiên hợp lệ','MoreScreen / ThemeProvider','Session đọc được user/role/permission.','Mở Tài khoản; xem quyền, API; đổi sáng/tối; lưu preference; logout.','Storage lỗi cho phép theme hệ thống và thông báo phù hợp; zero-tab vẫn mở được Tài khoản.','Theme và phiên hiện đúng sau mở lại; logout xóa dữ liệu phiên.','GET /auth/me; POST /auth/logout'),
('UC17','Theo dõi bãi và bàn giao','Người có yard.read hoặc handover.read','MonitorScreen / HandoverSummary','Dữ liệu bãi/visit thuộc phạm vi quyền.','Đọc occupancy và hồ sơ; từ chi tiết container mở handover summary ở chế độ chỉ xem.','Handover chưa có hoặc API bị cấm: thông báo trạng thái đúng, không hiện nút xác nhận đối tác.','Thông tin đọc khớp backend; Mobile không gửi Partner API key.','GET /yard/slots; GET /containers/:visitId/handover-summary'),
]
for idx,(code,title,actor,screen,pre,main,alt,post,api) in enumerate(USECASES,1):
    heading(f'3.3.{idx}',title,3)
    note('Mã và tác nhân.',f'{code}; {actor}. Màn: {screen}.')
    note('Tiền điều kiện.',pre)
    note('Luồng chính.',main)
    note('Ngoại lệ.',alt)
    note('Hậu điều kiện.',post)
    note('Đối chiếu API.',api+'.')

heading('3.4','Thiết kế luồng và trạng thái',2)
heading('3.4.1','Đăng nhập và phục hồi quyền',3)
figure('3.3','Trình tự đăng nhập và phân phối quyền cho tab Mobile',FIG['seq_login'])
para('Trong trình tự này, user và permission ở response auth/me là nguồn của bộ tab. Nếu ứng dụng restore sau khi tạm mất mạng, kết quả 500 hoặc không kết nối không đồng nghĩa phiên đã hết hạn. Chỉ trạng thái xác thực không hợp lệ mới dẫn tới xóa thông tin phiên. Điều kiện này được ghi trong audit vận hành như một ca regression.')

heading('3.4.2','Tiếp nhận Gate-in',3)
figure('3.4','Trình tự Gate-in giữa nhân viên, Mobile và backend',FIG['seq_gatein'])
para('Số container là khóa tra cứu ban đầu; command cuối cần visitId và truckVisitId đúng quan hệ. Màn review hiển thị seal, cân và tình trạng thực tế trước khi gửi. Sau commit backend trả receipt; app hiển thị kết quả đã xác nhận. Trong trường hợp một Truck Visit chứa nhiều container, chuyến có thể IN_PROGRESS sau container thứ nhất và COMPLETED khi container cuối đã tiếp nhận.')

heading('3.4.3','Gán vị trí bãi',3)
figure('3.5','Luồng kiểm tra slot và xác nhận gán tay',FIG['slot'])
para('Bản đồ bãi đọc trạng thái ô từ API, còn màu sắc lấy cùng mô hình Legend. Khi chọn một ô, Mobile gửi check để nhận eligible/blockers/warnings. Nút xác nhận gọi assign với source MANUAL; kết quả cuối đọc lại vị trí canonical. Nhánh gợi ý vị trí có trong một số ảnh audit lịch sử, nhưng phương án tối ưu chính thức vẫn được ghi “Đang trong quá trình phát triển”. Báo cáo không đánh giá thuật toán gợi ý như đã tối ưu.')

heading('3.4.4','Movement, inspection và booking',3)
figure('3.6','Vòng đời chung của tác nghiệp bãi',FIG['operations'])
para('Ba nhóm tác nghiệp dùng các route riêng và dữ liệu bổ sung khác nhau. Movement cần vị trí đích; inspection có loại và kết quả PASS/FAIL/HOLD; booking cần loại dịch vụ và thời gian đã chuẩn hóa. Hủy phải lưu lý do. Khi kết thúc, giao diện refetch detail để xác nhận state, bởi màn hình không thể suy ra transaction đã commit chỉ từ việc người dùng nhấn nút.')

heading('3.4.5','Quét Gate Pass và xác nhận xe ra',3)
figure('3.7','Trình tự Gate-out có bước kiểm lại điều kiện tại backend',FIG['seq_gateout'])
para('Mã QR chứa token do backend phát hành, không phải mã sinh ra từ số container. Khi quét, API kiểm tình trạng pass; lúc xác nhận, service kiểm lại readiness, hold, payment và trạng thái visit. PASS/USED/EXPIRED/CANCELLED cần hiển thị theo giá trị thật. Sau thành công, pass chuyển USED và visit EXITED; không gọi lại giao dịch tự động nếu response bị mất ở cuối đường truyền.')

heading('3.4.6','Thông báo sau giao dịch',3)
figure('3.8','Trình tự tạo và đọc thông báo theo user',FIG['seq_notification'])
para('Thông báo Gate-in, Gate-out và inspection HOLD được lưu cho người nhận nội bộ phù hợp. Màn danh sách dùng history và unread count do server cung cấp, có filter và retry. Lỗi push hoặc email sau commit không đảo ngược Gate-in/Gate-out. Việc click thông báo phải mở đúng visit hoặc operation; notification cũ không thể giữ ID của record khác.')

heading('3.5','Mô hình dữ liệu phục vụ Mobile',2)
figure('3.9','Quan hệ dữ liệu từ Container Visit đến các phân hệ hiện trường',FIG['erd'])
para('ContainerVisit là điểm neo cho reception, location log, tác nghiệp bãi, Hold, Gate Pass và handover. User, Role và Permission xác định actor có thể xem hoặc ghi từng bản ghi. YardSlot lưu khả năng chứa và trạng thái ô, còn ContainerLocationLog giữ lịch sử; lấy ô hiện tại từ dữ liệu canonical, không suy diễn bằng lịch sử bất kỳ. Sơ đồ là khung quan hệ rút gọn; schema thật có 53 model và khóa ngoại chi tiết ở Phụ lục B.')
table('3.3','Tập thực thể trọng tâm của ứng dụng hiện trường',['Thực thể','Khóa/quan hệ quan trọng','Nội dung Mobile cần đọc/ghi'],[
    ('ContainerVisit','id; containerId; icdId','State và mọi nghiệp vụ theo visit.'),
    ('TruckVisit / ContainerReception','truckVisitId; visitId','Xe đến và biên nhận Gate-in.'),
    ('YardSlot / ContainerLocationLog','yardSlotId; containerVisitId','Ô bãi hiện tại và lịch sử di chuyển.'),
    ('YardMovement','containerVisitId; fromSlotId; toSlotId','Trạng thái start/complete/cancel.'),
    ('ContainerInspection','containerVisitId; type; result','PASS/FAIL/HOLD và ghi chú.'),
    ('InYardBooking','containerVisitId; bookingType; scheduledAt','Lịch tác nghiệp và kết quả thực hiện.'),
    ('OperationalHold / GatePass','containerVisitId; status','Blocker readiness, QR và Gate-out.'),
    ('Notification','recipientUserId; related entity','Lịch sử, unread và điều hướng.'),
],[4.3,5.6,5.5],10)

heading('3.6','Thiết kế phân quyền và điều hướng',2)
para('RootNavigator tách luồng xác thực khỏi MainTabNavigator. Tab Cổng, Bãi, Giám định, Tra cứu, Việc ca và Monitor được đăng ký theo permission. OPERATOR ưu tiên Việc ca/Tra cứu/Monitor trên thanh tab; những màn tác nghiệp có quyền vẫn có thể mở từ task. Khi user không có tab, ứng dụng hiện lối sang Tài khoản/đăng xuất. Backend tiếp tục kiểm permission trên endpoint, kể cả khi UI vừa mở từ deep link.')
table('3.4','Ánh xạ tab sang permission sử dụng trong source',['Tab','Screen permission','Permission backend tiêu biểu'],[
    ('Cổng','gate.in hoặc gate.out','gate_in.create; gate_pass.use'),
    ('Bãi','yard.read, yard.assign hoặc yard.operations','yard.read; yard.update; yard.move; yard.inspect; yard.booking'),
    ('Giám định','yard.inspect','yard.inspect'),
    ('Tra cứu','container.read','container.read'),
    ('Việc ca','work_queue.read','Tập quyền tác vụ liên quan; server còn lọc queue'),
    ('Monitor','yard.read','yard.read'),
],[2.9,5.3,7.2])

heading('3.7','Khung biểu đồ cần giữ khi cập nhật luận văn',2)
para('Các hình 3.1–3.9 đã có nội dung để đọc trực tiếp. Khi thay ảnh hoặc vẽ lại trong công cụ UML, cần giữ các thành phần và điều kiện dưới đây; nếu thiếu nhánh từ chối, sơ đồ sẽ không giải thích được thao tác thực tế ở cổng/bãi.')
table('3.5','Đặc tả chi tiết để vẽ và đối chiếu từng biểu đồ',['Hình','Nút, luồng và dữ liệu bắt buộc','Cách kiểm tra'],[
    ('3.1 Use case','Actors Gate/Yard/Operator; Gate-in, Gate-out, assign, operation, lookup/queue; liên hệ theo permission.','Một actor thiếu quyền không có đường tới command tương ứng.'),
    ('3.2 State','PENDING, AUTHORIZED, IN_YARD, GATE_PASS_ISSUED, EXITED, CANCELLED; event và điều kiện trên cạnh.','Mỗi state có enum trong schema; không gán xếp bãi thành state mới.'),
    ('3.3 Login','Actor, App, Auth API, DB; JWT/user permission và nhánh lỗi 401/mạng.','Tab tính từ permission mới sau restore.'),
    ('3.4 Gate-in','Tra context, chuyến, reception, event và biên nhận sau commit.','Một request thất bại không có reception mới.'),
    ('3.5 Slot','Chọn visit/slot, check eligible, blocker/warning, assign MANUAL, refetch.','Slot không đủ điều kiện không xuất hiện trong kết quả assign.'),
    ('3.6 Operation','PENDING → IN_PROGRESS → COMPLETED; nhánh CANCELLED có lý do.','Movement/inspection/booking giữ payload riêng.'),
    ('3.7 Gate-out','Scan token, readiness, confirm, recheck transaction, EXITED/USED.','Phiếu hết hạn/đã dùng không commit Gate-out.'),
    ('3.8 Notification','Core event, recipient, record, history/unread, mark-read/deep link.','Gửi push lỗi không rollback core.'),
    ('3.9 ERD','Visit làm gốc; location, operation, hold, pass, notification gắn đúng khóa.','Tên field/khóa kiểm lại schema ở Phụ lục B.'),
],[2.2,9.5,3.7],9.6)


heading('4','TRIỂN KHAI ỨNG DỤNG VÀ KIỂM THỬ')
heading('4.1','Cấu trúc mã nguồn và khởi động',2)
para('Ứng dụng nằm tại apps/mobile. Thư mục src/features tổ chức theo auth, gate-in, gate-out, yard, containers, work-queue, notifications, handover và more. navigation giữ RootNavigator, MainTabNavigator và stack riêng của Cổng/Bãi. Các phần dùng chung như ScreenLayout, CodeScanner, ActionDialog, trạng thái loading/error và ThemeProvider nằm ở thư mục components/theme. Client dùng chung trong services/api; API caller của từng feature giữ route và kiểu dữ liệu gần màn hình.')
table('4.1','Thư mục Mobile và trách nhiệm',['Thư mục','Vai trò','Điểm đối chiếu'],[
    ('src/navigation','Tab, stack, route ID và permission.','RootNavigator.tsx; MainTabNavigator.tsx'),
    ('src/features/gate-in, gate-out','Quét/nhập và xác nhận cổng.','GateInScan/Form/Success; GatePassScan/GateOutConfirm'),
    ('src/features/yard','Sơ đồ bãi, assignment, movement/inspection/booking, monitor.','YardHome/Assignment/Operations/Detail/Survey/Monitor'),
    ('src/features/containers','Tìm kiếm và hồ sơ container.','ContainerSearch/Detail'),
    ('src/features/work-queue, notifications','Việc ca, thông báo và destination.','WorkQueueScreen; NotificationsScreen'),
    ('src/services, src/storage','API client, connection gate, cache đọc.','api-client.ts; read-cache.ts'),
    ('src/theme, src/components','Theme, layout, scanner, dialog, trạng thái UI.','ThemeProvider; CodeScanner; ScreenLayout'),
],[4.0,6.5,4.9],10)
figure('4.1','Sơ đồ triển khai môi trường thử nghiệm Mobile và API',FIG['deployment'])
para('Môi trường Expo Web dùng cổng 8081 và API thử nghiệm 3001 theo báo cáo smoke ngày 04/10. Khi mở Expo Go trên điện thoại thật, URL API phải trỏ đến IPv4 của máy trong LAN thay vì 127.0.0.1. URL đó là cấu hình môi trường; không chèn credential vào tài liệu hoặc ảnh báo cáo. Expo Web là bản React Native chạy trong trình duyệt, không thay bằng chứng camera và hệ thống quyền native.')

heading('4.2','Triển khai giao diện theo từng màn hình',2)
SCREEN_ROWS=[
    ('LoginScreen','Auth','Đăng nhập, lỗi trường/mạng, phục hồi phiên','auth.login/me/refresh'),
    ('GateInScanScreen','Cổng nhập','Camera/nhập tay, tra container','containers search'),
    ('GateInFormScreen','Cổng nhập','Context, chuyến xe, seal/cân/tình trạng, xác nhận','gate-in-context; truck arrive; gate-in'),
    ('GateInSuccessScreen','Cổng nhập','Biên nhận và tiếp tục hồ sơ mới','response gate-in'),
    ('GatePassScanScreen','Cổng xuất','Quét QR backend token, lý do từ chối','gate-pass/scan'),
    ('GateOutConfirmScreen','Cổng xuất','Recheck, xác nhận và biên nhận EXITED','gate-pass/scan; gate-out'),
    ('YardHomeScreen','Bãi','Tổng quan slot và lối mở tác nghiệp','yard/slots; containers'),
    ('YardAssignmentScreen','Bãi','Block/Bay/Row/Tier, Legend, check và assign MANUAL','yard/check; yard/assign'),
    ('YardOperationsScreen','Bãi','List/filter movement/inspection/booking','yard lists'),
    ('YardOperationDetailScreen','Bãi','Start/complete/cancel theo loại operation','yard/.../:id'),
    ('SurveyHomeScreen','Giám định','Loại, notes, request, lịch sử và kết quả','inspections'),
    ('ContainerSearchScreen','Tra cứu','Search, pagination, empty/error','containers'),
    ('ContainerDetailScreen','Tra cứu','Visit, events, slot, holds, readiness/pass, handover','container detail APIs'),
    ('WorkQueueScreen','Việc ca','Task, SLA, filter, destination theo quyền','containers/work-queue'),
    ('MonitorScreen','Điều phối','KPI bãi và container theo quyền','yard/container reads'),
    ('NotificationsScreen','Thông báo','Unread, lịch sử, read/read-all','notifications'),
    ('MoreScreen','Tài khoản','Thông tin user/quyền, theme, logout','auth/me; local storage'),
]
table('4.2','Danh sách đủ 17 file màn hình Mobile trong inventory',['Màn hình','Nhóm','Nhiệm vụ trên giao diện','Nguồn dữ liệu'],SCREEN_ROWS,[3.5,2.0,5.8,4.1],9.5)

OLD=ROOT/'test-artifacts/2026-10-01-mobile-operations-audit'
NEW=ROOT/'audit/runs/2026-10-03-improvement-02/screenshots/mobile-web-2026-10-04'
def audit_image(number,title,filename,source='old',note_text=''):
    p=(OLD if source=='old' else NEW)/filename
    assert p.exists(),p
    figure(number,title,p,note_text or 'Ảnh audit Expo Web theo bộ ngày 02/10/2026; bản trình duyệt dùng để minh họa giao diện và thao tác đã ghi.',width=10.0 if source=='old' else 9.5)

heading('4.2.1','Đăng nhập và thông tin phiên',3)
para('LoginScreen nhận email/mật khẩu, báo lỗi xác thực và tình huống mất kết nối. AuthProvider khôi phục user cùng permission; từ đó MainTabNavigator tạo danh sách tab. Nếu restore bị lỗi mạng, app cho phép thử lại. MoreScreen hiển thị tài khoản/quyền và đổi sáng/tối; lựa chọn theme được lưu cho lần mở sau. Tài khoản không có tab vẫn có đường vào Tài khoản và Đăng xuất.')
audit_image('4.2','Giao diện đăng nhập Mobile trong bộ audit','58-login-360.jpg')
audit_image('4.3','Tài khoản và quyền của người dùng','04-account-permissions-dialog.jpg')

heading('4.2.2','Cổng nhập và biên nhận Gate-in',3)
para('GateInScanScreen cho quét mã hoặc nhập số. Bản Web fallback nhập tay khi không thể dùng camera; CodeScanner native còn phải kiểm trên Android. Sau khi tìm container, GateInFormScreen đọc gate-in-context, cho chọn/đánh dấu chuyến xe đến và nhập thông tin thực tế. Trước khi gửi command có màn review và dialog xác nhận. GateInSuccessScreen dùng response đã commit để hiển thị biên nhận và reset form khi tiếp tục container khác.')
audit_image('4.4','Màn tiếp nhận vào cổng với các trường xác nhận thực tế','59-gate-in-light-360.jpg')
audit_image('4.5','Màn review Gate-in trước khi gửi giao dịch','11-gate-in-review.jpg')
audit_image('4.6','Biên nhận Gate-in sau phản hồi backend','12-gate-in-first-success.jpg')
note('Trạng thái cần phân biệt.', 'Một xe có thể chở nhiều container trong một Truck Visit. Sau lần tiếp nhận đầu, xe có thể còn IN_PROGRESS; biên nhận từng container không đồng nghĩa toàn bộ chuyến đã hoàn thành.')

heading('4.2.3','Quét phiếu và xác nhận Gate-out',3)
para('GatePassScanScreen nhận token từ QR do backend phát hành, rồi gọi gate-pass/scan. Màn GateOutConfirmScreen hiển thị số container, mã phiếu, xe, người nhận, hạn, readiness và blocker. Thao tác cuối gửi POST /gate-out; khi backend xác nhận, UI hiện thời điểm ra cổng và trạng thái USED/EXITED. Trên native có nhánh nhấn giữ theo thiết kế; báo cáo không nhận nhánh đó đã được kiểm từ ảnh Expo Web.')
audit_image('4.7','Khung quét QR phiếu ra cổng ở bản Expo Web','60-gate-out-scan-360.jpg')
audit_image('4.8','Màn xem lại điều kiện và xác nhận Gate-out','63-gate-out-review-dialog-360.jpg')
audit_image('4.9','Kết quả Gate-out và phiếu đã sử dụng','64-gate-out-success-360.jpg')

heading('4.2.4','Tổng quan bãi và xếp vị trí bằng tay',3)
para('YardHomeScreen tải block/slot và hồ sơ đang ở bãi. YardAssignmentScreen cho chọn container, xem vị trí theo Block–Bay–Row–Tier và Legend. Nút kiểm tra gửi yard/check; kết quả gồm eligible, blockers, warnings. Lệnh assign MANUAL chỉ gửi sau xác nhận và được đọc lại bằng API location. Thuật toán tối ưu xếp vị trí chưa có tiêu chí nghiệm thu, nên phần gợi ý trong giao diện được trình bày là đang phát triển.')
audit_image('4.10','Sơ đồ vị trí và trạng thái bãi đã tải dữ liệu','yard-populated-375.png',source='new',note_text='Ảnh smoke Expo Web ngày 04/10/2026: dữ liệu 56/56 ô của database test tại thời điểm chụp; đây không phải công suất bãi thực tế.')
audit_image('4.11','Kiểm tra ô đã chiếm và lý do từ chối','23-yard-occupied-refusal.jpg')
audit_image('4.12','Xếp vị trí MANUAL đã được backend ghi nhận','25-yard-manual-assigned.jpg')

heading('4.2.5','Danh sách và chi tiết tác nghiệp bãi',3)
para('YardOperationsScreen gộp list của movement, inspection và booking nhưng giữ filter theo loại/trạng thái. YardOperationDetailScreen mở đúng operation ID; bộ nút start, complete, cancel phụ thuộc status, type và quyền. Màn hủy yêu cầu lý do, còn form hoàn tất booking kiểm số kiện và trọng lượng. Phản hồi lỗi không được thay bằng trạng thái thành công khi request hậu xử lý thất bại.')
audit_image('4.13','Danh sách tác nghiệp bãi và bộ lọc','44-yard-operations-list.jpg')
audit_image('4.14','Chi tiết movement đang thực hiện','30-movement-in-progress.jpg')
audit_image('4.15','Booking hoàn tất với kết quả thực tế','54-booking-completed-results.jpg')

heading('4.2.6','Giám định hiện trường',3)
para('SurveyHomeScreen cho chọn loại giám định và nhập nội dung. Kiểm tra seal không hỏi mức độ hư hỏng; phần severity dành cho trường hợp kiểm hư hỏng. Lịch sử có loading/error/retry riêng, chỉ hiện “chưa có biên bản” sau response thành công. Inspection HOLD tác động đến readiness; chưa có endpoint release cho inspection HOLD được nghiệm thu trong bộ audit hiện tại.')
audit_image('4.16','Giám định với lựa chọn loại kiểm tra và ghi chú','survey-seal-375.png',source='new',note_text='Ảnh smoke Expo Web 04/10/2026 chứng minh loại Kiểm tra seal không hiển thị mức độ hư hỏng; chưa gửi biên bản trong lượt smoke này.')
audit_image('4.17','Biên bản giám định đã hoàn tất với kết quả','43-inspection-completed-notes.jpg')

heading('4.2.7','Tra cứu container và hồ sơ 360 độ',3)
para('ContainerSearchScreen tìm theo số container và phân trang. ContainerDetailScreen tải những nhóm dữ liệu có thể độc lập: thông tin visit, vị trí, events, holds, readiness, phiếu ra cổng và tóm tắt bàn giao. Mỗi phần có trạng thái loading/error/forbidden riêng để việc thiếu quyền đọc invoice hoặc hold không bị trình bày thành số không hoặc “không có”. Thông tin tài chính trong readiness được giới hạn theo quyền.')
audit_image('4.18','Tìm container theo số và chọn hồ sơ','18-lookup-results.jpg')
audit_image('4.19','Chi tiết container và vị trí trong bãi','container-detail-375.png',source='new',note_text='Ảnh smoke Expo Web 04/10/2026; vị trí A-01-01-1 được đối chiếu với sơ đồ bãi trong lượt kiểm thử này.')

heading('4.2.8','Việc ca và màn điều phối',3)
para('WorkQueueScreen đọc task theo API, hiển thị số nhiệm vụ và quá hạn, filter theo loại/urgency. Task cần mang visitId hoặc operationId đúng loại trước khi điều hướng. MonitorScreen ưu tiên cho OPERATOR có quyền xem bãi; thống kê lấy từ các trang dữ liệu đã tải thay vì chỉ một trang. Khi task hết quyền hoặc không còn tồn tại, màn quay về queue và thông báo rõ.')
audit_image('4.20','Việc ca có dữ liệu đã tải xong','work-queue-populated-375.png',source='new',note_text='Ảnh database test ngày 04/10/2026 có 12 task và 10 quá hạn tại thời điểm chụp, không phản ánh dữ liệu vận hành hiện tại.')
audit_image('4.21','Màn điều phối cho tài khoản Operator','80-operator-monitor-light.jpg')

heading('4.2.9','Thông báo, Tài khoản và các popup phụ',3)
para('NotificationsScreen lấy history, unread count; đọc một hoặc tất cả chỉ cập nhật sau response API. Popup thông báo mở đúng entity được gửi. MoreScreen lưu theme sáng/tối và hiển thị identity/permission. Dialog xác nhận, picker, màn lỗi mạng và khung quét là các màn phụ có ảnh riêng trong bộ audit; Back cần đóng overlay trước khi rời stack. Những case camera denied và TalkBack cần đo trên native.')
audit_image('4.22','Popup thông báo xe đến và liên kết hồ sơ','67-notification-arrival-popup.jpg')
audit_image('4.23','Các màn chính ở chế độ tối sau audit','98-final-gate-dark.jpg')

heading('4.3','Đối chiếu API, JSON và trạng thái dữ liệu',2)
para('Mobile gọi API nội bộ bằng Bearer JWT. Response dữ liệu thường nằm trong data; danh sách có meta. Lỗi được chuẩn hóa qua error.code, message, details và requestId. Với lỗi mạng hoặc outcome chưa xác định của write, ứng dụng giữ bối cảnh để người dùng kiểm lại trạng thái trước khi gửi một lệnh khác. Danh mục endpoint và DTO liên quan ở Phụ lục A được sinh từ inventory của source; đó là route có mặt trong mã, không phải kết quả chạy thử từng route.')
table('4.3','Payload mẫu cần đối chiếu trong thao tác hiện trường',['Thao tác','JSON/request chủ yếu','Điểm kiểm'],[
    ('Gate-in','{truckVisitId, actualSeal, actualWeight, conditionCode, conditionNotes}','visitId nằm ở path; xe thuộc chuyến hợp lệ.'),
    ('Yard check','{yardSlotId}','eligible, blockers, warnings được đọc trước assign.'),
    ('Yard assign','{yardSlotId, source: MANUAL}','Kết quả canonical từ backend.'),
    ('Inspection','{inspectionType, notes}; complete {result, notes}','Result PASS/FAIL/HOLD theo backend.'),
    ('Booking','{bookingType, scheduledAt, conditionNotes}','scheduledAt là ISO đúng instant UTC.'),
    ('Gate Pass scan','{qrToken}','Token do backend cấp, không từ số container.'),
    ('Gate-out','{visitId, qrToken}','Recheck pass/readiness và transaction EXITED/USED.'),
],[3.1,7.4,4.9],9.7)

heading('4.4','Kiểm thử đã có bằng chứng',2)
para('Bộ audit vận hành ngày 02/10/2026 chạy Mobile bằng Expo Web với API thật và hồ sơ QA riêng, ghi 100 ảnh. Backend regression 78/78 bài và mobile regression 50/50 bài ở lượt được báo cáo; hai sửa chữ cuối chỉ được typecheck/lint/export sau đó. Kết quả verify durable xác nhận trạng thái của bốn visit sau restart. Smoke ngày 04/10/2026 dùng API 3001 và DB test riêng: mở các tab, đọc bãi, việc ca, tra cứu, scanner fallback và theme. Những con số này được ghi theo phạm vi của mỗi báo cáo nguồn, không cộng thành tỷ lệ nghiệm thu Android.')
table('4.4','Kết quả có thể dẫn chứng và giới hạn',['Hạng mục','Kết quả ghi trong audit','Phạm vi bằng chứng'],[
    ('Expo Web vận hành 02/10','100 ảnh; nhiều luồng Gate/Yard/Inspection/Booking/Notification được thao tác.','UI trình duyệt và API/hồ sơ QA.'),
    ('Backend regression','16/16 bộ, 78/78 bài trong báo cáo audit ngày 02/10.','Phần regression tại thời điểm audit.'),
    ('Mobile regression','50/50 bài trong lượt suite đã ghi.','Không thay kiểm thử camera/native.'),
    ('Verify durable','57 invariant qua 45 API checks sau restart.','Bốn visit/ba truck của fixture QA.'),
    ('Expo Web smoke 04/10','Đăng nhập, mở năm tab, bãi, queue, lookup, seal, scanner fallback, theme.','Không gửi tất cả command trong lượt smoke.'),
    ('Native Android/iOS','Chưa có kết luận nghiệm thu trong nguồn dùng cho báo cáo.','Cần chạy ca ở Bảng 4.7.'),
],[3.1,6.1,6.2],9.8)
heading('4.4.1','Các mốc trạng thái sau giao dịch',3)
table('4.5','Ví dụ hồ sơ QA đã đọc lại sau thao tác',['Hồ sơ / sự kiện','Trạng thái theo audit','Điều được chứng minh'],[
    ('Xe 51C-QA0755','Sau Gate-in đầu IN_PROGRESS; sau container cuối COMPLETED.','Visit/reception gắn cùng chuyến.'),
    ('QAOU5007550','EXITED; Gate Pass USED; slot null; movement, inspection, booking đã hoàn tất trước exit.','Gate-out và dữ liệu liên quan bền vững sau restart.'),
    ('QAOU5007565','IN_YARD; slot A-01-03-2; movement và booking CANCELLED.','Hủy giữ vị trí nguồn/lý do.'),
    ('QAOU5007570','IN_YARD; inspection HOLD; operational hold ACTIVE; readiness false.','Hold chặn điều kiện ra cổng.'),
],[3.2,5.8,6.4],9.7)
note('Cách đọc ảnh.', 'Ảnh 4.4–4.23 được chọn từ các bộ audit; nhãn thời gian/trạng thái thuộc lúc chụp. Nếu người dùng chèn ảnh mới, cần cập nhật caption, case ID và kết quả tương ứng. Ảnh trước sửa trong gallery lịch sử chỉ dùng ở mục so sánh lỗi, không trình bày như UI cuối.')

heading('4.4.2','Ma trận kiểm thử giao diện và nghiệp vụ',3)
TESTED=[
    ('A01','Login','Nhập sai tài khoản','Có lỗi xác thực, không mở tab','Ảnh 06; audit 02/10'),
    ('A02','Session','Mất kết nối khi restore','Giữ phiên để retry, không xóa do 503','Ảnh 13–14; regression'),
    ('A03','Role','GATE_STAFF mở Tài khoản/quyền','Hiển thị đúng role và permission','Ảnh 03–04'),
    ('G01','Gate-in','Nhập container không hợp lệ','Có lỗi, không mở form giao dịch','Ảnh 07'),
    ('G02','Gate-in','Đánh dấu chuyến xe đến','Backend đổi xe sang ARRIVED','Ảnh 08; fixture gate'),
    ('G03','Gate-in','Thiếu seal thực tế','Trường bắt buộc được chỉ rõ','Ảnh 09'),
    ('G04','Gate-in','Xác nhận lượt đầu','Có reception; xe có thể IN_PROGRESS','Ảnh 11–12; verify Gate-in'),
    ('G05','Gate-in','Xác nhận container cuối của xe','Truck Visit COMPLETED, visit đúng','Ảnh 15–16; verify Gate-in'),
    ('G06','Gate-out','Dùng QR sai','Bị từ chối, không ghi Gate-out','Ảnh 61'),
    ('G07','Gate-out','Xác nhận phiếu hợp lệ','Visit EXITED và pass USED sau commit','Ảnh 62–64; durable verify'),
    ('G08','Gate-out','Quét lại phiếu USED','Bị từ chối và báo đã dùng','Ảnh 65'),
    ('Y01','Yard','Chọn slot đã chiếm','Check từ chối cùng lý do','Ảnh 23'),
    ('Y02','Yard','Assign MANUAL slot hợp lệ','Location canonical được ghi','Ảnh 24–25; verify follow-up'),
    ('Y03','Movement','Tạo / start / complete','State PENDING → IN_PROGRESS → COMPLETED','Ảnh 26–32; durable verify'),
    ('Y04','Movement','Hủy thiếu lý do / hủy hợp lệ','Validation và trạng thái CANCELLED','Ảnh 28, 87–88'),
    ('S01','Inspection','Thiếu field / chọn loại','Validation, loại kiểm tra đúng','Ảnh 37–39'),
    ('S02','Inspection','Complete PASS với notes','COMPLETED/PASS và notes được lưu','Ảnh 42–43; verify durable'),
    ('S03','Inspection','Complete HOLD với lý do','Hold tồn tại và readiness false','Ảnh 78–84; verify HOLD'),
    ('B01','Booking','Ngày sai hoặc số thực tế âm','Validation chặn command','Ảnh 47, 52'),
    ('B02','Booking','Create/complete/cancel','Trạng thái và kết quả/lý do lưu đúng','Ảnh 49, 53–57'),
    ('Q01','Queue','Mở booking từ task','Đúng operation ID và quyền','Ảnh 50–51; queue regression'),
    ('N01','Notification','Mark-one/read-all','Backend và unread count đồng bộ','Ảnh 66–72; verify notification'),
    ('R01','Readiness','Role Gate/Yard đọc blocker','Không lộ financial rows khi thiếu quyền','Verify permission'),
]
table('4.6','Các case đã được báo cáo trong bộ audit Expo Web/API',['ID','Nhóm','Thao tác','Kết quả cần khớp','Evidence'],TESTED,[1.0,2.2,4.0,5.1,3.1],9.2)
para('Bảng 4.6 dẫn kết quả ở báo cáo audit ngày 02/10 và hồ sơ QA đi kèm. “Ảnh” chứng minh màn tại thời điểm chụp; “verify” là file kiểm tra dữ liệu sau command. Trạng thái CASE nếu chạy lại có thể khác do hồ sơ test thay đổi. Khi cập nhật báo cáo cuối, cần gắn timestamp/run ID của lượt test mới và không giữ nguyên chữ Đạt khi nguồn đã đổi.')

NATIVE=[
    ('NT01','Android camera chưa cấp','Hỏi quyền đúng lúc; từ chối có nhập tay và mở Settings.','Ảnh/clip native, permission dialog'),
    ('NT02','Đọc QR container bằng camera','Mã nhận diện đúng, không nhận trùng liên tục.','Camera frame + visit ID'),
    ('NT03','Đọc Gate Pass QR thật','Token decode đúng, pass và visit khớp.','Scan response + camera video'),
    ('NT04','Quay từ Settings sau cấp quyền','Camera trở lại được, không kẹt màn cũ.','Trace app lifecycle'),
    ('NT05','Android Back khi dialog mở','Đóng dialog trước, không mất form chưa gửi.','Video + state form'),
    ('NT06','Font scale 1.5 và 2.0','Không tràn tiêu đề/tab, action chính còn thấy.','Ảnh 320/360/412dp'),
    ('NT07','TalkBack đọc màn Cổng','Tên, role, trạng thái, lỗi field được đọc đúng thứ tự.','Bản ghi lời đọc/người kiểm'),
    ('NT08','TalkBack đọc scanner','Nút bật đèn, nhập tay, thoát có label và focus.','Native hierarchy + lời đọc'),
    ('NT09','Touch target và khoảng cách','Android ≥48dp, khoảng cách độc lập ≥8dp.','XML bounds + density'),
    ('NT10','Bàn phím ảo ở form dài','Field/lỗi/nút xác nhận không bị che.','Ảnh với keyboard'),
    ('NT11','DateTimePicker booking','Chọn/hủy đúng UTC+7 và ISO sau gửi.','Native picker + payload'),
    ('NT12','Lưu theme qua cold launch','Sáng/tối không flash sai khi mở app.','Video cold launch'),
    ('NT13','Offline lúc chuẩn bị Gate-in','Write bị khóa, draft còn, không tự gửi lại.','Network trace + UI'),
    ('NT14','Reconnect và đổi permission','Revalidate user/visit trước khi mở write.','Response auth/me + request trace'),
    ('NT15','Timeout ngay sau Gate-out','Không gửi lại tự động; hiển thị cần kiểm trạng thái.','Fault harness + DB check'),
    ('NT16','Deep link task tới màn cần auth','Đăng nhập rồi mở đúng ID sau kiểm quyền.','URL/route trace không chứa JWT'),
    ('NT17','Read-only/zero-tab','Không có command trái quyền; vẫn có tài khoản/logout.','Tài khoản DB test + ảnh'),
    ('NT18','iOS VoiceOver/safe area','Điều hướng và vùng chạm đạt trên iPhone.','Chưa có thiết bị trong lượt này'),
]
table('4.7','Ca nghiệm thu native còn cần thực hiện',['ID','Tình huống','Kết quả kỳ vọng','Bằng chứng phải lưu'],NATIVE,[1.1,4.2,6.5,3.6],9.3)
para('Bảng 4.7 là kế hoạch nghiệm thu. Mỗi hàng hiện chưa có bằng chứng đạt native trong các tài liệu đã đối chiếu. NT18 thuộc iOS và được tách khỏi cổng nghiệm thu Android. Khi test cần lưu phiên bản app, Android/iOS, kích thước/density, role, theme và URL API; không ghi token hoặc mật khẩu.')

heading('4.5','Đánh giá UI/UX theo audit và phần còn thiếu',2)
para('Báo cáo re-audit ngày 04/10 ghi UX Gap Mobile điểm tối thiểu đã xác minh 60,6/100, điểm tối đa theo các observation chưa kiểm là 100; Nielsen quy đổi từ finding đã xác minh là 45/100. Các finding đã sửa ở source vẫn có thể chờ acceptance native. Những con số này phản ánh checkpoint trên bộ rubric và coverage của audit, không phải điểm hoàn thiện của mọi màn Android. Sau đó smoke Expo Web không tự nâng điểm native.')
table('4.8','Một số finding Mobile còn cần acceptance trên native',['ID','Nội dung ảnh hưởng','Cách kiểm tiếp'],[
    ('M-004 / M-005','Khoảng cách vùng chạm, màn hẹp và font lớn.','Đo bounds và chụp 320dp, font 150/200%.'),
    ('M-006 / M-007','Nhãn enum tiếng Việt và loại giám định không hỏi severity sai.','Đi qua từng enum và đổi loại form.'),
    ('M-008 / M-009','Lịch sử giám định loading/error; unread count.','Gây delay/error và đánh dấu đã đọc.'),
    ('M-010 / M-011','Picker booking và lưu theme.','Native picker, cold launch, restart app.'),
    ('M-012 / M-014','Lỗi theo field và heading accessibility.','Sửa input, TalkBack traversal.'),
    ('M-N-001 / M-N-002','Race logout/refresh và cache của phiên cũ.','Fault harness với response đến muộn.'),
],[3.1,7.3,5.0],9.7)
para('Điểm Web trong audit cùng đợt là 96,5/100 ở thang UX Gap tối thiểu, nhưng báo cáo này chỉ phân tích Mobile. Không dùng điểm Web để suy ra chất lượng ứng dụng React Native. Nguồn: [6], [7].')

heading('4.6','Triển khai môi trường và tái hiện kiểm thử',2)
para('Monorepo cài dependency bằng pnpm. API cần database MySQL riêng khi test giao dịch. README có các script BAT khởi động Local, Audit và Preview; script không tự seed hay migrate dữ liệu. Với thiết bị Android/Expo Go, máy điện thoại và máy chạy backend phải cùng mạng hoặc dùng định tuyến thích hợp, rồi đặt EXPO_PUBLIC_API_BASE_URL theo IPv4 LAN của máy backend. Với Android Emulator, host API có thể cần 10.0.2.2. Expo Web dùng localhost của máy tính thử nghiệm.')
table('4.9','Các bước lặp lại để chạy và đối chiếu một ca kiểm thử',['Bước','Thao tác','Dấu hiệu đúng'],[
    ('1','Xác định database và API target, kiểm tra health/readiness.','Không trỏ nhầm DB; API trả ready.'),
    ('2','Khởi chạy Mobile, xác nhận URL API và login role cần test.','Màn hiển thị đúng user/ICD/permission.'),
    ('3','Chọn fixture container, visit, truck/pass/slot có trạng thái đầu.','Ghi ID và ảnh trước thao tác.'),
    ('4','Thao tác trên app, giữ request/response theo requestId.','Không gửi lặp; lỗi hiển thị đúng.'),
    ('5','Đọc lại bằng API hoặc SQL read-only sau command.','State, event, location, pass khớp UI.'),
    ('6','Chụp ảnh màn và dialog, ghi timestamp/thiết bị/theme.','Bằng chứng liên kết được case ID.'),
],[1.2,8.7,5.5],10)
para('Các lệnh kiểm mã phù hợp với repo: pnpm --filter @icd/mobile typecheck; pnpm --filter @icd/mobile test; pnpm --filter @icd/api test và các bài e2e có DB test guard. Kết quả ghi trong audit cũ chỉ áp dụng cho commit và cấu hình của lượt đó. Người kiểm cần chạy lại sau khi thay mã, nguồn dữ liệu hoặc điều kiện backend.')

heading('4.7','Giới hạn hiện tại và hướng hoàn thiện',2)
para('Ảnh Expo Web đã giúp xác minh bố cục, dòng dữ liệu API và nhiều giao dịch trên hồ sơ QA. Nó không đo được camera native, TalkBack, pin, vòng đời app hoặc iOS. Bản báo cáo dùng từ “đã xác minh” khi có cả ảnh và trạng thái backend; dùng “chưa kiểm” khi còn thiếu một trong các lớp cần thiết.')
table('4.10','Danh sách việc còn mở theo mức ảnh hưởng',['Ưu tiên','Nội dung','Điều kiện đóng'],[
    ('P1','Nghiệm thu camera, QR, quyền thiết bị và Gate-in/Gate-out Android.','Case NT01–NT05, NT15 có ảnh/video + API/DB.'),
    ('P1','Kiểm lỗi mạng, race phiên và quyền sau reconnect.','NT13–NT17 có fault trace.'),
    ('P1','Kiểm dữ liệu readiness/HOLD/Gate Pass trên fixture khó.','Blockers, invoice và pass/status khớp backend.'),
    ('P2','TalkBack, font scale lớn, target/spacing, keyboard và theme.','NT06–NT12 đạt với số đo và lời đọc.'),
    ('P2','Hoàn thiện pipeline ảnh seal/vỏ nếu nghiệp vụ yêu cầu.','Có API lưu file, quyền, ảnh và kiểm readback.'),
    ('Đợt riêng','iOS/VoiceOver và phân tích hiệu năng field.','Thiết bị và dữ liệu đo thực tế.'),
],[2.0,8.5,4.9],9.8)

page_title('KẾT LUẬN')
para('Ứng dụng Mobile ICD hiện có các màn phục vụ nhân viên Cổng, Bãi, Giám định, Tra cứu và Việc ca, cùng những màn phụ cho thông báo, tài khoản và điều phối. Mã nguồn dùng React Native/Expo, kết nối API NestJS thật và để backend quyết định permission, readiness và giao dịch. Các ca audit Expo Web đã đi qua chuỗi Gate-in, xếp bãi, movement, inspection, booking và Gate-out trên hồ sơ QA; dữ liệu sau restart được kiểm lại trong báo cáo nguồn.')
para('Giới hạn lớn nhất đối với kết luận chất lượng Mobile nằm ở nghiệm thu thiết bị: camera/quyền, TalkBack, màn hẹp với cỡ chữ lớn, Android Back và hành vi mạng khi app chuyển foreground. Bộ ảnh audit đã được gắn vào báo cáo đúng với môi trường chụp. Bảng 4.7 đưa ra các ca cần làm để có thể cập nhật kết luận Android; iOS là đợt sau. Báo cáo cũng giữ ranh giới của Mobile: billing, quản trị và xác nhận Partner tiếp tục theo ứng dụng hoặc API được phân quyền tương ứng.')
para('Khi hoàn thiện bản nộp, cần chạy lại ma trận native trên build cuối, cập nhật mục lục và danh mục hình/bảng trong Word, thay ảnh nào không còn đúng phiên bản, rồi đối chiếu lại các số inventory nếu mã nguồn thay đổi sau ngày 04/10/2026.')

page_title('TÀI LIỆU THAM KHẢO')
REFS=[
    '[1] ICD Management System, README.md, cập nhật 04/10/2026; monorepo icd-management.',
    '[2] ICD_Mobile_Spec_v1.7_MySQL.md, đặc tả ứng dụng Mobile trong docs/business.',
    '[3] ICD_Business_Spec_v1.7_MySQL.md, đặc tả quy trình ICD trong docs/business.',
    '[4] apps/mobile/src, mã màn hình, navigation, theme, cache và API caller; đối chiếu 08/10/2026.',
    '[5] docs/system/system-inventory.json, snapshot mã nguồn ngày 04/10/2026.',
    '[6] test-artifacts/2026-10-01-mobile-operations-audit/REPORT.md và gallery 100 ảnh, kiểm ngày 02/10/2026.',
    '[7] audit/2026-10-04-mobile-web-preview-report.md và audit/2026-10-04-reaudit-01-report.md.',
    '[8] apps/api/prisma/schema.prisma; các controller/DTO/service tương ứng trong apps/api/src/modules.',
    '[9] ICD_Database_Design_v1.7_MySQL.md, docs/business; ICD_Partner_Handover_API_Spec_v1.7_MySQL.md, docs/api.',
]
for s in REFS:para(s)


page_title('PHỤ LỤC A  DANH MỤC API VÀ DTO MOBILE')
para('Bảng dưới tổng hợp endpoint có liên quan trực tiếp đến các luồng Mobile đã mô tả. Đường dẫn là route backend với prefix /api. Permission là điều kiện ở controller; service còn kiểm ownership, state và ICD. Nếu controller yêu cầu nhiều quyền, toàn bộ quyền được áp dụng. Đây là danh mục source ngày 04/10/2026, không phải bảng kết quả gọi API thành công.')
allowed_modules={'auth','containers','gate-in','gate-out','gate-pass','truck-visits','yard','work-queue','notifications','partner-handover'}
def mobile_endpoint(e):
    if e['module'] not in allowed_modules:return False
    path=e['path']
    if e['module']=='auth':return any(x in path for x in ('/auth/login','/auth/logout','/auth/me','/auth/refresh'))
    if e['module']=='containers':return e['method']=='GET'
    if e['module']=='gate-pass':return ('/scan' in path or '/readiness' in path or '/gate-passes' in path) and e['method'] in ('GET','POST')
    if e['module']=='truck-visits':return e['method']=='GET' or '/arrive' in path
    if e['module']=='yard':return not any(x in path for x in ('/yard/blocks','/yard/slots/:slotId')) or e['method']=='GET'
    if e['module']=='partner-handover':return e['method']=='GET' and ('/handover-summary' in path or path.startswith('/api/handovers/:'))
    return True

module_labels={
    'auth':'Xác thực','containers':'Container','gate-in':'Cổng nhập','gate-out':'Cổng xuất',
    'gate-pass':'Phiếu ra cổng','truck-visits':'Chuyến xe','yard':'Bãi/giám định',
    'work-queue':'Việc ca','notifications':'Thông báo','partner-handover':'Bàn giao (đọc)',
}
by_group=defaultdict(list)
for e in INV['endpoints']:
    if not mobile_endpoint(e):continue
    b=[x['type'] for x in e['inputs'] if x['location']=='Body']
    perm=', '.join(e['permissions']) or ('JWT + ownership' if e['auth']=='JWT' else e['auth'])
    by_group[e['module']].append((e['method']+' '+e['path'],perm,', '.join(b) or '—',e['file'].split('/')[-1]+':'+str(e['line'])))
for i,module in enumerate(['auth','containers','truck-visits','gate-in','yard','gate-pass','gate-out','work-queue','notifications','partner-handover'],1):
    rows=by_group.get(module,[])
    if not rows:continue
    table(f'A.{i}',f'Endpoint {module_labels[module]} liên quan Mobile',['Method và route','Quyền','Body DTO','Nguồn file:line'],rows,[6.4,3.25,2.85,2.9],8.7)

para('Một số caller trong apps/mobile xây route động hoặc sử dụng request chung, vì vậy đối chiếu cần xem cả file API của feature và controller backend. Không lấy số call site chia 205 endpoint để tính tỷ lệ “đã đấu backend”. API billing, cấu hình, tạo phiếu và external Partner nằm ngoài thao tác hiện trường Mobile đã có.')
DTO_NAMES={
    'LoginDto','RefreshTokenDto','CreateContainerReceptionDto','ScanGatePassDto','ConfirmGateOutDto',
    'CheckYardSlotDto','AssignYardSlotDto','RequestYardMovementDto','RequestContainerInspectionDto',
    'CompleteContainerInspectionDto','CreateInYardBookingDto','CompleteInYardBookingDto',
    'CancelYardMovementDto','CancelContainerInspectionDto','CancelInYardBookingDto','PaginationQueryDto',
}
dto_rows=[]
for dto in INV['dtos']:
    if dto['name'] not in DTO_NAMES:continue
    fields=[]
    for f in dto['fields']:
        flags=[]
        if f['optional']:flags.append('?')
        if f['validators']:flags.append('/'.join(v.split('(')[0] for v in f['validators'][:3]))
        fields.append(f['name']+(' '+f['type'] if f['type'] else '')+' '.join(flags))
    dto_rows.append((dto['name'],'; '.join(fields) or 'Kế thừa/không field cục bộ',dto['file'].split('/')[-1]+':'+str(dto['line'])))
table('A.11','Trường DTO đầu vào ảnh hưởng đến Mobile',['DTO','Field và validation chính','Nguồn'],dto_rows,[4.4,8.4,2.6],8.8)
para('Ví dụ JSON trong Bảng 4.3 minh họa cấu trúc request; DTO ở Bảng A.11 là nguồn để xác định field thực tế và validation. Dữ liệu trả về cần xem response/mapper của từng service. Trường null và danh sách rỗng có ý nghĩa khác lỗi load.')

page_title('PHỤ LỤC B  MÔ HÌNH CƠ SỞ DỮ LIỆU LIÊN QUAN MOBILE')
para('Schema Prisma của hệ thống có 53 model. Bảng B.1 nêu tên bảng theo nhóm để định vị; Bảng B.2 đi sâu vào các thực thể Mobile đọc hoặc ghi trong nghiệp vụ cổng/bãi. Các trường hệ thống createdAt/updatedAt và relation array được rút khỏi phần diễn giải khi không quyết định luồng Mobile; file schema là nguồn đầy đủ.')
groups={
    'Site, xác thực, RBAC':{'IcdSite','User','AuthSession','Role','Permission','UserRole','RolePermission','IcdSetting'},
    'Danh mục và chứng từ':{'ShippingLine','Consignee','ClearingAgent','Transporter','Manifest','MasterBl','HouseBl','Container','ContainerVisit','ContainerEvent','MovementOrder','TruckVisit','TruckVisitContainer','ContainerReception'},
    'Bãi và tác nghiệp':{'YardBlock','YardSlot','ContainerLocationLog','YardMovement','ContainerInspection','InYardBooking','YardRecommendation','YardRecommendationCandidate'},
    'Phí, Hold, phiếu':{'ServiceType','Tariff','TariffRule','ServiceOrder','ServiceOrderItem','Invoice','Payment','PaymentAllocation','OperationalHold','GatePass'},
    'Tích hợp, thông báo, audit':{'AuditLog','EdiRoute','EdiOutboxMessage','EdiAcknowledgement','EdiAlert','PartnerApiClient','CustomerWarehouse','TransportHandover','TransportConfirmation','PartnerApiLog','Notification','NotificationDevice','NotificationDelivery'},
}
model_by_name={m['name']:m for m in INV['models']}
assert set().union(*groups.values())==set(model_by_name)
table('B.1','Danh mục 53 bảng theo nhóm nghiệp vụ',['Nhóm','Model Prisma / tên bảng'],[
    (g,'; '.join(f'{n} ({model_by_name[n]["table"]})' for n in sorted(names)))
    for g,names in groups.items()
],[3.7,11.7],9.4)
detail_names=['IcdSite','User','AuthSession','Container','ContainerVisit','ContainerEvent','MovementOrder','TruckVisit','TruckVisitContainer','ContainerReception','YardBlock','YardSlot','ContainerLocationLog','YardMovement','ContainerInspection','InYardBooking','OperationalHold','GatePass','Notification','TransportHandover']
detail_rows=[]
for name in detail_names:
    m=model_by_name[name]
    scalar=[f for f in m['fields'] if not f['relation'] and f['name'] not in ('createdAt','updatedAt')]
    basics=[]
    for f in scalar:
        attr=f['attributes']
        label='PK' if '@id' in attr else ('FK' if f['name'].endswith('Id') else '')
        basics.append(f['name']+': '+f['type']+(f' [{label}]' if label else ''))
    detail_rows.append((m['name']+'\n'+m['table'],'; '.join(basics),str(m['line'])))
table('B.2','Trường vô hướng của thực thể trọng tâm',['Model / bảng','Field, kiểu và gợi ý khóa','Schema line'],detail_rows,[3.0,10.6,1.8],8.6)
para('Ký hiệu FK ở Bảng B.2 là gợi ý từ tên field kết thúc bằng Id. Khóa ngoại và onDelete/onUpdate đầy đủ phải đọc thuộc tính @relation và ràng buộc ở schema.prisma. Không suy diễn quan hệ chỉ từ tên trường nếu mô hình dùng khóa khác.')

enum_names={'ContainerVisitStatus','TruckVisitStatus','YardMovementStatus','ContainerInspectionStatus','ContainerInspectionResult','InYardBookingStatus','InYardBookingType','OperationalHoldStatus','OperationalHoldType','GatePassStatus'}
enum_rows=[(e['name'],', '.join(e['values']),str(e['line'])) for e in INV['enums'] if e['name'] in enum_names]
table('B.3','Enum trạng thái tác nghiệp cần giữ nguyên khi gọi API',['Enum','Giá trị','Schema line'],enum_rows,[4.2,9.9,1.3],9)

page_title('PHỤ LỤC C  QUYỀN VÀ MÀN HÌNH')
para('Các role dưới đây là bộ seed mặc định của dự án. DB có thể thay đổi cấu hình. Màn hiển thị là điều kiện navigation; backend vẫn giữ quyền quyết định từng API và trạng thái. AGENT/CONSIGNEE chưa có binding đến chủ hàng/đại lý ở schema, nên nghiệp vụ external-only bị chặn thay vì mở dữ liệu ngoài phạm vi.')
role_rows=[]
descriptions={
 'ADMIN':'Toàn bộ nhóm màn hiện có theo permission; dùng kiểm thử/điều hành.',
 'MANAGER':'Bãi, giám định, tra cứu, việc ca, theo quyền.',
 'OPERATOR':'Việc ca, tra cứu, Monitor; task có thể mở tab tác nghiệp theo quyền.',
 'GATE_STAFF':'Cổng, tra cứu, việc ca.',
 'YARD_STAFF':'Bãi, giám định, tra cứu, việc ca.',
 'AGENT':'Chỉ phần được scope hợp lệ; external-only đang bị chặn.',
 'CONSIGNEE':'Chỉ phần được scope hợp lệ; external-only đang bị chặn.',
}
for role in INV['roles']:
    role_rows.append((role['code'],str(len(role['permissions'])),descriptions.get(role['code'],'Theo permission backend')))
table('C.1','Role mặc định và khả năng mở Mobile',['Role','Số permission seed','Phạm vi Mobile'],role_rows,[3.1,1.6,10.7],10)
table('C.2','Từng màn TSX và vị trí mã nguồn',['Màn hình','Đường dẫn trong repo'],[(Path(s).stem,s) for s in INV['mobileScreens']],[4.0,11.4],9.0)
table('C.3','Bộ ảnh audit để cập nhật hình theo màn',['Nhóm màn','Ảnh đã dùng / ảnh bổ sung từ gallery','Điều cần kiểm khi thay'],[
    ('Đăng nhập / phiên','58; 03–06; 13–14; 100','Role, lỗi mạng, tên user và theme.'),
    ('Cổng nhập','07–16; 59; 69','Visit/chuyến, review và biên nhận.'),
    ('Cổng xuất','60–65; 85–86','QR, blocker, USED/EXITED; không lộ token.'),
    ('Bãi / vị trí','22–25; 33–35; 74–77; yard-populated-375','Slot canonical, Legend, nguồn MANUAL.'),
    ('Movement','26–32; 87–88','PENDING/IN_PROGRESS/COMPLETED/CANCELLED.'),
    ('Giám định','36–43; 78–79; survey-seal-375','Loại kiểm tra, HOLD, notes, severity.'),
    ('Booking','44–57','Ngày giờ, actual count/weight, lý do hủy.'),
    ('Tra cứu','17–20; 84; container-detail-375','Visit, location, hold và readiness.'),
    ('Việc ca / Monitor','50–51; 70–73; 80–84; work-queue-populated-375','Task ID, SLA, tổng các trang.'),
    ('Thông báo / Tài khoản','03–05; 21; 66–72; 81–83','Count, mark-read, popup và quyền.'),
    ('Năm tab sáng/tối','89–98','Ảnh cuối mỗi tab; không lấy ảnh before.'),
],[3.1,8.0,4.3],9.5)

page_title('PHỤ LỤC D  ĐƯỜNG DẪN MINH CHỨNG VÀ CÁCH CẬP NHẬT')
para('Ảnh audit vận hành nằm trong test-artifacts/2026-10-01-mobile-operations-audit, có gallery.html để xem 100 ảnh theo nhóm. Bộ smoke 04/10 nằm trong audit/runs/2026-10-03-improvement-02/screenshots/mobile-web-2026-10-04. Báo cáo kết quả tương ứng là REPORT.md và audit/2026-10-04-mobile-web-preview-report.md. Các file fixture-verify-*.json trong thư mục audit vận hành chứng minh trạng thái API/DB sau command.')
para('Sau khi thay ảnh: mở hình gốc ở cùng kích thước hiển thị; đọc ngày, role và database trong case; so khớp container/visit ID với nội dung đoạn viết; cập nhật caption của hình và mục kết quả test. Cuối cùng chọn toàn bộ Word và nhấn F9 để Word cập nhật Mục lục, Danh mục hình, Danh mục bảng và số trang. Nếu ảnh mới là Android native, ghi rõ thiết bị/phiên bản hệ điều hành thay vì nối nó với nhãn Expo Web cũ.')
para('Bản đồ truy vết: chương 1 dùng README và Business Spec; chương 2 dùng package.json, navigation, storage, Mobile Spec; chương 3 dùng permission mapping, API route và schema; chương 4 dùng TSX, caller, audit ảnh và verify JSON. Người đọc có thể kiểm từng nhận định bằng các nguồn [1]–[9].')

doc.core_properties.title='Xây dựng ứng dụng di động hỗ trợ vận hành ICD'
doc.core_properties.subject='Báo cáo Đồ án 4 Mobile ICD'
doc.core_properties.author='Lê Tiến Linh'
doc.core_properties.keywords='ICD; React Native; Expo; Gate-in; Gate-out; Yard'
settings=doc.settings.element
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUTPUT)
print(str(OUTPUT))
print('Paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'figures',len(doc.inline_shapes),'sections',len(doc.sections),'bytes',OUTPUT.stat().st_size)
