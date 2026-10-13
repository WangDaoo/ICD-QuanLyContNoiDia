from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

root = Path(r'D:\Project\Đồ Án 4 +Mobile')
work = root / 'icd-management/test-artifacts/2026-10-08-mobile-report'
path = root / '10123264 - Dao Minh Quang - Bao cao Mobile ICD - Theo mau goc.docx'
pdf = work / 'theo-mau-goc.pdf'
ref = work / 'DoAn3-mobile-reference.docx'
assert sha256(ref.read_bytes()).hexdigest().upper() == 'FA00C74F131462DDC97D675DE8087F82F84EF4F54ABB2EBEDF11B4164FB148EA'
with ZipFile(path) as z:
    assert z.testzip() is None
    media = [n for n in z.namelist() if n.startswith('word/media/')]
doc = Document(path)
paras = doc.paragraphs
body = '\n'.join(p.text for p in paras)
headings = [p.text for p in paras if p.style.name == 'Heading 3' and re.match(r'^3\.3\.\d+ Chức năng ', p.text)]
captions = [p.text for p in paras if p.style.name == 'HinhVe']
uc_figs = [p for p in captions if re.match(r'^Hình 3\.\d+[:.]',p) and 'Use Case' in p and '(UC' in p]
chapter3 = [int(re.match(r'^Hình 3\.(\d+)',p).group(1)) for p in captions if re.match(r'^Hình 3\.\d+[:.]',p)]
chapter4 = [int(re.match(r'^Hình 4\.(\d+)',p).group(1)) for p in captions if re.match(r'^Hình 4\.\d+\.',p)]
pages = PdfReader(str(pdf)).pages
short = [i + 1 for i,p in enumerate(pages) if len((p.extract_text() or '').strip()) < 100]
border = doc.sections[1].header.paragraphs[0]._p.get_or_add_pPr().find(qn('w:pBdr'))

assert len(headings) == 17, len(headings)
assert len(uc_figs) == 17, len(uc_figs)
assert sorted(chapter3) == list(range(1,27)), chapter3
assert sorted(chapter4) == list(range(1,24)), chapter4
assert len(media) == 53, len(media)
assert len(pages) == 113, len(pages)
assert not short, short
assert border is not None and border.find(qn('w:bottom')) is not None
assert doc.sections[1].footer.paragraphs[0].text.strip() == '1'
for expected in ['Đào Minh Quang','10123264','Vũ Xuân Thắng','Phát triển ứng dụng mobile đa nền tảng','1.4. Nội dung thực hiện','2.1. Quy trình phát triển ứng dụng','3.7. Tổng hợp các mô hình thiết kế']:
    assert expected.casefold() in body.casefold(), expected
for old in ['Lê Tiến Linh','12523040','Khung biểu đồ cần giữ khi cập nhật luận văn','Nếu người dùng chèn ảnh mới']:
    assert old.casefold() not in body.casefold(), old
assert body.count('•  Tác nhân (Actor):') == 17
assert body.count('•  Mục đích:') == 17
print('reference_sha=OK docx_zip=OK header_footer=OK')
print('use_cases=17 figures_ch3=26 figures_ch4=23 images=53')
print('pages=113 short_pages=[] identity=OK')
