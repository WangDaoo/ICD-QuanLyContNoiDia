from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document
from pypdf import PdfReader

root = Path(r'D:\Project\Đồ Án 4 +Mobile')
path = root / '10123264 - Dao Minh Quang - Bao cao Mobile ICD - Use Case.docx'
pdf = root / 'icd-management/test-artifacts/2026-10-08-mobile-report/usecase-final.pdf'

with ZipFile(path) as z:
    assert z.testzip() is None
    image_count = len([n for n in z.namelist() if n.startswith('word/media/')])

doc = Document(path)
paragraphs = [p.text.strip() for p in doc.paragraphs]
text = '\n'.join(paragraphs)
headings = [p.text.strip() for p in doc.paragraphs if p.style.name == 'Heading 3' and re.match(r'^3\.3\.\d+ Chức năng ', p.text)]
captions = [p.text.strip() for p in doc.paragraphs if p.style.name == 'HinhVe']
uc_captions = [p for p in captions if p.startswith('Hình 3.') and 'Use Case' in p and '(UC' in p]
caption_nums = [int(re.match(r'Hình 3\.(\d+)', p).group(1)) for p in captions if re.match(r'^Hình 3\.\d+\s*[:.]', p)]
pages = PdfReader(str(pdf)).pages
short_pages = [i + 1 for i, page in enumerate(pages) if len((page.extract_text() or '').strip()) < 100]

assert len(headings) == 17, (len(headings), headings)
assert len(uc_captions) == 17, (len(uc_captions), uc_captions)
assert sorted(caption_nums) == list(range(1, 27)), caption_nums
assert image_count == 53, image_count
assert len(pages) == 112, len(pages)
assert not short_pages, short_pages
for expected in ['ĐÀO MINH QUANG', '10123264', 'VŨ XUÂN THẮNG', 'PHÁT TRIỂN ỨNG DỤNG MOBILE ĐA NỀN TẢNG']:
    assert expected.casefold() in text.casefold(), expected
for old in ['Lê Tiến Linh', '12523040', 'nghe nhạc trực tuyến']:
    assert old.casefold() not in text.casefold(), old

print('docx_zip=OK')
print('uc_headings=', len(headings))
print('uc_captions=', len(uc_captions))
print('chapter_3_figures=', len(caption_nums), 'contiguous_1_to_26=OK')
print('images=', image_count)
print('pdf_pages=', len(pages), 'short_pages=', short_pages)
print('identity=OK')
