from __future__ import annotations

import copy
import hashlib
import json
import math
import re
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r'D:\Project\Đồ Án 4 +Mobile')
WORK = ROOT / 'icd-management/test-artifacts/2026-10-08-mobile-report'
SRC = ROOT / '10123264 - Dao Minh Quang - Bao cao Mobile ICD.docx'
OUT = ROOT / '10123264 - Dao Minh Quang - Bao cao Mobile ICD - Use Case.docx'
FIG = WORK / 'use-case-figures'
FIG.mkdir(exist_ok=True)
REFERENCE = WORK / 'DoAn3-mobile-reference.docx'
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest().upper() == 'FA00C74F131462DDC97D675DE8087F82F84EF4F54ABB2EBEDF11B4164FB148EA'

CASES = [
 dict(actor='Nhân viên ICD', title='Đăng nhập và khôi phục phiên', purpose='Xác thực nhân viên, nhận phạm vi ICD và quyền thao tác để mở đúng chức năng trong ứng dụng.', included=['Xác thực phiên với backend','Đọc người dùng và quyền','Khởi tạo các tab được phép'], optional=[('Làm mới token','Token hết hạn nhưng phiên có thể làm mới')], steps=[
 'Nhân viên mở ứng dụng. Khi chưa có phiên, màn đăng nhập hiển thị các trường email và mật khẩu.',
 'Nhân viên nhập thông tin và gửi yêu cầu đăng nhập. Ứng dụng hiển thị trạng thái đang xử lý, tránh gửi lặp cùng một yêu cầu.',
 'Backend kiểm tra tài khoản, trả thông tin phiên và người dùng khi thông tin hợp lệ.',
 'Ứng dụng đọc quyền của người dùng, xác định các tab được phép và mở màn phù hợp.',
 'Ở lần mở lại, ứng dụng kiểm tra phiên đã lưu và đọc hồ sơ hiện tại trước khi cho phép thao tác nghiệp vụ.'
 ], alternatives=[('Sai thông tin đăng nhập','Hiển thị lỗi xác thực, cho phép sửa thông tin và gửi lại; không mở màn nghiệp vụ.'),('Phiên cần làm mới','Thực hiện luồng refresh theo cơ chế phiên đang có. Khi backend từ chối phiên, xóa phiên cũ và quay về đăng nhập.'),('Mạng hoặc API không đáp ứng','Thông báo lỗi kết nối và cung cấp thao tác thử lại. Không suy diễn rằng người dùng có quyền chỉ từ dữ liệu cũ.')]),
 dict(actor='Nhân viên cổng', title='Đánh dấu xe đến', purpose='Ghi nhận xe đã có mặt tại ICD theo đúng Truck Visit, làm cơ sở tiếp nhận container.', included=['Tìm và đối chiếu chuyến xe','Kiểm tra quyền và trạng thái','Ghi nhận xe ARRIVED'], optional=[], steps=[
 'Nhân viên tìm chuyến xe theo thông tin đang có trong màn Cổng.',
 'Ứng dụng hiển thị Truck Visit và thông tin container liên quan để nhân viên đối chiếu với xe thực tế.',
 'Nhân viên xác nhận xe đã đến; ứng dụng gửi yêu cầu arrive cho đúng ID chuyến.',
 'Backend kiểm tra quyền, phạm vi ICD và trạng thái chuyến trước khi cập nhật.',
 'Ứng dụng nhận kết quả ARRIVED và tải lại bối cảnh tiếp nhận.'
 ], alternatives=[('Không tìm thấy chuyến phù hợp','Giữ màn tìm kiếm và hướng dẫn đối chiếu chuyến đã được bộ phận điều hành lập.'),('Chuyến đã được xử lý hoặc thiếu quyền','Hiển thị lý do từ backend và trạng thái hiện tại; không tự đổi trạng thái chuyến trên giao diện.')]),
 dict(actor='Nhân viên cổng', title='Tiếp nhận container vào cổng', purpose='Ghi thông tin container tại thời điểm vào ICD và tạo biên nhận Gate-in trên hồ sơ visit.', included=['Đọc bối cảnh Gate-in','Kiểm tra seal cân và tình trạng','Ghi nhận tiếp nhận container'], optional=[('Nhập số container thủ công','Không quét được mã hoặc camera không sẵn sàng')], steps=[
 'Nhân viên quét mã hoặc nhập số container; ứng dụng lấy bối cảnh Gate-in của visit tương ứng.',
 'Nhân viên đối chiếu container, lệnh và chuyến xe với thông tin thực tế tại cổng.',
 'Nhân viên nhập seal thực tế, khối lượng và tình trạng container theo biểu mẫu hiện có.',
 'Ứng dụng hiển thị phần xem lại để nhân viên kiểm tra trước khi xác nhận tiếp nhận.',
 'Backend kiểm tra điều kiện nghiệp vụ, ghi reception và event; ứng dụng chỉ hiển thị thành công sau phản hồi xác nhận.'
 ], alternatives=[('Nhập mã thay cho camera','Tại điểm đọc mã, nhân viên chuyển sang nhập thủ công; số container vẫn được đối chiếu với backend.'),('Dữ liệu thiếu hoặc không hợp lệ','Hiển thị lỗi gần trường liên quan, giữ thông tin đã nhập và cho phép sửa.'),('Visit đã tiếp nhận hoặc kết quả chưa xác định','Đọc lại trạng thái hồ sơ trước khi tiếp tục. Với timeout, không tự gửi lại giao dịch để tránh tiếp nhận trùng.')]),
 dict(actor='Nhân viên cổng', title='Quét phiếu ra cổng', purpose='Tra cứu Gate Pass bằng token để kiểm tra phiếu và mở bối cảnh xác nhận xe ra cổng.', included=['Đọc token Gate Pass','Kiểm tra trạng thái và thời hạn','Trả thông tin phiếu và container'], optional=[('Nhập token thủ công','Camera không có quyền hoặc không đọc được QR')], steps=[
 'Nhân viên mở màn quét Gate Pass và đọc QR trên phiếu do backend đã phát hành.',
 'Ứng dụng gửi token đến endpoint scan; token được dùng để tìm đúng phiếu, không suy ra dữ liệu chỉ từ hình QR.',
 'Backend kiểm tra trạng thái phiếu và trả thông tin container, thời hạn cùng các điều kiện liên quan.',
 'Nhân viên đối chiếu thông tin phiếu với xe và container đang tại cổng.',
 'Ứng dụng chỉ chuyển đến màn xác nhận khi phản hồi scan cung cấp phiếu và visit phù hợp.'
 ], alternatives=[('Thiếu quyền camera hoặc QR khó đọc','Nhân viên nhập token bằng luồng dự phòng và thực hiện cùng bước kiểm tra backend.'),('Phiếu hết hạn đã dùng hoặc bị hủy','Hiển thị trạng thái bị từ chối; không mở thao tác Gate-out.'),('Token không tìm thấy hoặc lỗi kết nối','Giữ khả năng nhập/quét lại và thông báo cách kiểm tra phiếu; không tạo phiếu mới từ Mobile.')]),
 dict(actor='Nhân viên cổng', title='Xác nhận Gate-out', purpose='Xác nhận container đã ra ICD bằng Gate Pass hợp lệ và ghi trạng thái sau giao dịch.', included=['Kiểm lại phiếu và readiness','Xác nhận giao dịch ra cổng','Ghi EXITED USED và giải phóng slot'], optional=[], steps=[
 'Nhân viên mở bối cảnh Gate-out từ phiếu vừa được backend kiểm tra.',
 'Nhân viên đối chiếu xe, container và phiếu với tình huống thực tế; xem các blocker được trả về.',
 'Nhân viên xác nhận xe ra cổng. Ứng dụng khóa thao tác gửi lặp trong khi yêu cầu đang xử lý.',
 'Backend kiểm lại token, thời hạn và readiness tại thời điểm xác nhận rồi ghi thay đổi trong transaction.',
 'Ứng dụng hiển thị biên nhận từ phản hồi server và tải lại hồ sơ visit, phiếu cùng vị trí bãi.'
 ], alternatives=[('Phát sinh Hold hoặc blocker sau khi quét','Backend từ chối Gate-out; ứng dụng hiển thị lý do hiện tại và yêu cầu xử lý blocker.'),('Phiếu đổi trạng thái hoặc bị thu hồi','Dừng xác nhận, đọc lại trạng thái phiếu; không dùng kết quả scan cũ để vượt kiểm tra.'),('Xung đột hoặc timeout','Giữ bối cảnh để đối chiếu kết quả. Không tự replay lệnh ra cổng khi chưa biết giao dịch đã được ghi hay chưa.')]),
 dict(actor='Nhân viên bãi', title='Kiểm tra và xếp vị trí bãi', purpose='Chọn vị trí hợp lệ và gán container vào slot bằng chế độ MANUAL.', included=['Đọc sơ đồ slot và Legend','Kiểm tra vị trí được chọn','Gán vị trí MANUAL'], optional=[], steps=[
 'Nhân viên mở bãi và chọn container có quyền xử lý.',
 'Ứng dụng tải danh sách vị trí cùng trạng thái và Legend để nhân viên quan sát vị trí hiện tại.',
 'Nhân viên chọn slot; ứng dụng gửi yêu cầu check và hiển thị blockers hoặc warnings được trả về.',
 'Khi vị trí được backend cho phép, nhân viên xác nhận gán bằng chế độ MANUAL.',
 'Backend kiểm lại điều kiện và ghi vị trí canonical cùng lịch sử; ứng dụng tải lại container và slot.'
 ], alternatives=[('Slot bị chiếm khóa hoặc bảo trì','Dừng gán, giữ container đang chọn và cho phép chọn slot khác.'),('Container hoặc block không phù hợp','Hiển thị kết quả check và điều kiện từ backend để nhân viên điều chỉnh lựa chọn.'),('Chọn chức năng tối ưu xếp vị trí','Hiển thị “Đang trong quá trình phát triển”. Ca sử dụng này chưa mô tả thuật toán tối ưu mới.')]),
 dict(actor='Nhân viên bãi', title='Yêu cầu di chuyển nội bộ', purpose='Tạo lệnh movement cho container giữa vị trí nguồn và đích trong bãi.', included=['Chọn container và slot đích','Kiểm tra yêu cầu di chuyển','Tạo movement PENDING'], optional=[], steps=[
 'Nhân viên chọn container đang trong bãi và mở chức năng tạo yêu cầu di chuyển.',
 'Ứng dụng hiển thị vị trí nguồn hiện tại; nhân viên chọn slot đích cần chuyển đến.',
 'Nhân viên kiểm tra thông tin và gửi yêu cầu movement cho visit đang chọn.',
 'Backend kiểm quyền, nguồn và điều kiện của đích; khi hợp lệ, ghi operation PENDING.',
 'Ứng dụng mở chi tiết operation mới để nhân viên tiếp tục theo dõi thực hiện.'
 ], alternatives=[('Nguồn không còn đúng hoặc đích không khả dụng','Hiển thị lý do, đọc lại vị trí hiện tại và giữ dữ liệu để chọn lại.'),('Không tạo được lệnh','Giữ bản nháp và cho phép thử lại sau khi xác định lỗi; vị trí container chưa được hiển thị như đã chuyển.')]),
 dict(actor='Nhân viên bãi', title='Bắt đầu và hoàn tất di chuyển', purpose='Theo dõi thực hiện movement và ghi vị trí mới sau khi container đã được chuyển ngoài bãi.', included=['Bắt đầu lệnh di chuyển','Xác nhận hoàn tất movement','Tải lại vị trí và lịch sử'], optional=[], steps=[
 'Nhân viên mở chi tiết movement và đối chiếu container, slot nguồn, slot đích.',
 'Với lệnh đủ điều kiện, nhân viên chọn bắt đầu; backend cập nhật trạng thái thực hiện.',
 'Nhân viên thực hiện di chuyển container ngoài hiện trường theo lệnh đang theo dõi.',
 'Sau khi đối chiếu vị trí thực tế, nhân viên xác nhận hoàn tất; ứng dụng gửi complete cho đúng operation ID.',
 'Backend ghi kết quả và location log; ứng dụng đọc lại chi tiết lệnh cùng vị trí mới.'
 ], alternatives=[('Lệnh đã đổi trạng thái','Hiển thị trạng thái mới và tải lại operation thay vì áp dụng phản hồi cũ.'),('Đích hoặc container phát sinh xung đột','Hiển thị điều kiện không đáp ứng, giữ liên kết operation để nhân viên xử lý.'),('Hoàn tất thất bại','Không báo đã chuyển thành công; đối chiếu trạng thái server trước lần gửi tiếp theo.')]),
 dict(actor='Nhân viên bãi', title='Yêu cầu giám định', purpose='Tạo phiếu kiểm tra container với loại kiểm tra và ghi chú phù hợp.', included=['Chọn loại kiểm tra','Nhập ghi chú hiện trường','Tạo inspection PENDING'], optional=[('Nhập mức độ hư hỏng','Loại kiểm tra yêu cầu thông tin hư hỏng')], steps=[
 'Nhân viên chọn container trong màn Giám định và kiểm tra đúng hồ sơ visit.',
 'Nhân viên chọn loại kiểm tra đang được hệ thống hỗ trợ.',
 'Biểu mẫu hiển thị các thông tin theo loại kiểm tra; nhân viên nhập ghi chú thực tế và mức hư hỏng khi phù hợp.',
 'Ứng dụng gửi yêu cầu tạo inspection; backend kiểm tra dữ liệu và quyền của người thực hiện.',
 'Sau khi tạo thành công, ứng dụng tải lại lịch sử kiểm tra của container.'
 ], alternatives=[('Kiểm tra seal hoặc tình trạng','Dùng nội dung phù hợp với loại kiểm tra; không bắt nhập mức hư hỏng chỉ vì biểu mẫu khác có trường này.'),('Thiếu trường bắt buộc','Chỉ rõ trường cần sửa và giữ nội dung đã nhập.'),('Tạo phiếu thất bại','Hiển thị lỗi riêng của thao tác; không dùng danh sách trống để thay cho lỗi tải lịch sử.')]),
 dict(actor='Nhân viên bãi', title='Kết thúc giám định', purpose='Ghi kết quả PASS FAIL hoặc HOLD và đọc lại các điều kiện nghiệp vụ liên quan.', included=['Đọc phiếu kiểm tra hiện tại','Ghi kết quả và ghi chú','Đọc lại readiness và blockers'], optional=[('Bổ sung lý do FAIL hoặc HOLD','Kết quả cần giải thích việc không đạt hoặc tạm giữ')], steps=[
 'Nhân viên mở inspection cần xử lý và đọc loại kiểm tra cùng trạng thái hiện tại.',
 'Nếu phiếu cần bắt đầu, nhân viên thực hiện start trước khi ghi hoàn tất.',
 'Nhân viên nhập kết quả kiểm tra thực tế và ghi chú; với FAIL hoặc HOLD, bổ sung lý do theo yêu cầu dữ liệu.',
 'Ứng dụng gửi complete cho inspection ID đang mở; backend kiểm trạng thái và ghi kết quả.',
 'Ứng dụng tải lại inspection cùng readiness để hiển thị tác động nghiệp vụ theo dữ liệu backend.'
 ], alternatives=[('Thiếu lý do cho kết quả không đạt','Giữ lựa chọn kết quả và chỉ rõ phần cần bổ sung.'),('Phiếu đã hoàn tất hoặc trạng thái khác','Đọc lại phiếu và thông báo trạng thái hiện tại; không ghi đè bằng nội dung cũ.'),('Backend từ chối hoặc kết nối gián đoạn','Không hiển thị hoàn tất. Nhân viên kiểm lại trạng thái trước khi tiếp tục.')]),
 dict(actor='Nhân viên bãi', title='Đặt lịch tác nghiệp bãi', purpose='Tạo booking cho rút hàng đóng hàng hoặc kiểm tra với thời gian được ghi đúng.', included=['Chọn loại booking','Chọn ngày giờ tác nghiệp','Tạo lịch booking'], optional=[], steps=[
 'Nhân viên chọn container còn trong bãi và mở biểu mẫu booking.',
 'Nhân viên chọn loại STRIPPING STUFFING hoặc INSPECTION theo danh sách hỗ trợ.',
 'Nhân viên chọn ngày và giờ bằng picker; ứng dụng giữ giá trị trước đó nếu hủy picker.',
 'Nhân viên nhập thông tin theo biểu mẫu và xác nhận. Thời gian hiển thị tại Việt Nam được chuyển thành instant gửi backend.',
 'Backend ghi booking hợp lệ; ứng dụng đọc lại lịch và trạng thái PENDING.'
 ], alternatives=[('Thời gian không hợp lệ','Chỉ rõ ngày giờ cần sửa và giữ bản nháp; không gửi giá trị sai để hiển thị như đã đặt lịch.'),('Hủy chọn ngày giờ','Giữ thời điểm đã có trước khi mở picker.'),('Thiếu dữ liệu theo loại booking','Hiển thị validation của trường liên quan theo DTO hiện tại.')]),
 dict(actor='Nhân viên bãi', title='Hoàn tất hoặc hủy booking', purpose='Ghi kết quả thực hiện hoặc lý do hủy lịch tác nghiệp đang được theo dõi.', included=['Đọc booking và trạng thái','Kiểm tra quyền chuyển trạng thái','Ghi kết quả xử lý booking'], optional=[('Nhập lý do hủy booking','Nhân viên chọn hủy lịch thay cho hoàn tất')], steps=[
 'Nhân viên mở chi tiết booking và đọc container, loại tác nghiệp, thời gian cùng trạng thái.',
 'Nhân viên bắt đầu tác nghiệp khi lịch đang ở trạng thái được phép.',
 'Sau khi thao tác ngoài bãi, nhân viên chọn hoàn tất và nhập kết quả thực tế; trường hợp dừng lịch thì chọn hủy và ghi lý do.',
 'Ứng dụng gửi start complete hoặc cancel theo đúng hành động; backend kiểm trạng thái trước khi lưu.',
 'Ứng dụng tải lại booking và chỉ hiển thị COMPLETED hoặc CANCELLED sau kết quả xác nhận.'
 ], alternatives=[('Hủy booking','Luồng mở rộng chỉ được kích hoạt khi người dùng chọn hủy; lý do được gửi cùng yêu cầu cancel.'),('Kết quả hoặc lý do không hợp lệ','Giữ dữ liệu hiện tại, hiển thị trường cần sửa; không tự điền số liệu thực tế.'),('Booking đã đổi trạng thái','Hiển thị xung đột và tải lại detail để tránh thao tác trên trạng thái cũ.')]),
 dict(actor='Nhân viên có quyền đọc', title='Tra cứu container', purpose='Tìm container và xem hồ sơ hiện trường trong phạm vi quyền của người dùng.', included=['Tìm kiếm và phân trang','Đọc hồ sơ theo quyền','Đọc vị trí sự kiện và trạng thái'], optional=[], steps=[
 'Nhân viên nhập số container hoặc tiêu chí tìm kiếm có trên màn Tra cứu.',
 'Ứng dụng lấy danh sách và thông tin phân trang từ backend.',
 'Nhân viên chọn record; ứng dụng mở chi tiết bằng đúng visit ID của record đó.',
 'Ứng dụng đọc các phần hồ sơ được phép như vị trí events holds readiness phiếu hoặc bàn giao.',
 'Nhân viên sử dụng thông tin hiện tại và thời điểm cập nhật để đối chiếu tình trạng container.'
 ], alternatives=[('Không có record phù hợp','Hiển thị trạng thái không có kết quả và cho phép thay hoặc xóa bộ lọc.'),('Một phần hồ sơ tải lỗi','Giữ phần dữ liệu đã có khi phù hợp, hiển thị lỗi và retry riêng; không báo hồ sơ đầy đủ.'),('Không có quyền hoặc dữ liệu cache không còn phù hợp','Ẩn dữ liệu bị hạn chế và không dùng cache để vượt kiểm tra quyền.')]),
 dict(actor='Nhân viên theo quyền nhiệm vụ', title='Theo dõi việc ca', purpose='Xem hàng đợi công việc và mở đúng chức năng để xử lý hồ sơ cần làm.', included=['Đọc hàng đợi và tổng hợp','Xác định đúng entity đích','Kiểm tra quyền khi mở nhiệm vụ'], optional=[('Lọc loại và mức ưu tiên','Nhân viên muốn thu hẹp danh sách việc ca')], steps=[
 'Nhân viên mở Việc ca; ứng dụng đọc queue trong phạm vi user và ICD.',
 'Ứng dụng hiển thị nhiệm vụ, số lượng và tình trạng quá hạn từ dữ liệu nhận được.',
 'Nhân viên có thể lọc theo loại hoặc mức ưu tiên trước khi chọn một nhiệm vụ.',
 'Ứng dụng xác định đúng entity ID và kiểm quyền màn đích trước khi điều hướng.',
 'Sau tác nghiệp, nhân viên quay lại queue và cập nhật danh sách để thấy trạng thái mới từ backend.'
 ], alternatives=[('Áp dụng bộ lọc','Danh sách phản ánh tiêu chí đang chọn; trường hợp không khớp filter khác với hàng đợi trống.'),('Nhiệm vụ đã đổi trạng thái hoặc mất quyền','Thông báo trạng thái mới và hướng quay lại queue, tránh mở record không liên quan.'),('Loại task chưa có producer backend','Không tạo nhiệm vụ giả. Việc mô hình hóa destination không chứng minh backend đã sinh mọi loại task.')]),
 dict(actor='Nhân viên ICD', title='Đọc và đánh dấu thông báo', purpose='Đọc thông báo của người dùng và đồng bộ số lượng chưa đọc sau thao tác.', included=['Đọc lịch sử thông báo','Đọc số lượng chưa đọc','Mở nội dung bản ghi'], optional=[('Đánh dấu một hoặc tất cả đã đọc','Nhân viên chọn thao tác đánh dấu đã đọc')], steps=[
 'Nhân viên mở Thông báo; ứng dụng tải history và unread count theo user.',
 'Nhân viên đọc danh sách, mở bản ghi và xem nội dung hoặc record liên quan khi được phép.',
 'Nhân viên chọn đánh dấu một thông báo hoặc đánh dấu tất cả theo thao tác có trên màn.',
 'Ứng dụng gửi command tương ứng và hiển thị pending khi đang xử lý.',
 'Sau phản hồi thành công, ứng dụng tải lại trạng thái đọc và count để đồng bộ với backend.'
 ], alternatives=[('Lịch sử chưa tải được','Hiển thị lỗi và nút retry; chỉ hiển thị empty sau response đọc thành công.'),('Đánh dấu đọc thất bại','Giữ count đã được xác nhận trước đó và cho phép thử lại, không tạo count âm hoặc giả bằng không.'),('Record liên quan không còn quyền truy cập','Thông báo giới hạn truy cập và giữ khả năng quay lại danh sách thông báo.')]),
 dict(actor='Nhân viên ICD', title='Xem tài khoản và thay theme', purpose='Xem thông tin phiên và cấu hình giao diện, đồng thời cho phép kết thúc phiên từ màn Tài khoản.', included=['Đọc thông tin tài khoản','Đọc vai trò và permission'], optional=[('Đổi giao diện sáng tối','Nhân viên chọn thay theme'),('Đăng xuất','Nhân viên chọn kết thúc phiên')], steps=[
 'Nhân viên mở Tài khoản kể cả trường hợp không có tab nghiệp vụ được cấp quyền.',
 'Ứng dụng hiển thị người dùng, vai trò, permission và thông tin kết nối đang áp dụng.',
 'Khi chọn đổi sáng tối, ứng dụng áp dụng theme và lưu preference bằng cơ chế lưu trữ đang có.',
 'Khi chọn đăng xuất, ứng dụng kết thúc phiên theo cơ chế auth và xóa dữ liệu phiên liên quan.',
 'Ở lần mở tiếp theo, ứng dụng nạp theme đã lưu và kiểm tra phiên trước khi mở màn nghiệp vụ.'
 ], alternatives=[('Đổi theme','Đây là hành động tùy chọn khi xem tài khoản; không thay đổi quyền hoặc trạng thái nghiệp vụ.'),('Lưu preference thất bại','Sử dụng lựa chọn khả dụng và hiển thị lỗi phù hợp; không xem lỗi lưu theme là lỗi xác nhận giao dịch.'),('Đăng xuất với tài khoản không có tab','Tài khoản vẫn truy cập được màn Tài khoản để kết thúc phiên; không bị kẹt trong màn không có thao tác.')]),
 dict(actor='Điều độ hoặc quản lý', title='Theo dõi bãi và bàn giao', purpose='Xem thông tin giám sát bãi và tóm tắt bàn giao theo quyền đọc được cấp.', included=['Đọc thông tin bãi theo quyền','Xem trạng thái container'], optional=[('Xem tóm tắt bàn giao','Người dùng có handover.read và chọn hồ sơ bàn giao')], steps=[
 'Người dùng mở màn Monitor hoặc nội dung theo dõi từ record đang chọn.',
 'Ứng dụng đọc dữ liệu bãi hoặc visit trong phạm vi permission và ICD.',
 'Người dùng xem các trạng thái, vị trí và thông tin container mà backend đã trả.',
 'Khi có quyền đọc bàn giao, người dùng mở summary của visit tương ứng.',
 'Ứng dụng hiển thị thông tin bàn giao từ API hiện có và giữ khả năng quay lại record nguồn.'
 ], alternatives=[('Không có quyền đọc một resource','Chỉ rõ phần bị giới hạn, không lấy dữ liệu qua Partner API thay cho quyền của nhân viên.'),('Bàn giao chưa có hoặc resource tải lỗi','Phân biệt không có hồ sơ với lỗi kết nối; cung cấp retry cho resource lỗi.'),('Dữ liệu đã thay đổi','Tải lại trạng thái hiện tại; không tạo hoặc xác nhận bàn giao chỉ bằng thao tác xem summary.')]),
]


class Canvas:
    def __init__(self, name, width=2400, height=1500):
        self.name, self.width, self.height = name, width, height
        self.im = Image.new('RGB', (width, height), 'white')
        self.d = ImageDraw.Draw(self.im)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="white"/>']

    def font(self, size, bold=False):
        return ImageFont.truetype(r'C:\Windows\Fonts\timesbd.ttf' if bold else r'C:\Windows\Fonts\times.ttf', size)

    def line(self, a, b, dashed=False, arrow=False, width=5):
        x1,y1=a; x2,y2=b
        if dashed:
            dist=math.hypot(x2-x1,y2-y1)
            for pos in range(0, int(dist), 26):
                lo=pos/dist; hi=min(pos+15,dist)/dist
                self.d.line((x1+(x2-x1)*lo,y1+(y2-y1)*lo,x1+(x2-x1)*hi,y1+(y2-y1)*hi),fill='black',width=width)
        else: self.d.line((x1,y1,x2,y2),fill='black',width=width)
        dash=' stroke-dasharray="15 11"' if dashed else ''
        self.svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="black" stroke-width="{width}"{dash}/>')
        if arrow:
            angle=math.atan2(y2-y1,x2-x1)
            for delta in [-.5,.5]:
                p=(x2-25*math.cos(angle+delta),y2-25*math.sin(angle+delta))
                self.line(b,p,width=width)

    def text(self, x,y,text,size=54,bold=False,background=False):
        f=self.font(size,bold); box=self.d.textbbox((x,y),text,font=f,anchor='mm')
        if background:
            self.d.rectangle((box[0]-9,box[1]-6,box[2]+9,box[3]+6),fill='white')
            self.svg.append(f'<rect x="{box[0]-9}" y="{box[1]-6}" width="{box[2]-box[0]+18}" height="{box[3]-box[1]+12}" fill="white"/>')
        self.d.text((x,y),text,font=f,fill='black',anchor='mm')
        self.svg.append(f'<text x="{x}" y="{y}" text-anchor="middle" dominant-baseline="central" font-family="Times New Roman" font-size="{size}" font-weight="{"bold" if bold else "normal"}">{escape(text)}</text>')

    def wrap(self, text, width, size=54):
        lines=[]
        for raw in text.split('\n'):
            line=''
            for word in raw.split():
                candidate=(line+' '+word).strip()
                if self.d.textlength(candidate,font=self.font(size))>width and line: lines.append(line);line=word
                else: line=candidate
            lines.append(line)
        return lines

    def oval(self, cx,cy,label,w=740,h=220,size=54):
        rect=(cx-w/2,cy-h/2,cx+w/2,cy+h/2)
        self.d.ellipse(rect,fill='#9DD7F0',outline='black',width=5)
        self.svg.append(f'<ellipse cx="{cx}" cy="{cy}" rx="{w/2}" ry="{h/2}" fill="#9DD7F0" stroke="black" stroke-width="5"/>')
        lines=self.wrap(label,w-150,size)
        assert len(lines)*(size+7)<h-14, (label,lines)
        for n,line in enumerate(lines): self.text(cx,cy+(n-(len(lines)-1)/2)*(size+7),line,size)
        return (cx,cy,w/2,h/2)

    def actor(self,x,y,label,compact=False):
        if compact:
            radius=31; head=y-65; top=y-32; waist=y+65; foot=y+150
            arm=75; leg=72; label_y=y+210; label_size=46; lineheight=53
        else:
            radius=42; head=y-83; top=y-40; waist=y+105; foot=y+215
            arm=88; leg=80; label_y=y+270; label_size=52; lineheight=59
        self.d.ellipse((x-radius,head-radius,x+radius,head+radius),outline='black',width=5)
        self.svg.append(f'<circle cx="{x}" cy="{head}" r="{radius}" fill="none" stroke="black" stroke-width="5"/>')
        self.line((x,top),(x,waist)); self.line((x-arm,y+10),(x+arm,y+10))
        self.line((x,waist),(x-leg,foot));self.line((x,waist),(x+leg,foot))
        for n,line in enumerate(self.wrap(label,370,label_size)):
            self.text(x,label_y+n*lineheight,line,label_size)

    def boundary(self,title):
        self.d.rectangle((425,35,self.width-30,self.height-30),outline='black',width=5)
        self.svg.append(f'<rect x="425" y="35" width="{self.width-455}" height="{self.height-65}" fill="none" stroke="black" stroke-width="5"/>')
        self.text((425+self.width-30)/2,100,title,58,True)

    def relation(self,a,b,kind):
        dx=b[0]-a[0];dy=b[1]-a[1]
        fa=1/math.sqrt((dx/a[2])**2+(dy/a[3])**2)
        fb=1/math.sqrt((dx/b[2])**2+(dy/b[3])**2)
        p=(a[0]+dx*fa,a[1]+dy*fa);q=(b[0]-dx*fb,b[1]-dy*fb)
        self.line(p,q,dashed=True,arrow=True,width=4)
        mx=(p[0]+q[0])/2;my=(p[1]+q[1])/2
        if abs(dx)<60: mx+=165
        else: my-=28
        self.text(mx,my,'«'+kind+'»',46,background=True)

    def save(self):
        self.svg.append('</svg>')
        self.im.save(FIG/(self.name+'.png'),dpi=(300,300))
        (FIG/(self.name+'.svg')).write_text('\n'.join(self.svg),encoding='utf-8')
        return FIG/(self.name+'.png')


def diagram(i,case):
    c=Canvas(f'UC{i:02d}')
    c.boundary(f'Hệ thống ICD phục vụ Mobile - UC{i:02d}')
    c.actor(190,610,case['actor'])
    main=(850,680,320,130)
    c.line((190,660),(530,680))
    children=[]
    ys=[340,680,1020] if len(case['included'])==3 else [390,860]
    for label,y in zip(case['included'],ys): children.append((1900,y,370,110))
    for target in children: c.relation(main,target,'include')
    optionals=[]
    if len(case['optional'])==1: optionals=[(850,1260,330,100)]
    elif len(case['optional'])==2: optionals=[(800,1260,310,100),(1500,1260,310,100)]
    for optional in optionals:c.relation(optional,main,'extend')
    c.oval(850,680,f'UC{i:02d}\n'+case['title'],w=640,h=260,size=53)
    for label,y in zip(case['included'],ys):c.oval(1900,y,label,size=54)
    for (label,_),box in zip(case['optional'],optionals):c.oval(box[0],box[1],label,w=box[2]*2,h=box[3]*2,size=50)
    return c.save()


def overview():
    c=Canvas('UC-tong-quat',2400,1680)
    c.boundary('Các nhóm Use Case của hệ thống ICD trên Mobile')
    actors=[(170,180,'Nhân viên ICD'),(170,520,'Nhân viên cổng'),(170,900,'Nhân viên bãi'),(170,1370,'Điều độ / Quản lý')]
    groups=[(1050,250,'UC01 UC15 UC16\nPhiên thông báo tài khoản'),(1820,500,'UC02 UC03\nTiếp nhận vào cổng'),(1050,690,'UC04 UC05\nQuét phiếu và Gate-out'),(1820,930,'UC06 UC07 UC08\nVị trí và di chuyển bãi'),(1050,1140,'UC09 UC10 UC11 UC12\nGiám định và booking'),(1820,1430,'UC13 UC14 UC17\nTra cứu việc ca giám sát')]
    for ai,gi in [(0,0),(1,1),(1,2),(2,3),(2,4),(3,5)]:
        ax,ay,_=actors[ai];gx,gy,_=groups[gi];c.line((ax+75,ay+10),(gx-340,gy))
    for x,y,label in groups:c.oval(x,y,label,w=680,h=240,size=50)
    for x,y,label in actors:c.actor(x,y,label,compact=True)
    return c.save()


doc=Document(SRC)
old_sections=[]
heads=[p for p in doc.paragraphs if p.style.name=='Heading 3' and re.match(r'3\.3\.\d+ ',p.text)]
assert len(heads)==17
for h in heads:
    values={}
    current=h._p.getnext()
    while current is not None:
        if current.tag==qn('w:p'):
            from docx.text.paragraph import Paragraph
            p=Paragraph(current,doc._body)
            if p.style.name in ('Heading 2','Heading 3'):break
            text=p.text
            for label,key in [('Mã và tác nhân.','actor'),('Tiền điều kiện.','pre'),('Luồng chính.','main'),('Ngoại lệ.','alt'),('Hậu điều kiện.','post'),('Đối chiếu API.','api')]:
                if text.startswith(label):values[key]=text[len(label):].strip()
        current=current.getnext()
    assert set(values)=={'actor','pre','main','alt','post','api'},(h.text,values)
    old_sections.append(values)

# Existing chapter 3 diagrams after the state model follow the 17 new use-case figures.
def renumber(container):
    for p in container.paragraphs:
        for run in p.runs:
            updated=re.sub(r'Hình 3\.([3-9])(?=[^0-9]|$)',lambda m:'Hình 3.'+str(int(m.group(1))+17),run.text)
            if updated != run.text: run.text=updated
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:renumber(cell)
renumber(doc)

def newp(anchor,text='',style='Nội dung',bold=False,keep=False):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next=keep
    p.paragraph_format.keep_together=True
    if text:
        r=p.add_run(text);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    anchor.addprevious(p._p)
    return p

def label(anchor,text):
    return newp(anchor,text,bold=True,keep=True)

def picture(anchor,path,description):
    p=newp(anchor,style='Normal',keep=True)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(3)
    run=p.add_run(); inline=run.add_picture(str(path),width=Cm(15.5))
    inline._inline.docPr.set('descr',description)
    return p

def caption(anchor,text):
    p=newp(anchor,text,'HinhVe')
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.3
    p.paragraph_format.space_after=Pt(7)
    for r in p.runs:r.font.size=Pt(12)
    return p

# Replace the old flowchart-like actor picture with an actual actor/ellipse overview.
overview_path=overview()
caption1=next(p for p in doc.paragraphs if p.style.name=='HinhVe' and p.text.startswith('Hình 3.1.'))
old_picture=caption1._p.getprevious()
assert old_picture is not None and old_picture.findall('.//'+qn('w:drawing'))
old_picture.getparent().remove(old_picture)
picture(caption1._p,overview_path,'Sơ đồ Use Case tổng quát: tác nhân nhân viên ICD, cổng, bãi và điều độ hoặc quản lý liên kết với các nhóm UC01 đến UC17.')
caption1.runs[0].text='Hình 3.1: Use Case tổng quát của hệ thống ICD trên Mobile'

intro=next(p for p in doc.paragraphs if p.text.startswith('Mỗi ca sử dụng dưới đây'))
intro.text='Các chức năng dưới đây được trình bày theo cùng mẫu: hình Use Case, tác nhân, mục đích, điều kiện tiên quyết, luồng sự kiện chính, các luồng mở rộng và điều kiện kết thúc. Mã UC01–UC17 liên kết với màn hình và API của Mobile. Sơ đồ mô tả chức năng và điều kiện sử dụng; kết quả kiểm thử vẫn được đánh giá riêng ở Chương 4.'
first_head=heads[0]._p
newp(first_head,'Quy ước ký hiệu UML',bold=True,keep=True)
newp(first_head,'Tác nhân được vẽ bằng hình người; hình elip biểu diễn ca sử dụng. Khung chữ nhật chỉ phạm vi hệ thống ICD phục vụ Mobile. Đường liền biểu diễn tác nhân tham gia; «include» chỉ hành vi cần thực hiện trong ca sử dụng; «extend» chỉ hành vi bổ sung có điều kiện. Các đường quan hệ không biểu diễn thứ tự thời gian. Phiên đăng nhập hợp lệ là điều kiện trước của chức năng nghiệp vụ, không yêu cầu người dùng đăng nhập lại sau mỗi thao tác.')

for i,(heading,case,old) in enumerate(zip(heads,CASES,old_sections),1):
    anchor=heading._p.getnext()
    while anchor is not None:
        from docx.text.paragraph import Paragraph
        if anchor.tag==qn('w:p') and Paragraph(anchor,doc._body).style.name in ('Heading 2','Heading 3'):break
        following=anchor.getnext();anchor.getparent().remove(anchor);anchor=following
    assert anchor is not None
    heading.text=f'3.3.{i} Chức năng {case["title"].lower()}'
    heading.paragraph_format.keep_with_next=True
    path=diagram(i,case)
    picture(anchor,path,f'Sơ đồ Use Case UC{i:02d}: {case["title"]}. Tác nhân: {case["actor"]}.')
    caption(anchor,f'Hình 3.{i+2}: Use Case {case["title"].lower()} (UC{i:02d})')
    if case['optional']:
        newp(anchor,'Chú thích quan hệ mở rộng: '+ '; '.join(label+' khi '+condition[0].lower()+condition[1:] for label,condition in case['optional'])+'.')
    label(anchor,f'Mô tả Use Case: {case["title"]} (UC{i:02d})')
    actor_text=re.sub(r'^UC\d+;\s*','',old['actor'].split('Màn:')[0].strip()).rstrip('.')
    newp(anchor,'Tác nhân (Actor): '+actor_text)
    newp(anchor,'Mục đích: '+case['purpose'])
    label(anchor,'1. Điều kiện tiên quyết')
    newp(anchor,old['pre'])
    label(anchor,'2. Luồng sự kiện chính')
    for j,step in enumerate(case['steps'],1):
        p=newp(anchor)
        p.add_run(f'Bước {j}: ').bold=True;p.add_run(step)
    label(anchor,'3. Các luồng mở rộng')
    for j,(title,body) in enumerate(case['alternatives'],1):
        label(anchor,f'Luồng E{j} - {title}:')
        newp(anchor,body)
    label(anchor,'4. Điều kiện kết thúc')
    newp(anchor,old['post'])
    newp(anchor,'Nếu backend chưa xác nhận thành công, ứng dụng phải thể hiện lỗi hoặc kết quả cần đối chiếu; không tự ghi nhận đã hoàn thành thao tác nghiệp vụ.' if i not in [13,14,15,16,17] else 'Thông tin hiển thị và hành động được phép phải phản ánh dữ liệu phiên cùng response của resource liên quan.')
    label(anchor,'Đối chiếu triển khai')
    screen=old['actor'].split('Màn:',1)[1].strip() if 'Màn:' in old['actor'] else ''
    newp(anchor,'Màn hình: '+screen).alignment=WD_ALIGN_PARAGRAPH.LEFT
    newp(anchor,'API liên quan: '+old['api']).alignment=WD_ALIGN_PARAGRAPH.LEFT

# Keep the next analytical section separate after the final use case.
next_section=next(p for p in doc.paragraphs if p.style.name=='Heading 2' and p.text.startswith('3.4 '))
next_section.paragraph_format.page_break_before=True
doc.save(OUT)
(FIG/'use-case-index.json').write_text(json.dumps([{'id':f'UC{i:02d}','title':c['title'],'actor':c['actor'],'figure':f'Hình 3.{i+2}','png':f'UC{i:02d}.png','svg':f'UC{i:02d}.svg'} for i,c in enumerate(CASES,1)],ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'new_use_cases':len(CASES),'figures':len(list(FIG.glob('*.png'))),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'inline_shapes':len(doc.inline_shapes)},ensure_ascii=False))
