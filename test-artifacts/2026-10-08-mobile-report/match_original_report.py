from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import hashlib
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(r'D:\Project\Đồ Án 4 +Mobile')
WORK = ROOT / 'icd-management/test-artifacts/2026-10-08-mobile-report'
REFERENCE = WORK / 'DoAn3-mobile-reference.docx'
SOURCE = ROOT / '10123264 - Dao Minh Quang - Bao cao Mobile ICD - Use Case.docx'
OUTPUT = ROOT / '10123264 - Dao Minh Quang - Bao cao Mobile ICD - Theo mau goc.docx'
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest().upper() == 'FA00C74F131462DDC97D675DE8087F82F84EF4F54ABB2EBEDF11B4164FB148EA'

ref = Document(REFERENCE)
doc = Document(SOURCE)


def set_text(p: Paragraph, value: str, *, bold_prefix: str | None = None) -> None:
    p.clear()
    if bold_prefix and value.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(value[len(bold_prefix):])
    else:
        p.add_run(value)


def insert_after(anchor: Paragraph, value: str, *, style='Nội dung', bold_prefix: str | None = None, bullet=False) -> Paragraph:
    element = OxmlElement('w:p')
    anchor._p.addnext(element)
    p = Paragraph(element, anchor._parent)
    p.style = doc.styles[style]
    if bullet:
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        value = '•  ' + value
    set_text(p, value, bold_prefix=bold_prefix)
    return p


def insert_before(anchor: Paragraph, value: str, *, style='Nội dung') -> Paragraph:
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    p = Paragraph(element, anchor._parent)
    p.style = doc.styles[style]
    set_text(p, value)
    return p


def find_prefix(prefix: str) -> Paragraph:
    return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix) and p.style.name.startswith('Heading'))


def add_bottom_border(p: Paragraph) -> None:
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn('w:pBdr'))
    if pbdr is None:
        pbdr = OxmlElement('w:pBdr')
        ppr.append(pbdr)
    bottom = pbdr.find(qn('w:bottom'))
    if bottom is None:
        bottom = OxmlElement('w:bottom')
        pbdr.append(bottom)
    for key, value in [('val','single'),('sz','4'),('space','7'),('color','999999')]:
        bottom.set(qn('w:'+key), value)


# Restore the recurring page furniture and paragraph rhythm of the supplied thesis.
for name in ('Normal', 'Nội dung', 'Heading 1', 'Heading 2', 'Heading 3', 'Tiểu mục', 'Mục lớn', 'HinhVe', 'Bang'):
    dst = doc.styles[name]
    src = ref.styles[name]
    dst.font.name = src.font.name
    if src.font.size: dst.font.size = src.font.size
    dst.font.bold = src.font.bold
    dst.font.italic = src.font.italic
    for prop in ('alignment','space_before','space_after','line_spacing','first_line_indent','left_indent','right_indent'):
        setattr(dst.paragraph_format, prop, getattr(src.paragraph_format, prop))

header = doc.sections[1].header.paragraphs[0]
reference_header = ref.sections[1].header.paragraphs[0]
src_ppr = reference_header._p.pPr
if src_ppr is not None:
    dst_ppr = header._p.get_or_add_pPr()
    old = dst_ppr.find(qn('w:pBdr'))
    if old is not None: dst_ppr.remove(old)
    original = src_ppr.find(qn('w:pBdr'))
    if original is not None: dst_ppr.append(deepcopy(original))
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_bottom_border(header)
header_border = header._p.get_or_add_pPr().find(qn('w:pBdr')).find(qn('w:bottom'))
header_border.set(qn('w:color'),'000000')
header_border.set(qn('w:space'),'2')

footer = doc.sections[1].footer.paragraphs[0]
footer.clear()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text = ' PAGE '
separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
display = OxmlElement('w:t'); display.text = '1'
end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
for el in (begin,instr,separate,display,end): run._r.append(el)

# Chapter 1 adopts the same academic sequence as the source while keeping ICD facts.
renames = {
    '1.1 Bối cảnh vận hành cảng cạn':'1.1. Lý do chọn đề tài',
    '1.2 Mục tiêu và kết quả cần có':'1.2. Mục tiêu của đề tài',
    '1.3 Đối tượng và phạm vi':'1.3. Giới hạn và phạm vi của đề tài',
    '1.4 Phương pháp thực hiện và nguồn bằng chứng':'1.5. Phương pháp tiếp cận',
    '3.2 Yêu cầu chức năng':'3.2. Đặc tả yêu cầu phần mềm',
    '3.7 Khung biểu đồ cần giữ khi cập nhật luận văn':'3.7. Tổng hợp các mô hình thiết kế',
    '4.2 Triển khai giao diện theo từng màn hình':'4.2. Triển khai các chức năng trên ứng dụng Mobile',
    '4.4 Kiểm thử đã có bằng chứng':'4.4. Kiểm thử ứng dụng',
    '4.7 Giới hạn hiện tại và hướng hoàn thiện':'4.7. Đánh giá kết quả triển khai',
    'PHỤ LỤC D  ĐƯỜNG DẪN MINH CHỨNG VÀ CÁCH CẬP NHẬT':'PHỤ LỤC D  NGUỒN MINH CHỨNG',
}
for p in doc.paragraphs:
    key = p.text.strip()
    if key in renames: set_text(p, renames[key])

method = find_prefix('1.5. Phương pháp tiếp cận')
content = [
    ('Khảo sát nghiệp vụ: ', 'đối chiếu luồng Manifest, chuyến xe, cổng, bãi, giám định và Gate-out với đặc tả ICD v1.7 và README của hệ thống.'),
    ('Phân tích và thiết kế: ', 'mô hình hóa tác nhân, ca sử dụng, dữ liệu và trạng thái; tách rõ hành động của ứng dụng với quyết định của backend.'),
    ('Xây dựng ứng dụng: ', 'triển khai các màn hiện trường bằng React Native và Expo, kết nối API theo quyền của tài khoản đang đăng nhập.'),
    ('Kiểm thử và đối chiếu: ', 'dùng hồ sơ QA để kiểm giao diện, phản hồi API và dữ liệu đọc lại; những phép thử chưa thực hiện trên thiết bị native được ghi riêng.'),
]
heading = insert_before(method, '1.4. Nội dung thực hiện', style='Heading 2')
cursor = heading
for lead, body in content:
    cursor = insert_after(cursor, lead + body, bold_prefix=lead, bullet=True)

# The reference devotes its first theory section to the development process.
theory = find_prefix('2.1 Kiến trúc ứng dụng di động kết nối API')
theory_parts = [
    ('Khảo sát và xác định yêu cầu: ', 'từ quy trình ICD, xác định các tác vụ cần làm ngay tại cổng hoặc bãi và những tác vụ vẫn thuộc Web.'),
    ('Phân tích và thiết kế: ', 'đặc tả ca sử dụng, permission, điều kiện trạng thái và nguồn dữ liệu của từng màn hình trước khi nối API.'),
    ('Lập trình và tích hợp: ', 'phát triển các feature Mobile theo nhóm nghiệp vụ, đọc dữ liệu từ backend và chỉ hiển thị thành công sau phản hồi xác nhận.'),
    ('Kiểm thử: ', 'đối chiếu giao diện với ảnh audit, kiểm tra API và dữ liệu lưu; tách kết quả Expo Web khỏi những mục cần thử trên Android hoặc iOS.'),
]
new_theory = insert_before(theory, '2.1. Quy trình phát triển ứng dụng', style='Heading 2')
cursor = insert_after(new_theory, 'Quy trình thực hiện bám theo các mốc có thể kiểm chứng trong mã nguồn và tài liệu nghiệp vụ. Mỗi màn hình được xem cùng API, quyền và trạng thái dữ liệu mà nó sử dụng.')
for lead, body in theory_parts:
    cursor = insert_after(cursor, lead + body, bold_prefix=lead, bullet=True)
for old,new in [('2.1 Kiến trúc','2.2. Kiến trúc'),('2.2 Dữ liệu','2.3. Dữ liệu'),('2.3 Thiết kế','2.4. Thiết kế'),('2.4 Các tiêu chí','2.5. Các tiêu chí')]:
    p=find_prefix(old); set_text(p,p.text.replace(old,new,1))

# Match the source's Use Case write-up: role/purpose bullets, a separator,
# numbered main flow, and indented extension cases.
for p in list(doc.paragraphs):
    text = p.text.strip()
    if text.startswith('Mô tả Use Case:'):
        p.runs[0].bold = True
        p.paragraph_format.space_before = Pt(5)
    elif text.startswith('Tác nhân (Actor):'):
        set_text(p, '•  ' + text, bold_prefix='•  Tác nhân (Actor):')
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.3)
    elif text.startswith('Mục đích:'):
        set_text(p, '•  ' + text, bold_prefix='•  Mục đích:')
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(12)
        add_bottom_border(p)
    elif re.fullmatch(r'[1-4]\. (Điều kiện tiên quyết|Luồng sự kiện chính|Các luồng mở rộng|Điều kiện kết thúc)',text):
        for r in p.runs: r.bold=True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next=True
    elif re.match(r'^Bước \d+:',text):
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        set_text(p,re.sub(r'^Bước (\d+):',r'\1.',text))
    elif re.match(r'^Luồng E\d+ - ',text):
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        set_text(p,'•  '+text.replace(' - ',' – ',1))
        for r in p.runs: r.bold=True
    elif text in ('Đối chiếu triển khai','Quy ước ký hiệu UML'):
        for r in p.runs: r.bold=True
    elif text.startswith(('Màn hình:','API liên quan:')):
        p.paragraph_format.left_indent = Cm(0.4)
        p.runs[0].italic=True

# Restore the reference's blue visual treatment for tabular evidence.
for table in doc.tables:
    if not table.rows: continue
    for cell in table.rows[0].cells:
        tcpr=cell._tc.get_or_add_tcPr()
        shd=tcpr.find(qn('w:shd'))
        if shd is None: shd=OxmlElement('w:shd');tcpr.append(shd)
        shd.set(qn('w:fill'),'4472C4')
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb=RGBColor(255,255,255)
                run.bold=True

# Replace the repeated generic audit caption note with figure-specific explanation.
# Every statement stays within the visible screenshot or the existing feature text.
FIGURE_NOTES = {
 2: ('Màn đăng nhập là điểm vào chung cho các vai trò nhân viên.', [
     ('Trường nhập liệu: ', 'email và mật khẩu được nhập trước khi gửi yêu cầu xác thực.'),
     ('Phản hồi: ', 'lỗi đăng nhập hoặc mất kết nối xuất hiện tại màn này; chỉ phiên được backend xác nhận mới mở các tab nghiệp vụ.')]),
 3: ('Màn Tài khoản tập hợp thông tin phiên và các quyền đang có.', [
     ('Thông tin người dùng: ', 'danh tính, phạm vi ICD và quyền được trình bày để nhân viên biết mình có thể thao tác ở đâu.'),
     ('Thiết lập cá nhân: ', 'lựa chọn sáng hoặc tối được lưu; lối đăng xuất vẫn có khi tài khoản không được mở tab tác nghiệp.')]),
 4: ('Biểu mẫu Gate-in đặt thông tin container cạnh các dữ kiện cần đối chiếu tại cổng.', [
     ('Bối cảnh: ', 'nhân viên chọn chuyến xe, kiểm tra container và trạng thái đến của xe trước khi xác nhận.'),
     ('Thông tin thực tế: ', 'seal, trọng lượng và tình trạng được nhập theo hàng đang tiếp nhận; dữ liệu lỗi cần được sửa trước khi gửi lệnh.')]),
 5: ('Hộp thoại xem lại xuất hiện trước khi gửi lệnh tiếp nhận.', [
     ('Nội dung đối chiếu: ', 'số container, chuyến xe và dữ liệu hiện trường được nhắc lại để tránh thao tác nhầm.'),
     ('Quyết định: ', 'nhân viên có thể quay lại biểu mẫu; nút xác nhận chỉ hoàn tất khi backend ghi nhận giao dịch.')]),
 6: ('Biên nhận biểu diễn kết quả mà API trả về sau Gate-in.', [
     ('Kết quả: ', 'mã hồ sơ và trạng thái sau tiếp nhận giúp nhân viên đối chiếu với container vừa xử lý.'),
     ('Thao tác tiếp: ', 'mở hồ sơ hoặc bắt đầu container khác chỉ sau khi phản hồi thành công đã được nhận.')]),
 7: ('Màn quét phiếu là cửa vào của luồng Gate-out.', [
     ('Quét và nhập tay: ', 'Expo Web trong ảnh dùng ô nhập mã; kiểm thử camera và quyền truy cập cần thực hiện riêng trên native.'),
     ('Kiểm tra phiếu: ', 'mã được gửi tới API scan để lấy trạng thái hiện tại, thay vì suy ra tính hợp lệ từ nội dung QR tại thiết bị.')]),
 8: ('Trước khi xe ra, nhân viên xem lại phiếu, container và điều kiện còn hiệu lực.', [
     ('Trạng thái: ', 'hệ thống trình bày blocker hoặc cảnh báo readiness để nhân viên xử lý đúng hồ sơ.'),
     ('Xác nhận: ', 'lệnh Gate-out được gửi một lần; trạng thái hiển thị cuối cùng dựa trên phản hồi backend.')]),
 9: ('Màn kết quả cho biết giao dịch ra cổng đã được ghi.', [
     ('Phiếu: ', 'Gate Pass chuyển sang USED sau khi hoàn tất; phiếu không còn được dùng cho lượt quét tiếp theo.'),
     ('Đối chiếu: ', 'container visit và vị trí bãi được tải lại để kiểm trạng thái EXITED.')]),
 10: ('Sơ đồ bãi đưa trạng thái ô và Legend vào cùng một vùng nhìn.', [
     ('Màu trạng thái: ', 'màu của ô theo cùng quy ước với Legend, tránh hiểu một ô đang chiếm là còn trống.'),
     ('Thao tác: ', 'người dùng chọn ô để xem điều kiện; tính năng đề xuất vị trí tối ưu vẫn hiển thị là đang phát triển.')]),
 11: ('Kết quả kiểm tra ô bãi giải thích vì sao vị trí không thể sử dụng.', [
     ('Dữ liệu kiểm tra: ', 'blocker từ backend được hiển thị bên cạnh ô đã chọn.'),
     ('Hướng xử lý: ', 'nhân viên quay lại và chọn vị trí khác; ứng dụng không gửi lệnh gán cho ô bị từ chối.')]),
 12: ('Ảnh cho thấy kết quả gán vị trí bằng phương thức MANUAL.', [
     ('Vị trí đã ghi: ', 'mã slot được hiển thị lại từ dữ liệu sau giao dịch.'),
     ('Giới hạn: ', 'bản này chưa triển khai thuật toán xếp bãi tự động; điều kiện slot tiếp tục do backend kiểm tra.')]),
 13: ('Danh sách tác nghiệp gom movement, booking và các hồ sơ theo dõi trong ca.', [
     ('Bộ lọc: ', 'nhân viên thu hẹp danh sách theo nhóm nghiệp vụ và trạng thái.'),
     ('Điều hướng: ', 'mỗi hàng mở chi tiết đúng record; dữ liệu rỗng và lỗi tải cần được phân biệt.')]),
 14: ('Chi tiết movement hiển thị vị trí nguồn, vị trí đích và tiến trình xử lý.', [
     ('Trạng thái: ', 'PENDING và IN_PROGRESS cho biết lệnh đang ở bước nào.'),
     ('Hành động: ', 'bắt đầu hoặc hoàn tất chỉ xuất hiện khi quyền và trạng thái cho phép; phản hồi server quyết định kết quả.')]),
 15: ('Màn booking hiển thị kết quả tác nghiệp đã hoàn tất.', [
     ('Hồ sơ: ', 'loại tác nghiệp, lịch và container được giữ cùng một ngữ cảnh.'),
     ('Kết thúc: ', 'kết quả thực tế được đối chiếu với trạng thái booking sau khi API xác nhận.')]),
 16: ('Biểu mẫu giám định hỗ trợ chọn loại kiểm tra và ghi nhận hiện trạng.', [
     ('Loại kiểm tra: ', 'hư hỏng, seal hoặc tình trạng được chọn theo nội dung công việc.'),
     ('Dữ liệu ghi nhận: ', 'ghi chú và mức độ chỉ xuất hiện khi phù hợp với loại kiểm tra; biểu mẫu giữ dữ liệu nếu yêu cầu lỗi.')]),
 17: ('Biên bản sau giám định hiển thị kết quả đã được lưu.', [
     ('Kết quả: ', 'trạng thái và ghi chú giúp nhân viên xác nhận hồ sơ mình vừa xử lý.'),
     ('Ảnh hưởng nghiệp vụ: ', 'nếu giám định tạo blocker, điều kiện cấp phiếu hoặc ra cổng phải được đọc lại từ backend.')]),
 18: ('Màn Tra cứu bắt đầu từ số container và danh sách hồ sơ phù hợp.', [
     ('Tìm kiếm: ', 'người dùng nhập số container để nhận kết quả từ API.'),
     ('Chọn hồ sơ: ', 'khi có nhiều visit, nhân viên mở đúng hồ sơ trước khi đọc vị trí và sự kiện.')]),
 19: ('Hồ sơ chi tiết gom các thông tin cần kiểm ở hiện trường.', [
     ('Bối cảnh: ', 'sự kiện, vị trí bãi, holds và điều kiện ra cổng được trình bày theo quyền.'),
     ('Trạng thái: ', 'kết quả mới nhất được lấy từ backend; dữ liệu đọc từ cache cần có thời điểm cập nhật.')]),
 20: ('Việc ca cho thấy những công việc đang chờ cùng mức ưu tiên.', [
     ('Danh sách: ', 'loại việc, hạn xử lý và trạng thái giúp nhân viên chọn việc cần làm.'),
     ('Đi tới tác vụ: ', 'mục được chọn chuyển tới đúng container hoặc record mà backend trả về.')]),
 21: ('Màn Điều phối phục vụ theo dõi tổng hợp khi tài khoản có quyền.', [
     ('Phạm vi nhìn: ', 'các nhóm trạng thái container và công việc được dùng để nhận biết tình hình bãi.'),
     ('Giới hạn thao tác: ', 'màn tổng hợp không thay quyền của từng lệnh cổng hoặc bãi.')]),
 22: ('Thông báo xe đến đưa nhân viên tới hồ sơ liên quan.', [
     ('Nội dung: ', 'popup nêu sự kiện và container để người nhận biết lý do cần mở hồ sơ.'),
     ('Tương tác: ', 'có thể đóng thông báo hoặc mở chi tiết; trạng thái đã đọc chỉ đổi khi API xác nhận.')]),
 23: ('Chế độ tối áp dụng trên những màn chính của ứng dụng.', [
     ('Hiển thị: ', 'nền, chữ và màu trạng thái được đổi theo theme nhưng vẫn giữ cùng ý nghĩa nghiệp vụ.'),
     ('Phạm vi ảnh: ', 'đây là ảnh Expo Web của bộ audit; kiểm tra tương phản trên Android native là bước nghiệm thu riêng.')]),
}

for p in list(doc.paragraphs):
    match=re.match(r'^Hình 4\.(\d+)\.',p.text.strip())
    if not match or p.style.name!='HinhVe': continue
    number=int(match.group(1))
    if number not in FIGURE_NOTES: continue
    next_el=p._p.getnext()
    if next_el is not None:
        next_p=Paragraph(next_el,p._parent)
        if next_p.text.strip().startswith('Ảnh audit Expo Web theo bộ ngày'):
            next_el.getparent().remove(next_el)
    intro,bullets=FIGURE_NOTES[number]
    cursor=insert_after(p,intro)
    for lead,body in bullets:
        cursor=insert_after(cursor,lead+body,bold_prefix=lead,bullet=True)

for p in list(doc.paragraphs):
    if p.text.strip().startswith(('Ảnh audit Expo Web theo bộ ngày','Ảnh đầu Expo Web','Ảnh smoke Expo Web','Ảnh dataset test')):
        p._element.getparent().remove(p._element)

# A single source note is more useful than repeating the same note after 22 figures.
interface=find_prefix('4.2. Triển khai các chức năng')
insert_after(interface,'Các ảnh trong mục 4.2 được lấy từ bộ audit Expo Web ngày 02/10/2026. Chúng minh họa bố cục và trạng thái tại thời điểm chụp; camera, TalkBack và vòng đời ứng dụng vẫn cần kiểm chứng trên Android native.')

# Remove draft-language from completed analytical sections.
for p in doc.paragraphs:
    if p.text.strip().startswith('Các hình 3.1–3.9 đã có nội dung'):
        set_text(p,'Các Hình 3.1–3.26 liên kết nghiệp vụ hiện trường với quyền, trạng thái và dữ liệu. Use Case xác định người tham gia; sơ đồ luồng và trạng thái giải thích thứ tự xử lý; các bảng ánh xạ ở cuối chương giúp truy nguyên từ giao diện tới API.')
    if p.text.strip().startswith('Khi hoàn thiện bản nộp, cần chạy lại ma trận native'):
        set_text(p,'Phạm vi kiểm chứng hiện tại dừng ở các lượt audit đã có bằng chứng. Nghiệm thu Android cần bổ sung camera, TalkBack, cỡ chữ lớn, Android Back và reconnect; iOS được đánh giá trong một lượt riêng. Những kết quả đó sẽ là cơ sở để cập nhật đánh giá cuối cùng của ứng dụng.')
    if p.text.strip().startswith('Cách đọc ảnh. Ảnh 4.4–4.23'):
        set_text(p,'Các Hình 4.4–4.23 được chọn từ bộ audit vận hành và phản ánh trạng thái tại thời điểm chụp. Ảnh giao diện chỉ xác nhận màn đã hiển thị; kết quả giao dịch được đối chiếu bằng phản hồi API và dữ liệu đọc lại trong hồ sơ QA.')
    if p.text.strip().startswith('Sau khi thay ảnh: mở hình gốc'):
        set_text(p,'Bộ ảnh trong báo cáo giữ nhãn môi trường Expo Web và ngày chụp. Những ảnh này đại diện cho luồng giao diện, còn bằng chứng camera, TalkBack và quyền thiết bị trên Android native chưa có trong tập tư liệu được sử dụng.')
    if p.text.strip()=='Bảng 3.5. Đặc tả chi tiết để vẽ và đối chiếu từng biểu đồ':
        set_text(p,'Bảng 3.5. Liên hệ mô hình thiết kế với nghiệp vụ Mobile')

# The source ends with clearly labelled achieved results, limits and development work.
conclusion=find_prefix('KẾT LUẬN')
reference_section=find_prefix('TÀI LIỆU THAM KHẢO')
elements=[]
el=conclusion._p.getnext()
while el is not None and el is not reference_section._p:
    elements.append(el)
    el=el.getnext()
for el in elements: el.getparent().remove(el)
cursor=insert_after(conclusion,'1. Kết quả đạt được',style='Heading 2')
cursor=insert_after(cursor,'Đồ án đã hình thành ứng dụng Mobile ICD cho các công việc tại Cổng, Bãi và Giám định. Nhân viên còn có màn Tra cứu, Việc ca, Thông báo, Tài khoản và Điều phối theo permission của phiên đăng nhập.')
for lead,body in [
 ('Giao diện hiện trường: ','các màn quét, biểu mẫu, kết quả giao dịch và màn chi tiết được đặt theo luồng công việc của nhân viên.'),
 ('Kết nối nghiệp vụ: ','ứng dụng dùng API NestJS; backend quyết định quyền, readiness và việc ghi giao dịch.'),
 ('Bằng chứng kiểm thử: ','bộ audit Expo Web đã đi qua Gate-in, xếp bãi, movement, inspection, booking và Gate-out trên hồ sơ QA; dữ liệu sau restart được kiểm lại trong báo cáo nguồn.')]:
    cursor=insert_after(cursor,lead+body,bold_prefix=lead,bullet=True)
cursor=insert_after(cursor,'2. Hạn chế còn tồn tại',style='Heading 2')
cursor=insert_after(cursor,'Ảnh Expo Web chứng minh giao diện và một số giao dịch trong phạm vi audit, nhưng không thay được phép thử native. Camera và quyền truy cập, TalkBack, cỡ chữ lớn, Android Back, chuyển nền và mất mạng vẫn cần đo trên Android. Các nghiệp vụ billing, quản trị và Partner thuộc Web hoặc API tương ứng, không được trình bày như chức năng Mobile đã hoàn tất.')
cursor=insert_after(cursor,'3. Hướng hoàn thiện',style='Heading 2')
cursor=insert_after(cursor,'Chạy lại ma trận native trên bản build cuối và lưu ảnh, cấu hình thiết bị, vai trò, theme cùng kết quả API cho mỗi ca. Sau nghiệm thu Android, kiểm tra riêng iOS; khi mã nguồn hoặc dữ liệu thay đổi, cập nhật ảnh và các con số inventory từ lần chụp mới.')

doc.save(OUTPUT)
print('saved',OUTPUT.name,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
