from collections import Counter
from pathlib import Path
from docx import Document

root = Path(r'D:\Project\Đồ Án 4 +Mobile')
reference = root / 'icd-management/test-artifacts/2026-10-08-mobile-report/DoAn3-mobile-reference.docx'
current = root / '10123264 - Dao Minh Quang - Bao cao Mobile ICD - Use Case.docx'

for label, path in [('REFERENCE', reference), ('CURRENT', current)]:
    doc = Document(path)
    out = [f'# {label}: {path}', f'paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}']
    out.append('STYLES ' + repr(Counter(p.style.name for p in doc.paragraphs).most_common(25)))
    out.append('## HEADINGS')
    for i, p in enumerate(doc.paragraphs):
        if p.style.name in ('Heading 1', 'Heading 2', 'Heading 3', 'Tiểu mục', 'Mục lớn', 'MucLon', 'MucNho', 'Ý lớn'):
            out.append(f'{i:04d} | {p.style.name} | {p.text.strip()}')
    out.append('## FIRST BODY SAMPLES')
    selected = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if p.style.name in ('Nội dung', 'Normal', 'No Spacing') and len(t) > 65 and i > 200:
            selected.append((i,p.style.name,t))
        if len(selected) >= 30:
            break
    for i, style, t in selected:
        out.append(f'{i:04d} | {style} | {t[:800]}')
    if label == 'CURRENT':
        out.append('## USE CASE ONE AND CHAPTER FOUR')
        for lo, hi in [(285, 315), (790, 820), (880, 925)]:
            for i in range(lo, hi):
                p = doc.paragraphs[i]
                out.append(f'{i:04d} | {p.style.name} | {p.text.strip()[:800]}')
    target = path.parent / f'{label.lower()}-inspection.txt'
    target.write_text('\n'.join(out), encoding='utf-8')
    print(label, 'inspection written')
