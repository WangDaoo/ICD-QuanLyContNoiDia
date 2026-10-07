from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\Project\Đồ Án 4 +Mobile")
RUN = ROOT / "icd-management" / "test-artifacts" / "2026-10-07-week6-report"
TEMPLATE = RUN / "Bao_cao_tuan_6-template.docx"
OUTPUT = ROOT / "Bao_cao_tuan_6_ICD.docx"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_text(doc, text, style="normal", align=None, bold=False, size=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return p


def add_bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style="normal")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.add_run("• ").bold = True
        p.add_run(value)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        if widths:
            cell.width = Cm(widths[idx])
        set_cell_shading(cell, "D9EAD3")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            if widths:
                cell.width = Cm(widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            r = p.add_run(str(value))
            r.font.size = Pt(9)
    doc.add_paragraph()
    return table


def main():
    doc = Document(TEMPLATE)
    clear_body(doc)
    section = doc.sections[0]
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.0)

    # Preserve the visual role of the source cover while replacing the old subject.
    for line, size, bold, before in [
        ("BỘ GIÁO DỤC VÀ ĐÀO TẠO", 14, True, 0),
        ("TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT HƯNG YÊN", 14, True, 2),
        ("", 12, False, 10),
        ("ĐỒ ÁN 4", 16, True, 8),
        ("BÁO CÁO TIẾN ĐỘ TUẦN 6", 16, True, 6),
        ("", 12, False, 20),
        ("XÂY DỰNG HỆ THỐNG QUẢN LÝ ICD", 18, True, 8),
        ("", 12, False, 16),
        ("KHOA: CÔNG NGHỆ THÔNG TIN", 13, True, 3),
        ("NGÀNH: KỸ THUẬT PHẦN MỀM", 13, True, 3),
        ("CHUYÊN NGÀNH: PHÁT TRIỂN PHẦN MỀM ỨNG DỤNG", 13, True, 3),
        ("", 12, False, 16),
        ("SINH VIÊN: TRẦN MAI LAN", 13, False, 3),
        ("MÃ SV: 10123196", 13, False, 3),
        ("MÃ LỚP: 12523W.4", 13, False, 3),
        ("HƯỚNG DẪN: TS. HOÀNG QUỐC VIỆT", 13, False, 3),
    ]:
        add_text(doc, line, style="Bìa", align=WD_ALIGN_PARAGRAPH.CENTER, bold=bold, size=size, before=before)
    add_text(doc, "HƯNG YÊN - 2026", style="Bìa", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, before=28)
    # The supplied "Tiêu đề" style already starts on a new page.  Adding a
    # manual break here would therefore create an empty page between the cover
    # and the report body.

    add_text(doc, "BÁO CÁO TIẾN ĐỘ TUẦN 6", style="Tiêu đề", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, after=8)
    add_text(doc, "Đề tài Xây dựng hệ thống quản lý ICD", style="normal", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=12)
    add_text(doc, "Báo cáo tổng hợp kết quả triển khai, kiểm thử và kế hoạch công việc tiếp theo. Nội dung được đối chiếu với README hệ thống ICD, đặc tả nghiệp vụ, đặc tả Web Mobile và mã nguồn trong monorepo.", after=8)
    add_table(doc, ["Nội dung", "Thông tin"], [
        ["Tên hệ thống", "ICD Management System"],
        ["Phạm vi", "Web React, Mobile Expo React Native, API NestJS và MySQL"],
        ["Mốc báo cáo", "Tuần 6"],
        ["Nguồn đối chiếu", "README hệ thống, tài liệu nghiệp vụ ICD, mã nguồn và kết quả kiểm thử Local"],
    ], [4.2, 12.2])

    add_text(doc, "MỤC TIÊU VÀ PHẠM VI TUẦN 6", style="Heading 1", bold=True, after=6)
    add_text(doc, "1.1. Mục tiêu", style="Mục lớn", bold=True)
    add_text(doc, "Tuần 6 tập trung hoàn thiện phần mô tả hệ thống ICD theo chuỗi vận hành container, đồng thời chuẩn hóa tài liệu kỹ thuật để có thể theo dõi thống nhất giữa nghiệp vụ, giao diện Web, ứng dụng Mobile, API và cơ sở dữ liệu. Báo cáo cũng ghi nhận tình trạng kiểm thử thủ công trên môi trường Local nhằm kiểm tra tính liên kết của các bước nghiệp vụ quan trọng.")
    add_bullets(doc, [
        "Tổng hợp kiến trúc monorepo và các thành phần triển khai thực tế.",
        "Đối chiếu chuỗi nghiệp vụ từ Manifest đến Gate out với giao diện và API.",
        "Lập danh mục chức năng Web, Mobile, quyền truy cập, API JSON và dữ liệu MySQL.",
        "Thực hiện kiểm thử E2E ban đầu trên dữ liệu tổng hợp có gắn nhãn E2E để không lẫn với dữ liệu vận hành thông thường.",
    ])
    add_text(doc, "1.2. Phạm vi", style="Mục lớn", bold=True)
    add_text(doc, "Phạm vi báo cáo gồm Web frontend tại apps/web, Mobile app tại apps/mobile, API tại apps/api và các package dùng chung. Báo cáo mô tả đúng những phần đã được xác nhận trong mã nguồn và môi trường Local. Các tính năng chưa có giao diện thao tác hoặc chưa hoàn tất kiểm thử được nêu rõ ở phần giới hạn.")

    add_text(doc, "TỔNG QUAN HỆ THỐNG ICD", style="Heading 1", bold=True, after=6)
    add_text(doc, "2.1. Bài toán nghiệp vụ", style="Mục lớn", bold=True)
    add_text(doc, "ICD Management System hỗ trợ vận hành cảng cạn và kho bãi container. Hệ thống quản lý thông tin lược khai, vận đơn, container, lệnh vận chuyển, chuyến xe, tiếp nhận cổng, vị trí bãi, dịch vụ, hóa đơn, phiếu ra cổng và bàn giao cho đối tác logistics. Mỗi thay đổi trạng thái quan trọng được backend kiểm tra trước khi ghi nhận để tránh giao diện tự xác nhận một giao dịch chưa được hệ thống nghiệp vụ chấp nhận.")
    add_text(doc, "2.2. Chuỗi nghiệp vụ chính", style="Mục lớn", bold=True)
    add_table(doc, ["Bước", "Đầu vào", "Kết quả nghiệp vụ"], [
        ["1. Khai báo hàng", "Manifest, MBL, HBL, container", "Container Visit ở trạng thái PENDING"],
        ["2. Cho phép vào ICD", "Movement Order được authorize", "Container đủ điều kiện gắn Truck Visit"],
        ["3. Tiếp nhận cổng", "Xe ARRIVED, seal và trọng lượng thực tế", "Container Visit chuyển IN_YARD"],
        ["4. Vận hành bãi", "Xếp slot, đảo chuyển, giám định, booking", "Vị trí và tác nghiệp được lưu lịch sử"],
        ["5. Tài chính", "Đơn dịch vụ, hóa đơn, thanh toán", "Readiness xác định các blocker tài chính"],
        ["6. Ra cổng", "Readiness đạt, Gate Pass hợp lệ, quét QR", "Visit EXITED, Gate Pass USED, vị trí bãi được đóng"],
        ["7. Bàn giao", "Handover và xác nhận Partner hoặc kho", "Theo dõi giao nhận ngoài ICD mà không thay đổi core state"],
    ], [3.1, 5.6, 7.7])

    add_text(doc, "KIẾN TRÚC VÀ CÔNG NGHỆ", style="Heading 1", bold=True, after=6)
    add_text(doc, "3.1. Kiến trúc monorepo", style="Mục lớn", bold=True)
    add_text(doc, "Dự án tổ chức theo monorepo. Các frontend không truy cập trực tiếp MySQL mà gửi request tới API. API thực hiện xác thực, kiểm tra quyền, kiểm tra trạng thái nghiệp vụ và lưu dữ liệu bằng Prisma trên MySQL. Thiết kế này giúp giữ backend là nguồn quyết định đối với readiness, thanh toán, hold, vị trí bãi và Gate out.")
    add_table(doc, ["Thành phần", "Công nghệ", "Vai trò"], [
        ["Web", "React 19, Vite, TypeScript, Tailwind", "Điều hành và quản trị trên trình duyệt"],
        ["Mobile", "Expo 57, React Native 0.86, React Navigation", "Tác nghiệp hiện trường: cổng, bãi, giám định, tra cứu, việc ca"],
        ["Backend", "NestJS 12, TypeScript, Prisma", "API, xác thực JWT, RBAC, quy tắc nghiệp vụ, audit và tích hợp"],
        ["Database", "MySQL 8", "Lưu nghiệp vụ, quyền, dữ liệu vận hành, EDI và notification"],
        ["Package dùng chung", "API client, shared types, validation", "Giảm sai khác contract giữa client và server"],
    ], [3.2, 5.6, 7.6])
    add_text(doc, "3.2. Cấu trúc xử lý request", style="Mục lớn", bold=True)
    add_text(doc, "Mọi API được phục vụ dưới tiền tố /api. Request đi qua JWT guard, Permission guard, validation pipe, request ID và response envelope. Dữ liệu thành công dùng trường data; danh sách có thể có meta; lỗi dùng error gồm code, message, details và requestId. Cách đóng gói này giúp Web và Mobile xử lý nhất quán loading, validation và thông báo lỗi.")

    add_text(doc, "KẾT QUẢ TRIỂN KHAI CHỨC NĂNG", style="Heading 1", bold=True, after=6)
    add_text(doc, "4.1. Web frontend", style="Mục lớn", bold=True)
    add_text(doc, "Web cung cấp 18 nhóm màn hình điều hành. Các màn hình sử dụng API thực, có loading, empty, error và trạng thái stale hoặc forbidden ở các vùng dữ liệu quan trọng. Điều hướng sử dụng URL và query state để giữ context như visitId, handoverId, gatePassId, truckVisitId và các bộ lọc an toàn.")
    add_table(doc, ["Nhóm", "Các chức năng chính"], [
        ["Tổng quan và việc ca", "Dashboard, work queue, cảnh báo lưu bãi, thông báo"],
        ["Hàng hóa và cổng", "Manifest, MBL HBL, Container 360, Movement Order, Truck Visit, Gate in"],
        ["Bãi", "Sơ đồ 2D, slot, xếp vị trí thủ công, đảo chuyển, kiểm định, booking"],
        ["Tài chính và ra cổng", "Service order, invoice, payment, readiness, Gate Pass, Gate out"],
        ["Tích hợp và quản trị", "Handover, Partner API, EDI, danh mục, người dùng, quyền, audit, báo cáo"],
    ], [4.2, 12.2])
    add_text(doc, "4.2. Mobile app", style="Mục lớn", bold=True)
    add_text(doc, "Mobile hiện có 17 screen files. Người dùng thực địa thao tác theo nhóm Cổng, Bãi, Giám định, Tra cứu và Việc ca; OPERATOR có thêm Điều phối. Mobile sử dụng cache đọc theo user, API và ICD với thời hạn khác nhau; khi offline, dữ liệu cache được gắn nhãn và các thao tác ghi Gate in hoặc Gate out bị khóa để không tạo giao dịch ngoại tuyến không kiểm soát.")
    add_table(doc, ["Nhóm màn", "Chức năng"], [
        ["Cổng", "Scan hoặc nhập QR, Gate in form, Gate out confirm, thông báo kết quả"],
        ["Bãi", "Yard home, assignment, operations, booking, giám định"],
        ["Tra cứu", "Tìm container, Container detail, readiness, timeline"],
        ["Việc ca", "Danh sách task theo quyền và điều hướng đến nghiệp vụ"],
        ["Tài khoản", "Thông tin user, theme, logout và các trạng thái zero tab"],
    ], [4.2, 12.2])

    add_text(doc, "API VÀ HỢP ĐỒNG JSON", style="Heading 1", bold=True, after=6)
    add_text(doc, "5.1. Thống kê", style="Mục lớn", bold=True)
    add_text(doc, "Danh mục hiện tại ghi nhận 38 controller đăng ký, 22 nhóm module API, 205 endpoint và 119 DTO. Đây là danh mục endpoint từ controller thực tế, không phải tỷ lệ bao phủ giao diện. Các endpoint được chia theo các module nghiệp vụ chính và dùng contract JSON thống nhất.")
    add_table(doc, ["Nhóm API", "Số endpoint", "Mục đích"], [
        ["Yard", "29", "Block, slot, assign, movement, inspection, booking, recommendation"],
        ["Master data", "27", "Shipping line, consignee, transporter, tariff, service type và danh mục"],
        ["Partner handover", "25", "Handover nội bộ và API đối tác"],
        ["Billing", "23", "Dịch vụ, invoice, payment, allocation"],
        ["Manifest", "14", "Manifest, MBL, HBL và container liên quan"],
        ["EDI", "14", "Route, outbox, acknowledgement, alert"],
        ["Các module còn lại", "73", "Auth, cổng, pass, container, report, role, setting, user, work queue, audit"],
    ], [4.5, 3.0, 8.9])
    add_text(doc, "5.2. Ví dụ contract", style="Mục lớn", bold=True)
    add_text(doc, "Ví dụ đăng nhập trả về data gồm accessToken, refreshToken, tokenType, expiresIn, refreshExpiresIn và user. Các lỗi validation trả về error.details.fields với tên field và danh sách constraint. Contract này được Mobile và Web unwrap tại API boundary, tránh dùng dữ liệu mạng chưa kiểm kiểu trực tiếp trong component.")

    add_text(doc, "CƠ SỞ DỮ LIỆU VÀ PHÂN QUYỀN", style="Heading 1", bold=True, after=6)
    add_text(doc, "6.1. Thiết kế dữ liệu", style="Mục lớn", bold=True)
    add_text(doc, "Schema Prisma hiện có 53 model và 41 enum. Dữ liệu được chia theo các nhóm site và auth, master data, manifest, container và hold, lệnh và xe, yard, billing, gate pass, EDI, handover và notification. Các bảng dùng khóa UUID, timestamp UTC và kiểu Decimal cho trọng lượng hoặc tiền, từ đó tránh sai lệch khi tính nghiệp vụ.")
    add_table(doc, ["Nhóm dữ liệu", "Số bảng", "Ví dụ"], [
        ["Site, auth và RBAC", "8", "User, Role, Permission, AuthSession, IcdSetting"],
        ["Master data", "4", "ShippingLine, Consignee, ClearingAgent, Transporter"],
        ["Manifest và container", "7", "Manifest, MasterBl, HouseBl, Container, ContainerVisit, Event, Hold"],
        ["Lệnh, xe và cổng", "4", "MovementOrder, TruckVisit, TruckVisitContainer, Reception"],
        ["Yard", "6", "YardBlock, YardSlot, location log, movement, inspection, booking"],
        ["Billing và Gate Pass", "9", "ServiceOrder, Invoice, Payment, Allocation, GatePass"],
        ["EDI, handover và notification", "12", "Outbox, acknowledgement, API client/log, handover, device/delivery"],
        ["Audit và recommendation", "3", "AuditLog, YardRecommendation, Candidate"],
    ], [4.8, 2.4, 9.2])
    add_text(doc, "6.2. Phân quyền", style="Mục lớn", bold=True)
    add_text(doc, "Hệ thống có 7 role mặc định với 56 permission. Permission guard yêu cầu quyền phù hợp ở backend; giao diện chỉ hiển thị action khi người dùng có quyền, đồng thời giữ nội dung đọc ở trường hợp read only. ADMIN được seed đầy đủ quyền một cách tường minh, không dựa vào bypass ngầm.")
    add_table(doc, ["Role", "Trách nhiệm chính"], [
        ["ADMIN", "Quản trị toàn hệ thống, dữ liệu, quyền, vận hành và tích hợp"],
        ["MANAGER", "Giám sát, yard, handover, report, audit, EDI đọc và xử lý alert"],
        ["OPERATOR", "Điều phối container, work queue, yard và dữ liệu vận hành theo quyền"],
        ["GATE_STAFF", "Truck visit, Gate in và sử dụng Gate Pass tại cổng"],
        ["YARD_STAFF", "Xếp vị trí, đảo chuyển, kiểm định và booking"],
        ["AGENT", "Phạm vi đại lý theo contract và dữ liệu được cấu hình"],
        ["CONSIGNEE", "Phạm vi chủ hàng theo contract và dữ liệu được cấu hình"],
    ], [3.5, 12.5])

    add_text(doc, "KẾT QUẢ KIỂM THỬ E2E BAN ĐẦU", style="Heading 1", bold=True, after=6)
    add_text(doc, "7.1. Môi trường và nguyên tắc", style="Mục lớn", bold=True)
    add_text(doc, "Kiểm thử được thực hiện trên môi trường Local MySQL icd_management, API cổng 3000 và Web cổng 5173. Dữ liệu tạo mới mang nhãn E2E-20261004-2147 để dễ truy vết. Các tuyến EDI Local dùng mock://local và SMTP không được cấu hình, vì vậy không phát sinh email hay gửi thông điệp tới đối tác ngoài môi trường.")
    add_text(doc, "7.2. Các case đã xác nhận", style="Mục lớn", bold=True)
    add_table(doc, ["Case", "Thao tác chính", "Kết quả"], [
        ["AUTH-01", "Đăng nhập ADMIN Local", "PASS: Dashboard tải dữ liệu qua API 3000"],
        ["MAN-01 đến MAN-03", "Tạo Manifest, MBL, HBL, submit", "PASS: liên kết chứng từ và trạng thái SUBMITTED đúng"],
        ["CON-01 và MO-01", "Tạo Container Visit và authorize Movement Order", "PASS: PENDING chuyển đủ điều kiện vào cổng"],
        ["TRUCK-01 và GATE-IN-01", "Tạo xe, ARRIVED, Gate in", "PASS: visit IN_YARD sau khi đối chiếu seal và trọng lượng"],
        ["YARD-01 đến YARD-04", "Xếp slot, đảo chuyển, inspection PASS, booking", "PASS: ghi nhận đúng vị trí, lịch sử và kết quả tác nghiệp"],
        ["READY-01", "Kiểm readiness trước billing", "PASS: chặn vì chưa có hồ sơ tính phí và dịch vụ chưa tính phí"],
        ["HOLD-01", "Tạo Customs Hold", "PASS: readiness thêm blocker Operational Hold"],
        ["HOLD-02", "Release Hold trên browser kiểm thử", "Chưa xác nhận: UI dùng prompt nhập lý do, cần kiểm tiếp bằng phiên tương tác phù hợp"],
    ], [3.0, 7.2, 5.8])
    add_text(doc, "7.3. Phần cần kiểm tiếp", style="Mục lớn", bold=True)
    add_bullets(doc, [
        "Hoàn tất Release Hold rồi xác nhận readiness được cập nhật lại.",
        "Tạo Service Order, issue Invoice, record Payment nội bộ và kiểm tra số dư từ backend.",
        "Cấp Gate Pass, kiểm token QR, Gate out, trạng thái EXITED, Gate Pass USED và đóng vị trí bãi.",
        "Kiểm thử luồng Handover Partner, EDI outbox và notification trên fixture Local.",
        "Khởi động Mobile với API Local và kiểm các thao tác hiện trường trên Expo hoặc Android emulator.",
    ])

    add_text(doc, "VẤN ĐỀ GHI NHẬN VÀ KẾ HOẠCH TUẦN TIẾP THEO", style="Heading 1", bold=True, after=6)
    add_text(doc, "8.1. Giới hạn hiện tại", style="Mục lớn", bold=True)
    add_text(doc, "Một số nghiệp vụ có API nhưng chưa có đủ thao tác trên Web hoặc Mobile. Ví dụ, Reports UI hiện tổng hợp các collection nhưng chưa nối đủ toàn bộ endpoint report hoặc nút export XLSX; Mobile chưa có màn billing và payment riêng. Role AGENT và CONSIGNEE chưa có liên kết User với công ty trong backend, do đó phạm vi portal khách hàng chưa thể kết luận hoàn chỉnh. Các hạn chế này được giữ minh bạch để tránh coi danh mục API là chức năng giao diện đã nghiệm thu.")
    add_text(doc, "8.2. Kế hoạch thực hiện", style="Mục lớn", bold=True)
    add_table(doc, ["Ưu tiên", "Công việc", "Tiêu chí hoàn thành"], [
        ["P1", "Hoàn thiện E2E Billing, Gate Pass và Gate out", "Visit EXITED, pass USED, vị trí bãi đóng và dữ liệu persisted"],
        ["P1", "Kiểm thử Mobile với API Local", "Gate, Yard, inspection, lookup và work queue có evidence runtime"],
        ["P2", "Kiểm thử RBAC theo từng role", "Không lộ action trái quyền và backend từ chối request không đủ quyền"],
        ["P2", "Đối chiếu report, EDI, handover", "Dữ liệu hiển thị trung thực, retry và error state rõ ràng"],
        ["P3", "Bổ sung tài liệu triển khai và checklist nghiệm thu", "README, test log và hướng dẫn khởi động nhất quán"],
    ], [2.0, 8.4, 5.6])

    add_text(doc, "KẾT LUẬN", style="Heading 1", bold=True, after=6)
    add_text(doc, "Trong tuần 6, hệ thống ICD đã có nền tảng tài liệu kỹ thuật thống nhất và chuỗi nghiệp vụ cốt lõi đã được đối chiếu từ lược khai tới vận hành bãi. Kết quả E2E ban đầu xác nhận các bước tạo chứng từ, tạo container, lệnh vận chuyển, tiếp nhận cổng, xếp bãi, đảo chuyển, kiểm định và booking làm việc với dữ liệu Local. Giai đoạn tiếp theo tập trung khép kín phần billing, Gate Pass và Gate out, sau đó mở rộng kiểm thử Mobile và phân quyền theo role.")

    add_text(doc, "TÀI LIỆU ĐỐI CHIẾU", style="Tiêu đề", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, before=14)
    for source in [
        "README.md - ICD Management System Tài liệu tổng hợp hệ thống.",
        "ICD Business Spec v1.7 MySQL - đặc tả nghiệp vụ.",
        "ICD Web Spec v1.7 MySQL và ICD Mobile Spec v1.7 MySQL.",
        "ICD Database Design v1.7 MySQL và ICD Partner Handover API Spec v1.7 MySQL.",
        "Kết quả kiểm thử Local lưu tại test-artifacts/2026-10-04-local-business-e2e.",
    ]:
        add_text(doc, source, style="normal")

    # Replace stale header and footer retained from the source template.
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            p.clear()
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("Đồ án 4  Hệ thống quản lý ICD")
        hr.font.size = Pt(9)
        for p in sec.footer.paragraphs:
            p.clear()
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("Trang ")
        fr.font.size = Pt(9)
        add_page_field(fp)

    doc.core_properties.title = "Báo cáo tiến độ tuần 6 Hệ thống quản lý ICD"
    doc.core_properties.subject = "Đồ án 4"
    doc.core_properties.comments = "Nội dung tổng hợp từ README và tài liệu ICD"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
